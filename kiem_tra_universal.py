"""Tự kiểm tra LỚP BẢO HIỂM PHỦ 100% — chạy: python kiem_tra_universal.py

Script này KHÔNG cần engine gốc: nó dựng thẳng một file .docx chứa các khối lạ đã được
TeX dựng, để bạn mở ra xem ngay có giống bản gốc không.  Nếu truyền thêm đường dẫn một
file .tex thì nó kiểm tra chính file đó.

    python kiem_tra_universal.py                      # dùng mẫu có sẵn
    python kiem_tra_universal.py de_cua_toi.tex       # dùng file của bạn
    python kiem_tra_universal.py de_cua_toi.tex safe  # ép chính sách an toàn tuyệt đối
"""
from __future__ import annotations

import json
import os
import sys
import time
import traceback
from pathlib import Path

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

MAU = r"""\documentclass[12pt,a4paper]{article}
\usepackage[utf8]{vietnam}
\usepackage{amsmath,amssymb}
\usepackage{tikz,pgfplots,tkz-tab}
\usetikzlibrary{arrows.meta,patterns}
\pgfplotsset{compat=1.18}
\newcommand{\bonpa}[4]{\par\noindent\textbf{A.} #1\quad\textbf{B.} #2\quad\textbf{C.} #3\quad\textbf{D.} #4\par}
\begin{document}
\section*{Kiểm thử cấu trúc lạ}

\textbf{Câu 1.} Bảng biến thiên vẽ tay bằng TikZ (không phải tkz-tab):
\begin{center}
\begin{tikzpicture}[scale=0.9,>=stealth]
  \draw[thick] (0,0) rectangle (10,3.4);
  \draw[thick] (0,2.7) -- (10,2.7);
  \draw[thick] (1.4,0) -- (1.4,3.4);
  \node at (0.7,3.05) {$x$};  \node at (0.7,1.35) {$f(x)$};
  \node at (2.4,3.05) {$-2$}; \node at (5,3.05) {$1$}; \node at (8.6,3.05) {$3$};
  \node (a) at (2.4,0.6) {$-1$}; \node (b) at (5,2.2) {$4$}; \node (c) at (8.6,0.9) {$0$};
  \draw[dashed] (2.4,2.7) -- (a); \draw[dashed] (5,2.7) -- (b); \draw[dashed] (8.6,2.7) -- (c);
  \draw[->,thick] (a) -- (b);  \draw[->,thick] (b) -- (c);
\end{tikzpicture}
\end{center}
Số điểm cực trị của $y=|f(x)|$ là
\bonpa{$2$.}{$3$.}{$4$.}{$5$.}

\textbf{Câu 2.} Đồ thị pgfplots:
\begin{center}
\begin{tikzpicture}
\begin{axis}[axis lines=middle,width=7cm,height=5cm,xmin=-3,xmax=3,ymin=-3,ymax=3]
\addplot[thick,blue,domain=-2.2:2.2,samples=120]{x^3-3*x};
\end{axis}
\end{tikzpicture}
\end{center}
\bonpa{$1$.}{$2$.}{$3$.}{$4$.}

\textbf{Câu 3.} Môi trường chưa ai khai báo bao giờ:
\begin{bangxetdaucuatoi}
\hang{$x$}{$-\infty$}{$0$}{$+\infty$}
\end{bangxetdaucuatoi}

\end{document}
"""


def main() -> int:
    print("=" * 70)
    print(" TỰ KIỂM TRA LỚP BẢO HIỂM PHỦ 100% ".center(70, "="))
    print("=" * 70)
    args = [a for a in sys.argv[1:] if a]
    if len(args) > 1 and args[1] in ("safe", "balanced", "native_first"):
        os.environ["L2W_TRUTH_POLICY"] = args[1]
    try:
        from backend.structure_registry import registry_size
        from backend.tex_truth import describe_toolchain
        from backend.universal_fallback import protect, restore, scan_coverage
    except Exception as exc:
        print(f"✗ Không nạp được module mới: {type(exc).__name__}: {exc}")
        traceback.print_exc()
        return 1

    tc = describe_toolchain()
    print(f"Trình biên dịch TeX : {tc['tex_engine'] or '✗ CHƯA CÓ — hãy cài MiKTeX hoặc TeX Live'}")
    print(f"Định dạng ảnh       : {tc['best_vector_format'].upper()}"
          f"   (dvisvgm={bool(tc.get('dvisvgm'))}, pdftocairo={bool(tc.get('pdftocairo'))}, "
          f"inkscape={bool(tc.get('inkscape'))}, soffice={bool(tc.get('soffice'))})")
    rs = registry_size()
    print(f"Sổ đăng ký cấu trúc : {rs['environments']} môi trường + {rs['commands']} lệnh")
    print(f"Chính sách          : {os.environ.get('L2W_TRUTH_POLICY', 'balanced')}")
    print("-" * 70)

    src = Path(args[0]).read_text(encoding="utf-8-sig", errors="replace") if args else MAU
    nguon = args[0] if args else "(mẫu dựng sẵn)"
    print(f"Nguồn: {nguon}  —  {len(src):,} ký tự")

    rep = scan_coverage(src)
    d = rep.to_dict()
    print(f"\n── ĐỘ PHỦ ──")
    print(f"  Tổng số cấu trúc              : {d['tong_so_cau_truc']}")
    print(f"  Engine gốc dựng (sửa được)    : {d['dung_engine_goc']}")
    print(f"  Macro của bạn                 : {d['macro_nguoi_dung']}")
    print(f"  Giao cho TeX dựng (ảnh vector): {d['dung_anh_tex']}  ({d['so_khoi_chup_anh']} khối)")
    print(f"  Chưa biết                     : {d['chua_biet']}  {d['lenh_la'] or ''} {d['moi_truong_la'] or ''}")
    print(f"  → PHỦ NGUYÊN BẢN {d['do_phu_nguyen_ban_phan_tram']}%   |   PHỦ TỔNG {d['do_phu_tong_phan_tram']}%")

    out = ROOT / "kiem_tra_universal_ket_qua.docx"
    print(f"\n── DỰNG THỬ ──")
    t0 = time.time()
    _marked, ctx = protect(src, workdir=ROOT / "_kiem_tra_truth")
    print(f"  Dựng ảnh bằng TeX : {time.time() - t0:.2f}s — {ctx.reason}")
    for w in ctx.warnings[:10]:
        print(f"    ⚠ {w}")
    if not ctx.blocks:
        print("  (không có khối nào cần TeX dựng — nghĩa là engine gốc tự lo được hết)")
        print("=" * 70)
        return 0

    try:
        import docx
        from backend.docx_media import make_marker
        doc = docx.Document()
        doc.add_heading("Kiểm thử các khối do TeX dựng", level=1)
        for i, info in ctx.blocks.items():
            doc.add_paragraph(f"Khối {i + 1} — {info['label']}:")
            doc.add_paragraph(make_marker(i))
        doc.save(str(out))
        st = restore(out, ctx)
        print(f"  Chèn vào Word     : {st['so_anh_chen']}/{len(ctx.blocks)} ảnh"
              f"{', hỏng ' + str(st['so_khoi_hong']) if st['so_khoi_hong'] else ''}")
        for w in st.get("canh_bao", [])[:6]:
            print(f"    ⚠ {w}")
    except Exception as exc:
        print(f"  ✗ {type(exc).__name__}: {exc}")
        traceback.print_exc()
        return 1

    ok = st["so_anh_chen"] == len(ctx.blocks) and st["so_khoi_hong"] == 0
    print("\n" + "=" * 70)
    print(("✓ ĐẠT — mở file sau và so với bản PDF gốc:" if ok else
           "⚠ CHƯA ĐẠT HẾT — xem cảnh báo bên trên; file vẫn tạo được:"))
    print(f"   {out}")
    print("=" * 70)
    return 0 if ok else 2


if __name__ == "__main__":
    raise SystemExit(main())
