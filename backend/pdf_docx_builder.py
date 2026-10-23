"""`backend/pdf_docx_builder.py` — dựng .docx từ manifest khối của `pdf_reader`.

* chữ → đoạn Word (giữ đậm/nghiêng/cỡ); toán nhận dạng được → OMML qua `omml_math`;
* hình, bảng khó, công thức không chắc, trang scan → ảnh vùng cắt từ CHÍNH PDF gốc (SVG + PNG = một đối tượng);
* bảng kẻ đều → bảng Word (gộp ô theo lưới), vai trò ghi bằng metadata `l2w:role=…;block=…`;
* trắc nghiệm → bảng phương án không viền theo đúng số cột đọc được từ PDF (answer_key_unknown);
* hình đặt cạnh chữ → bảng dàn trang 2 cột (layout_immini);
* mỗi trang PDF bắt đầu một trang Word mới (page-aligned) để đối chiếu hiển thị theo trang.
Mọi ảnh mang `descr = "l2w:<loại>:<block_id>"` để kiểm định đối chiếu từng khối.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

from .docx_media import MediaInjector
from .ooxml_tools import set_table_role

__all__ = ["build_docx_from_manifest"]

EMU_PER_PT = 12700


class _Builder:
    def __init__(self, manifest: dict[str, Any]) -> None:
        self.m = manifest
        self.doc = docx.Document()
        st = self.doc.styles["Normal"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(12)
        st.paragraph_format.space_after = Pt(2)
        self.inj = MediaInjector(self.doc)
        self.placed: dict[str, str] = {}
        self.text_width_pt = 470.0
        self.counts = {"paragraphs": 0, "omml": 0, "images": 0, "tables": 0, "mcq": 0, "layouts": 0, "notices": 0}

    # ── trang ──
    def setup_page(self) -> None:
        pages = self.m["pages"]
        if not pages:
            return
        p0 = pages[0]
        sec = self.doc.sections[0]
        sec.page_width = Pt(p0["width"])
        sec.page_height = Pt(p0["height"])
        xs0 = [b["bbox"][0] for b in self.m["blocks"] if b["type"] not in ("page",) and b.get("bbox")]
        xs1 = [b["bbox"][2] for b in self.m["blocks"] if b["type"] not in ("page",) and b.get("bbox")]
        ys0 = [b["bbox"][1] for b in self.m["blocks"] if b["type"] not in ("page",) and b.get("bbox")]
        ys1 = [b["bbox"][3] for b in self.m["blocks"] if b["type"] not in ("page",) and b.get("bbox")]
        left = min(xs0) if xs0 else 56
        right = p0["width"] - max(xs1) if xs1 else 56
        top = min(ys0) if ys0 else 56
        bottom = p0["height"] - max(ys1) if ys1 else 56
        clamp = lambda v: max(28.0, min(90.0, v - 4))  # noqa: E731
        sec.left_margin, sec.right_margin = Pt(clamp(left)), Pt(clamp(right))
        sec.top_margin, sec.bottom_margin = Pt(clamp(top)), Pt(clamp(bottom))
        if any(b["type"] == "page" and b["method"] == "page_image" for b in self.m["blocks"]):
            # trang scan dùng ảnh nguyên trang: lề nhỏ để ảnh không bị thu quá mức
            sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(20)
        self.left_pt = int(sec.left_margin) / EMU_PER_PT
        self.page_w_pt = float(p0["width"])
        self.text_width_pt = (int(sec.page_width) - int(sec.left_margin) - int(sec.right_margin)) / EMU_PER_PT
        self.text_height_pt = (int(sec.page_height) - int(sec.top_margin) - int(sec.bottom_margin)) / EMU_PER_PT

    # ── tiện ích ──
    def _image_run(self, paragraph, asset: dict[str, Any], block_id: str, kind: str, *, max_w: float | None = None,
                   max_h: float | None = None):
        png = Path(asset["png"]).read_bytes() if asset.get("png") else b""
        svg = Path(asset["svg"]).read_bytes() if asset.get("svg") else b""
        w, h = float(asset.get("width_pt") or 50), float(asset.get("height_pt") or 20)
        mw = max_w or self.text_width_pt
        mh = max_h or getattr(self, "text_height_pt", 700)
        r = min(1.0, mw / w if w else 1, mh / h if h else 1)
        w, h = w * r, h * r
        drawing = self.inj.build_drawing(main_data=svg or png, main_ext="svg" if svg else "png", png_data=png if svg else b"",
                                         cx=int(w * EMU_PER_PT), cy=int(h * EMU_PER_PT), name=f"L2W {block_id}",
                                         descr=f"l2w:{kind}:{block_id}")
        run = paragraph.add_run()
        run._r.append(drawing)
        self.counts["images"] += 1
        return run

    def _runs(self, paragraph, runs: list[dict[str, Any]], *, default_size: float | None = None,
              block_id: str = "") -> None:
        from .omml_math import latex_to_omml_element
        for r in runs:
            if r.get("join"):
                continue
            if r.get("math") is not None and "math" in r:
                try:
                    paragraph._p.append(latex_to_omml_element(r["math"], display=False))
                    self.counts["omml"] += 1
                except Exception:
                    rr = paragraph.add_run(f"[L2W: công thức lỗi dựng: {r['math']}]")
                    self.counts["notices"] += 1
                continue
            if r.get("image"):
                self._image_run(paragraph, r["image"], r.get("image_id") or block_id,
                                "math" if r.get("math_image") else "glyph", max_w=self.text_width_pt)
                continue
            if r.get("math_image") or r.get("glyph_image"):
                paragraph.add_run("[L2W: vùng toán chưa cắt được ảnh]")
                self.counts["notices"] += 1
                continue
            txt = r.get("text", "")
            if not txt:
                continue
            run = paragraph.add_run(txt)
            if r.get("bold"):
                run.bold = True
            if r.get("italic"):
                run.italic = True
            size = r.get("size") or default_size
            if size:
                run.font.size = Pt(max(5.0, min(40.0, float(size))))

    # ── khối ──
    def paragraph_block(self, container, b: dict[str, Any]) -> None:
        lines = b.get("lines") or []
        if b["type"] == "math":
            from .omml_math import latex_to_omml_element
            for latex in b.get("latex") or b.get("inline_math") or []:
                p = container.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p._p.append(latex_to_omml_element(latex, display=True))
                self.counts["omml"] += 1
            self.placed[b["block_id"]] = "omml_paragraph"
            return
        p = container.add_paragraph()
        if b.get("align") == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for i, part in enumerate(lines):
            if i > 0 and not (lines[i - 1]["runs"] and lines[i - 1]["runs"][-1].get("join")):
                p.add_run(" ")
            self._runs(p, part["runs"], default_size=b.get("size"), block_id=b["block_id"])
        if b["type"] == "heading":
            for r in p.runs:
                r.bold = True
        if b.get("ocr"):
            p.add_run(b.get("text", ""))
        self.counts["paragraphs"] += 1
        self.placed[b["block_id"]] = "paragraph"

    def image_block(self, container, b: dict[str, Any]) -> None:
        asset = b.get("assets") or {}
        kind = {"figure": "fig", "image": "img", "table": "table", "page": "page", "math": "math"}.get(b["type"], "img")
        if not asset.get("png") and not asset.get("svg"):
            self.notice(container, b, "không cắt được ảnh vùng")
            return
        p = container.add_paragraph()
        self._align_like_source(p, b)
        self._image_run(p, asset, b["block_id"], kind)
        self.placed[b["block_id"]] = "image"

    def _align_like_source(self, p, b: dict[str, Any]) -> None:
        """Căn ảnh như trên trang nguồn: giữa nếu tâm khối gần tâm vùng chữ; ngược lại canh trái + thụt đầu dòng."""
        bb = b.get("bbox")
        left = getattr(self, "left_pt", 56.0)
        if not bb or b["type"] == "page":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return
        center_text = left + self.text_width_pt / 2
        cx = (bb[0] + bb[2]) / 2
        if abs(cx - center_text) < 0.08 * self.text_width_pt:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            indent = max(0.0, min(bb[0] - left, self.text_width_pt - (bb[2] - bb[0])))
            if indent > 2:
                p.paragraph_format.left_indent = Pt(indent)

    def notice(self, container, b: dict[str, Any], reason: str) -> None:
        p = container.add_paragraph()
        r = p.add_run(f"[L2W: VÙNG CHƯA DỰNG ĐƯỢC — trang {b['page']}, {b['block_id']}: {reason}]")
        r.bold = True
        self.counts["notices"] += 1
        self.placed[b["block_id"]] = "notice"

    def table_block(self, container, b: dict[str, Any]) -> None:
        nr, nc = b["n_rows"], b["n_cols"]
        t = container.add_table(rows=nr, cols=nc)
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = b.get("col_widths_pt")
        if widths and len(widths) == nc:
            from .math_tables import apply_fixed_table_grid
            scale = min(1.0, self.text_width_pt / max(1.0, sum(widths)))
            apply_fixed_table_grid(t, [max(200, int(w * scale * 20)) for w in widths])
        align = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                 "left": WD_ALIGN_PARAGRAPH.LEFT}
        for row in b["rows"]:
            for cell in row:
                r0, c0 = cell["row"], cell["col"]
                tgt = t.cell(r0, c0)
                if cell["rowspan"] > 1 or cell["colspan"] > 1:
                    tgt = tgt.merge(t.cell(r0 + cell["rowspan"] - 1, c0 + cell["colspan"] - 1))
                first = True
                for para in cell["paragraphs"]:
                    p = tgt.paragraphs[0] if first else tgt.add_paragraph()
                    first = False
                    if cell.get("align") in align:
                        p.alignment = align[cell["align"]]
                    self._runs(p, para["runs"], block_id=b["block_id"])
        set_table_role(t._tbl, "data", block=b["block_id"], source="pdf")
        self.counts["tables"] += 1
        self.placed[b["block_id"]] = "table"

    def mcq_block(self, container, b: dict[str, Any]) -> None:
        cols = max(1, min(4, int(b.get("layout_cols") or 1)))
        opts = b["options"]
        rows = (len(opts) + cols - 1) // cols
        t = container.add_table(rows=rows, cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        from .math_tables import set_table_borderless
        set_table_borderless(t)
        for i, o in enumerate(opts):
            cell = t.cell(i // cols, i % cols)
            p = cell.paragraphs[0]
            lr = p.add_run(o["label"] + " ")
            lr.bold = bool(o.get("label_bold"))
            self._runs(p, o["runs"], block_id=b["block_id"])
        set_table_role(t._tbl, "layout_mcq_pdf", block=b["block_id"], cols=cols, answer_key="unknown")
        self.counts["mcq"] += 1
        self.placed[b["block_id"]] = "mcq_table"

    def layout_block(self, container, b: dict[str, Any], by_id: dict[str, dict[str, Any]]) -> None:
        t = container.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        from .math_tables import set_table_borderless
        set_table_borderless(t)
        left, right = t.cell(0, 0), t.cell(0, 1)
        lw = sum((by_id[x]["bbox"][2] - by_id[x]["bbox"][0]) for x in b["children_left"][:1] if x in by_id) or 1
        rw = sum((by_id[x]["bbox"][2] - by_id[x]["bbox"][0]) for x in b["children_right"][:1] if x in by_id) or 1
        frac = rw / (lw + rw)
        for cell, ids, share in ((left, b["children_left"], 1 - frac), (right, b["children_right"], frac)):
            first_par = cell.paragraphs[0]
            for cid in ids:
                child = by_id.get(cid)
                if child is None:
                    continue
                self.render(cell, child, by_id, max_w=self.text_width_pt * share)
            if first_par.text == "" and len(cell.paragraphs) > 1 and not first_par._p.findall(".//*"):
                first_par._p.getparent().remove(first_par._p)
        set_table_role(t._tbl, "layout_immini", block=b["block_id"], source="pdf")
        self.counts["layouts"] += 1
        self.placed[b["block_id"]] = "layout_table"

    def render(self, container, b: dict[str, Any], by_id: dict[str, dict[str, Any]], max_w: float | None = None) -> None:
        method, typ = b["method"], b["type"]
        if typ == "page_number":
            self._page_field_footer()
            self.placed[b["block_id"]] = "footer_page_field"
            return
        if typ == "layout":
            self.layout_block(container, b, by_id)
        elif method == "unresolved":
            self.notice(container, b, b.get("fallback_reason") or "chưa xử lý")
        elif method in ("vector_image", "raster_image", "page_image"):
            if max_w:
                asset = b.get("assets") or {}
                p = container.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                kind = {"figure": "fig", "image": "img", "table": "table", "page": "page", "math": "math"}.get(typ, "img")
                if asset.get("png") or asset.get("svg"):
                    self._image_run(p, asset, b["block_id"], kind, max_w=max_w)
                    self.placed[b["block_id"]] = "image"
                else:
                    self.notice(container, b, "không cắt được ảnh vùng")
            else:
                self.image_block(container, b)
        elif typ == "table" and method == "table_native":
            self.table_block(container, b)
        elif typ == "mcq":
            self.mcq_block(container, b)
        else:
            self.paragraph_block(container, b)

    def _page_field_footer(self) -> None:
        if getattr(self, "_footer_done", False):
            return
        self._footer_done = True
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        footer = self.doc.sections[0].footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        p._p.append(fld)

    def build(self, out: Path) -> dict[str, Any]:
        self.setup_page()
        by_id = {b["block_id"]: b for b in self.m["blocks"]}
        pages = sorted({b["page"] for b in self.m["blocks"]})
        body = self.doc
        for k, pno in enumerate(range(1, self.m["page_count"] + 1)):
            if k > 0:
                pb = body.add_paragraph()
                pb.add_run().add_break(WD_BREAK.PAGE)
            for b in [x for x in self.m["blocks"] if x["page"] == pno]:
                if b.get("parent"):
                    continue
                try:
                    self.render(body, b, by_id)
                except Exception as exc:
                    self.notice(body, b, f"lỗi dựng khối ({type(exc).__name__}: {exc})")
                    b.setdefault("warnings", []).append(f"lỗi dựng: {type(exc).__name__}: {exc}")
                    b["status"] = "unresolved"
        # bỏ đoạn trống đầu tài liệu do python-docx tạo sẵn? (Document() mới không có đoạn nào)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out))
        missing = [b["block_id"] for b in self.m["blocks"] if b["block_id"] not in self.placed and b["type"] != "page"
                   and not (b["type"] == "page" and b.get("method") == "text_native")]
        return {"counts": self.counts, "placed": self.placed, "not_placed": missing, "media": dict(self.inj.stats),
                "pages_built": len(pages)}


def build_docx_from_manifest(manifest: dict[str, Any], out_docx: Path) -> dict[str, Any]:
    return _Builder(manifest).build(Path(out_docx))
