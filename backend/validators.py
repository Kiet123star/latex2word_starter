"""Bộ kiểm định 5 nhóm (`backend/validators.py`) — giữ API cũ, sửa các kết quả gây hiểu nhầm.

Nhóm / Layer:
  A (L1) Nội dung   : token nguồn ↔ Word (độ phủ + thứ tự), mất dấu tiếng Việt, lặp nội dung, chỗ đánh dấu
                      nội bộ lộ ra, lệnh TeX rò rỉ ngoài vùng mã hợp lệ, thông báo [L2W: …] còn trong bản nháp.
  B (L2) Toán       : từng khối toán nguồn được đối chiếu với từng m:oMath theo thứ tự (chữ ký ký tự), khối
                      dựng bằng ảnh tính riêng; KHÔNG suy từ tổng số OMML.
  C (L3) Hình       : hình nguồn (TikZ/axis/includegraphics, trừ bảng biến thiên) ↔ nhóm DrawingML + ảnh LOGIC
                      (SVG+PNG của cùng hình = 1); đối chiếu theo block_id khi pipeline cung cấp manifest.
  D (L4) Bảng       : so theo TỪNG vai trò (phương án, đúng/sai, biến thiên, số liệu), phát hiện bảng rỗng.
  E (L5) Hiển thị   : render Word → PDF, so với CHÍNH PDF đầu vào (hoặc PDF biên dịch từ nguồn TeX gốc).
                      Không có nguồn so sánh → NOT_MEASURED, giá trị null (không còn 95.0 / 92.0 mặc định).

Trạng thái lớp: PASS | WARN | FAIL | SKIP | NOT_MEASURED.  Tổng hợp (`summarize_quality`) KHÔNG coi SKIP /
NOT_MEASURED là đạt: kiểm tra bắt buộc chưa chạy → review_required.
"""
from __future__ import annotations

import difflib
import re
import shutil
import subprocess
import unicodedata
import uuid
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

try:
    import numpy as np
    from PIL import Image, ImageFilter

    HAS_VISION = True
except ImportError:  # pragma: no cover
    HAS_VISION = False

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "wpg": "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
}

# Ngưỡng (không hạ so với bản cũ; đều CAO hơn). Ý nghĩa ghi trong details của từng lớp.
TEXT_RECALL_PASS = 0.97
TEXT_RECALL_FAIL = 0.50
MATH_MATCH_FAIL = 0.50
VISUAL_TEXT_RECALL_PASS = 0.97
DUPLICATE_RATIO_WARN = 0.03


@dataclass
class LayerResult:
    """Kết quả kiểm tra từng lớp."""

    layer: int
    name: str
    status: str  # PASS | WARN | FAIL | SKIP | NOT_MEASURED
    summary: str
    details: dict[str, Any] = field(default_factory=dict)
    group: str = ""
    required: bool = True

    def to_dict(self) -> dict:
        return asdict(self)


# ───────────────────────────── trích nguồn TeX ─────────────────────────────
_VERB_ENVS = r"verbatim\*?|Verbatim\*?|lstlisting|minted|comment"
_MATH_ENVS = r"equation\*?|align\*?|gather\*?|multline\*?|eqnarray\*?|flalign\*?|alignat\*?|displaymath|math"
_FIG_ENVS = r"tikzpicture|pspicture|picture|circuitikz|tikzcd"
_NON_RENDERED_ARG_CMDS = (
    "label", "ref", "eqref", "pageref", "cite", "usepackage", "documentclass", "input", "include", "includegraphics",
    "vspace", "hspace", "setlength", "addtolength", "color", "pagestyle", "thispagestyle", "hypersetup", "bibliography",
    "bibliographystyle", "graphicspath", "usetikzlibrary", "pgfplotsset", "tikzset", "newcommand", "renewcommand",
    "providecommand", "newenvironment", "renewenvironment", "def", "let", "setcounter", "addtocounter", "linespread",
    "geometry", "newif", "definecolor", "fontsize", "rule", "includeonly", "subfile", "import",
    "subimport", "phantom", "hphantom", "vphantom", "url", "newtheorem", "setstretch", "rowcolor", "cellcolor",
    "columncolor", "arrayrulecolor", "pagenumbering", "enlargethispage", "addvspace",
)


#: lệnh có N đối số đầu là đặc tả/kích thước (không hiển thị); đối số còn lại là nội dung
_SPEC_ARG_CMDS = {"multirow": 2, "multicolumn": 2, "cline": 1, "hhline": 1, "parbox": 1, "raisebox": 1,
                  "rotatebox": 1, "scalebox": 1, "resizebox": 2, "makebox": 0, "framebox": 0}
_SPEC_ARG_ENVS = {"minipage": 1, r"multicols\*?": 1, "thebibliography": 1, "wrapfigure": 2, "wraptable": 2,
                  "adjustwidth": 2, "spacing": 1, r"tabular\*": 1, "tabularx": 1}


def _mask_comments_verbatim(tex: str) -> tuple[str, list[str]]:
    """Bỏ chú thích; trả về (văn bản, danh sách nội dung verbatim) — verbatim được thay bằng chỗ trống."""
    verb: list[str] = []

    def keep_verb(m: re.Match) -> str:
        verb.append(m.group(0))
        return " "

    tex = re.sub(r"\\begin\{(" + _VERB_ENVS + r")\}.*?\\end\{\1\}", keep_verb, tex, flags=re.S)
    tex = re.sub(r"\\verb\*?([^A-Za-z\s*])(.*?)\1", keep_verb, tex)
    tex = re.sub(r"(?<!\\)%.*$", "", tex, flags=re.M)
    return tex, verb


def _body(tex: str) -> str:
    m = re.search(r"\\begin\{document\}", tex)
    if not m:
        return tex
    e = re.search(r"\\end\{document\}", tex[m.end():])
    return tex[m.end(): m.end() + e.start()] if e else tex[m.end():]


def _find_brace_end(s: str, i: int) -> int:
    depth = 0
    while i < len(s):
        c = s[i]
        if c == "\\":
            i += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return i
        i += 1
    return len(s) - 1


def extract_math_segments(tex: str) -> list[dict[str, Any]]:
    """Các khối toán của thân tài liệu theo thứ tự, NGOÀI hình vẽ (TikZ…). Mỗi khối: {"tex", "display", "pos"}."""
    body, _verb = _mask_comments_verbatim(_body(tex))
    # che hình vẽ: toán bên trong nhãn hình thuộc nhóm hình, không thuộc nhóm toán
    fig_spans = [(m.start(), m.end()) for m in
                 re.finditer(r"\\begin\{(" + _FIG_ENVS + r")\}.*?\\end\{\1\}", body, re.S)]
    out: list[dict[str, Any]] = []
    i = 0
    n = len(body)
    env_re = re.compile(r"\\begin\{(" + _MATH_ENVS + r")\}")
    while i < n:
        if any(a <= i < b for a, b in fig_spans):
            i = next(b for a, b in fig_spans if a <= i < b)
            continue
        c = body[i]
        if c == "\\":
            if body.startswith("\\[", i):
                e = body.find("\\]", i + 2)
                if e > 0:
                    out.append({"tex": body[i + 2: e], "display": True, "pos": i})
                    i = e + 2
                    continue
            if body.startswith("\\(", i):
                e = body.find("\\)", i + 2)
                if e > 0:
                    out.append({"tex": body[i + 2: e], "display": False, "pos": i})
                    i = e + 2
                    continue
            m = env_re.match(body, i)
            if m:
                end_tok = "\\end{" + m.group(1) + "}"
                e = body.find(end_tok, m.end())
                if e > 0:
                    out.append({"tex": body[m.end(): e], "display": True, "pos": i, "env": m.group(1)})
                    i = e + len(end_tok)
                    continue
            i += 2
            continue
        if c == "$":
            if body.startswith("$$", i):
                e = body.find("$$", i + 2)
                if e > 0:
                    out.append({"tex": body[i + 2: e], "display": True, "pos": i})
                    i = e + 2
                    continue
            j = i + 1
            while j < n:
                if body[j] == "\\":
                    j += 2
                    continue
                if body[j] == "$":
                    break
                j += 1
            if j < n:
                seg = body[i + 1: j]
                if seg.strip():
                    out.append({"tex": seg, "display": False, "pos": i})
                i = j + 1
                continue
        i += 1
    return out


def _latex_symbol_maps() -> tuple[dict[str, str], dict[str, str]]:
    try:
        from .omml_math import BLACKBOARD_BOLD, FUNC_NAMES, LATEX_SYMBOLS, LIM_OPS, NARY_OPS
        sym = dict(LATEX_SYMBOLS)
        sym.update(NARY_OPS)
        sym.update(FUNC_NAMES)
        sym.update(LIM_OPS)
        return sym, BLACKBOARD_BOLD
    except Exception:  # pragma: no cover
        return {}, {}


_SYM, _BB = _latex_symbol_maps()


def math_signature(tex: str) -> str:
    """Chuỗi ký tự chữ/số/Hy Lạp của biểu thức theo thứ tự (lệnh → ký hiệu Unicode như omml_math)."""
    s = re.sub(r"\\mathbb\s*\{\s*([A-Z])\s*\}", lambda m: _BB.get(m.group(1), m.group(1)), tex)
    s = re.sub(r"\\(?:text|textbf|textit|mathrm|mathbf|operatorname|mbox)\s*\{([^{}]*)\}", r" \1 ", s)
    s = re.sub(r"\\(?:begin|end)\s*\{[^{}]*\}(?:\{[^{}]*\})?", " ", s)
    s = re.sub(r"\\(?:pmod|bmod|mod)(?![A-Za-z])", " mod ", s)
    s = re.sub(r"\\tag\*?\{([^{}]*)\}", r" \1 ", s)  # Word in số thẻ "(n)" trong công thức
    # (?![A-Za-z]): "\\right" không được ăn tiền tố của "\\rightarrow", "\\rightleftharpoons"…
    s = re.sub(r"\\(?:(?:left|right|big|Big|bigg|Bigg|bigl|bigr|Bigl|Bigr|displaystyle|limits|nolimits|quad|qquad)"
               r"(?![A-Za-z])|label\{[^{}]*\}|color\{[^{}]*\}|hspace\*?\{[^{}]*\}|phantom\{[^{}]*\})", " ", s)

    def cmd(m: re.Match) -> str:
        name = "\\" + m.group(1)
        return " " + _SYM.get(name, "") + " "

    s = re.sub(r"\\([A-Za-z]+)", cmd, s)
    return "".join(ch for ch in unicodedata.normalize("NFC", s) if ch.isalnum())


def _norm_tokens(text: str) -> list[str]:
    text = unicodedata.normalize("NFC", text)
    toks = re.findall(r"[^\W_]+", text, flags=re.UNICODE)
    return [t.lower() for t in toks if len(t) >= 2 or t.isdigit()]


def tex_text_tokens(tex: str) -> tuple[list[str], list[str]]:
    """(token chữ của thân tài liệu ngoài toán/hình, các lệnh xuất hiện trong verbatim)."""
    body, verb = _mask_comments_verbatim(_body(tex))
    body = re.sub(r"\\begin\{(" + _FIG_ENVS + r"|axis)\}.*?\\end\{\1\}", " ", body, flags=re.S)
    # bỏ toán
    body = re.sub(r"\$\$.*?\$\$|\\\[.*?\\\]|\\\(.*?\\\)", " ", body, flags=re.S)
    body = re.sub(r"\\begin\{(" + _MATH_ENVS + r")\}.*?\\end\{\1\}", " ", body, flags=re.S)
    body = re.sub(r"(?<!\\)\$(?:\\.|[^$\\])*\$", " ", body, flags=re.S)
    # \ce{…} (mhchem) được dựng thành công thức OMML → không tính là chữ thường
    body = re.sub(r"\\(?:ce|num|si|unit)\s*\{(?:[^{}]|\{[^{}]*\})*\}", " ", body)
    body = re.sub(r"\\(?:SI|qty)\s*(?:\[[^\]]*\])?\s*\{(?:[^{}]|\{[^{}]*\})*\}\s*\{(?:[^{}]|\{[^{}]*\})*\}", " ", body)
    # đặc tả cột của bảng (không hiển thị)
    body = re.sub(r"\\begin\{(tabular\*?|tabularx|longtable|array)\}\s*(?:\{[^{}]*\})?\s*\{(?:[^{}]|\{[^{}]*\})*\}",
                  " ", body)
    # đối số không hiển thị
    for c in _NON_RENDERED_ARG_CMDS:
        # (?![A-Za-z]) — trước đây "\include" ăn tiền tố của "\includegraphics", "\color" ăn "\colorbox"
        body = re.sub(r"\\" + c + r"(?![A-Za-z@])\*?\s*(?:\[[^\]]*\])?\s*(\{(?:[^{}]|\{[^{}]*\})*\})?"
                      r"(\{(?:[^{}]|\{[^{}]*\})*\})?", " ", body)
    # đối số KÍCH THƯỚC/ĐẶC TẢ không hiển thị: chỉ bỏ N đối số đầu, giữ đối số nội dung
    for cmd, nargs in _SPEC_ARG_CMDS.items():
        body = re.sub(r"\\" + cmd + r"(?![A-Za-z@])\*?\s*(?:\[[^\]]*\])?" + r"\s*\{(?:[^{}]|\{[^{}]*\})*\}" * nargs,
                      " ", body)
    for env, nargs in _SPEC_ARG_ENVS.items():
        body = re.sub(r"\\begin\{" + env + r"\}\s*(?:\[[^\]]*\])?" + r"\s*\{(?:[^{}]|\{[^{}]*\})*\}" * nargs,
                      " ", body)
    body = re.sub(r"\\begin\{tasks\}\s*\(\d+\)", " ", body)
    body = re.sub(r"\\cmidrule\s*(?:\([^)]*\))?\s*\{[^{}]*\}", " ", body)
    body = re.sub(r"\\(?:textcolor|colorbox)\s*(?:\[[^\]]*\])?\s*\{[^{}]*\}", " ", body)
    body = re.sub(r"\\fcolorbox\s*(?:\[[^\]]*\])?\s*\{[^{}]*\}\s*\{[^{}]*\}", " ", body)
    body = re.sub(r"\\\\\s*\[[^\]]*\]", " ", body)
    body = re.sub(r"\\(?:begin|end)\s*\{[^{}]*\}(?:\[[^\]]*\])?", " ", body)
    body = re.sub(r"\\[A-Za-z@]+\*?(?:\[[^\]]*\])?", " ", body)
    body = re.sub(r"\[L2W:[^\]]*\]", " ", body)
    verb_cmds = [c for v in verb for c in re.findall(r"\\[A-Za-z]{2,}", v)]
    return _norm_tokens(body.replace("~", " ")), verb_cmds


def strip_latex_commands(tex: str) -> str:
    """(API cũ) văn bản thuần ước lượng của nguồn."""
    return " ".join(tex_text_tokens(tex)[0])


# ───────────────────────────── đọc DOCX ─────────────────────────────
def _docx_roots(docx_path: Path) -> tuple[ET.Element, list[str]]:
    with zipfile.ZipFile(docx_path, "r") as zf:
        xml_bytes = zf.read("word/document.xml")
        media = [f for f in zf.namelist() if f.startswith("word/media/")]
    return ET.fromstring(xml_bytes), media


def _w_text(root: ET.Element) -> str:
    """Chữ w:t theo ĐOẠN: các run trong một đoạn nối liền (một từ có thể bị tách thành nhiều run do định
    dạng); đoạn cách nhau bằng xuống dòng. Chữ trong công thức (m:t) không tính ở đây."""
    W = NS["w"]
    parent_map = {c: p for p in root.iter() for c in p}
    owner: dict[Any, list[str]] = {}
    order: list[Any] = []
    want = {f"{{{W}}}t", f"{{{W}}}tab", f"{{{W}}}br"}
    for el in root.iter():
        if el.tag not in want:
            continue
        anc = parent_map.get(el)
        while anc is not None and anc.tag != f"{{{W}}}p":
            anc = parent_map.get(anc)
        if anc is None:
            continue
        if anc not in owner:
            owner[anc] = []
            order.append(anc)
        owner[anc].append((el.text or "") if el.tag == f"{{{W}}}t" else " ")
    return "\n".join("".join(owner[p]) for p in order)


def _omath_texts(root: ET.Element) -> list[str]:
    """Văn bản m:t của từng m:oMath ngoài cùng (oMath lồng không tính hai lần)."""
    out = []
    tag = f"{{{NS['m']}}}oMath"
    parent_map = {c: p for p in root.iter() for c in p}
    for om in root.iter(tag):
        p = parent_map.get(om)
        nested = False
        while p is not None:
            if p.tag == tag:
                nested = True
                break
            p = parent_map.get(p)
        if not nested:
            out.append("".join(t.text or "" for t in om.iter(f"{{{NS['m']}}}t")))
    return out


def _images(root: ET.Element) -> list[dict[str, str]]:
    out = []
    for d in root.iter(f"{{{NS['w']}}}drawing"):
        gd = d.find(f".//{{{NS['a']}}}graphicData")
        uri = gd.get("uri") if gd is not None else ""
        dp = d.find(f".//{{{NS['wp']}}}docPr")
        out.append({"kind": "picture" if uri.endswith("/picture") else ("group" if "Group" in uri else
                                                                       ("shape" if "Shape" in uri else "other")),
                    "name": (dp.get("name") or "") if dp is not None else "",
                    "descr": (dp.get("descr") or "") if dp is not None else ""})
    return out


# ───────────────────────────── Layer 1 — Nội dung ─────────────────────────────
def _content_check(src_tokens: list[str], docx_tokens: list[str], *, verb_cmds: list[str] | None,
                   docx_text: str, inventory: dict[str, Any] | None,
                   generated: Counter | None = None) -> dict[str, Any]:
    sc, dc = Counter(src_tokens), Counter(docx_tokens)
    inter = sum((sc & dc).values())
    total = max(1, sum(sc.values()))
    recall = inter / total
    missing = sc - dc
    gen = generated or Counter()
    # token do chính pipeline thêm (nhãn "Câu n.", "Bảng n.", "Lời giải.", nhãn danh sách) không tính là lặp
    excess = sum(max(0, dc[t] - sc[t] - gen.get(t, 0)) for t in sc if dc[t] > sc[t])
    dup_ratio = excess / total
    order = difflib.SequenceMatcher(None, src_tokens, docx_tokens, autojunk=False).ratio() \
        if len(src_tokens) * max(1, len(docx_tokens)) < 9_000_000 else None

    def strip_accents(t: str) -> str:
        return "".join(ch for ch in unicodedata.normalize("NFD", t) if unicodedata.category(ch) != "Mn")

    lost_diacritics = []
    plain_docx = Counter(t for t in docx_tokens if strip_accents(t) == t)
    for t, cnt in missing.items():
        st = strip_accents(t)
        if st != t and plain_docx.get(st, 0) > 0:
            lost_diacritics.append(t)
    residue = re.findall(r"\\[A-Za-z]{2,}", docx_text)
    allowed = Counter(verb_cmds or [])
    rc = Counter(residue)
    residue_out = {k: v - allowed.get(k, 0) for k, v in rc.items() if v - allowed.get(k, 0) > 0}
    leaks = (inventory or {}).get("marker_leaks_total", 0)
    notices = (inventory or {}).get("l2w_notices", [])
    return {
        "source_tokens": total, "docx_tokens": len(docx_tokens), "matched_tokens": inter,
        "token_recall": round(recall, 4), "order_ratio": round(order, 4) if order is not None else None,
        "duplicate_excess_ratio": round(dup_ratio, 4),
        "generated_label_tokens": sum(gen.values()),
        "missing_sample": [t for t, _ in missing.most_common(25)],
        "lost_diacritics_sample": lost_diacritics[:20], "lost_diacritics": len(lost_diacritics),
        "tex_residue": sum(residue_out.values()), "tex_residue_sample": sorted(residue_out)[:20],
        "marker_leaks": leaks, "l2w_notices": notices[:20],
        "metric": "token_recall = |token nguồn ∩ token Word| / |token nguồn| (đa tập, NFC, chữ thường); "
                  "order_ratio = difflib ratio trên dãy token — tín hiệu thứ tự, không phải độ chính xác.",
    }


def _content_status(d: dict[str, Any]) -> tuple[str, str]:
    if d["source_tokens"] >= 20 and d["docx_tokens"] == 0:
        return "FAIL", "Word không có chữ trong khi nguồn có"
    if d["token_recall"] < TEXT_RECALL_FAIL:
        return "FAIL", f"giữ {d['token_recall']:.0%} token nguồn"
    reasons = []
    if d["token_recall"] < TEXT_RECALL_PASS:
        reasons.append(f"thiếu {1 - d['token_recall']:.1%} token")
    if d["duplicate_excess_ratio"] > DUPLICATE_RATIO_WARN:
        reasons.append(f"lặp {d['duplicate_excess_ratio']:.1%}")
    if d["lost_diacritics"]:
        reasons.append(f"{d['lost_diacritics']} từ mất dấu")
    if d["tex_residue"]:
        reasons.append(f"{d['tex_residue']} lệnh TeX rò rỉ")
    if d["marker_leaks"]:
        reasons.append(f"{d['marker_leaks']} chỗ đánh dấu nội bộ lộ ra")
    if d["l2w_notices"]:
        reasons.append(f"{len(d['l2w_notices'])} thông báo [L2W] cần xử lý")
    return ("WARN" if reasons else "PASS"), (", ".join(reasons) or "đủ nội dung chữ")


def run_layer1_text(tex_content: str, root: ET.Element, inventory: dict[str, Any] | None = None, *,
                    normalized_tex: str | None = None) -> LayerResult:
    src_tokens, verb_cmds = tex_text_tokens(tex_content)
    generated = None
    if normalized_tex:
        generated = Counter(tex_text_tokens(normalized_tex)[0]) - Counter(src_tokens)
    w_text = _w_text(root)
    d = _content_check(src_tokens, _norm_tokens(w_text), verb_cmds=verb_cmds, docx_text=w_text, inventory=inventory,
                       generated=generated)
    d["paragraphs"] = len(root.findall(".//w:p", NS))
    status, why = _content_status(d)
    return LayerResult(1, "A. Nội dung (chữ, thứ tự, dấu tiếng Việt, rò rỉ)", status,
                       f"Giữ {d['token_recall']:.1%} token nguồn ({d['matched_tokens']}/{d['source_tokens']}); {why}.",
                       d, group="A")


# ───────────────────────────── Layer 2 — Toán ─────────────────────────────
def match_math(src_sigs: list[str], docx_sigs: list[str], lookahead: int = 0) -> dict[str, Any]:
    """Ghép khối toán nguồn ↔ m:oMath theo THỨ TỰ (difflib trên dãy chữ ký, chịu được OMML thừa như ô bảng
    biến thiên), sau đó ghép phần còn lại theo đa tập ký tự (vd. C^k_n ↔ C_n^k) gần vị trí nhất."""
    sm = difflib.SequenceMatcher(None, src_sigs, docx_sigs, autojunk=False)
    pair: dict[int, int] = {}
    for blk in sm.get_matching_blocks():
        for k in range(blk.size):
            pair[blk.a + k] = blk.b + k
    used = set(pair.values())
    for i, sig in enumerate(src_sigs):
        if i in pair:
            continue
        prev_b = max((pair[x] for x in pair if x < i), default=-1)
        next_b = min((pair[x] for x in pair if x > i), default=len(docx_sigs))
        for k in range(prev_b + 1, next_b):
            if k in used:
                continue
            if sig and sorted(docx_sigs[k]) == sorted(sig):
                pair[i] = k
                used.add(k)
                break
    unmatched = [i for i in range(len(src_sigs)) if i not in pair]
    return {"matched": len(pair), "unmatched_indices": unmatched}


def run_layer2_math(tex_content: str, root: ET.Element, math_images: int = 0) -> LayerResult:
    segs = extract_math_segments(tex_content)
    src_sigs = [math_signature(s["tex"]) for s in segs]
    docx_texts = _omath_texts(root)
    docx_sigs = ["".join(ch for ch in unicodedata.normalize("NFC", t) if ch.isalnum()) for t in docx_texts]
    res = match_math(src_sigs, docx_sigs)
    n = len(segs)
    matched = res["matched"]
    unmatched = res["unmatched_indices"]
    covered = matched + min(math_images, len(unmatched))
    details = {
        "source_math_blocks": n, "omml_equations": len(docx_texts), "matched_native": matched,
        "math_as_image": math_images, "unmatched": len(unmatched) - min(math_images, len(unmatched)),
        "unmatched_sample": [segs[i]["tex"][:120] for i in unmatched[:12]],
        "fractions": len(root.findall(".//m:f", NS)), "radicals": len(root.findall(".//m:rad", NS)),
        "nary_ops": len(root.findall(".//m:nary", NS)), "matrices": len(root.findall(".//m:m", NS)),
        "accents": len(root.findall(".//m:acc", NS)),
        "metric": "Mỗi khối toán nguồn (ngoài hình) được ghép theo thứ tự với một m:oMath có cùng chữ ký ký tự "
                  "(chữ/số/ký hiệu). Khớp chữ ký chưa chứng minh cấu trúc đúng; cấu trúc kiểm bằng fixture.",
    }
    if n == 0:
        return LayerResult(2, "B. Toán (OMML theo từng khối)", "PASS", "Nguồn không có công thức.", details, group="B")
    ratio = covered / n
    status = "PASS" if covered == n else ("FAIL" if ratio < MATH_MATCH_FAIL else "WARN")
    return LayerResult(2, "B. Toán (OMML theo từng khối)", status,
                       f"{matched}/{n} khối toán khớp OMML theo thứ tự"
                       + (f", {math_images} khối giữ bằng ảnh" if math_images else "")
                       + (f"; {details['unmatched']} khối chưa khớp" if details["unmatched"] else "") + ".",
                       details, group="B")


# ───────────────────────────── Layer 3 — Hình ─────────────────────────────
def _source_figures(tex: str) -> dict[str, int]:
    body, _ = _mask_comments_verbatim(_body(tex))
    tikz = re.findall(r"\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}", body, re.S)
    bbt = sum(1 for b in tikz if re.search(r"\\tkzTab(?:Init|Line|Var)", b))
    try:
        from .math_tables import is_tkz_tab_block
        bbt = sum(1 for b in tikz if is_tkz_tab_block(b))
    except Exception:
        pass
    body_no_tikz = re.sub(r"\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}", " ", body, flags=re.S)
    axis_alone = len(re.findall(r"\\begin\{axis\}", body_no_tikz))
    incl = len(re.findall(r"\\includegraphics\*?(?:\[[^\]]*\])?\{[^{}]+\}", body_no_tikz))
    other = len(re.findall(r"\\begin\{(?:pspicture|picture|circuitikz|tikzcd)\}", body))
    return {"tikz": len(tikz), "bbt_in_tikz": bbt, "standalone_axis": axis_alone, "includegraphics": incl,
            "other_picture_envs": other, "expected_figures": len(tikz) - bbt + axis_alone + incl + other}


def run_layer3_illustrations(tex_content: str, root: ET.Element, media_files: list[str],
                             manifest: dict[str, Any] | None = None) -> LayerResult:
    src = _source_figures(tex_content)
    imgs = _images(root)
    groups = sum(1 for i in imgs if i["kind"] == "group")
    pictures = [i for i in imgs if i["kind"] == "picture"]
    fig_pictures = [p for p in pictures if not p["descr"].startswith(("l2w:math", "l2w:bbt", "l2w:table"))]
    produced = groups + len(fig_pictures)
    expected = src["expected_figures"]
    per_block = []
    missing_blocks = []
    if manifest:
        descrs = {p["descr"] for p in pictures}
        for b in manifest.get("figure_blocks", []):
            ok = any(b["block_id"] in d for d in descrs) if b.get("method") in ("vector_image", "raster_image",
                                                                                "page_image") else None
            per_block.append({"block_id": b["block_id"], "method": b.get("method"), "found": ok,
                              "status": b.get("status")})
            if ok is False or b.get("status") == "unresolved":
                missing_blocks.append(b["block_id"])
    details = {**src, "drawingml_shape_groups": groups,
               "drawingml_editable_shapes": len(root.findall(".//wps:wsp", NS)),
               "logical_images": len(pictures), "figure_images": len(fig_pictures),
               "media_parts": len(media_files), "produced_figures": produced, "per_block": per_block[:200],
               "missing_blocks": missing_blocks,
               "note": "Ảnh SVG kèm PNG của cùng một hình được tính là MỘT hình. Nhóm DrawingML tồn tại chưa "
                       "chứng minh hình đúng; đối chiếu nét vẽ cần kiểm tra hiển thị (nhóm E)."}
    if expected == 0 and produced == 0:
        return LayerResult(3, "C. Hình (DrawingML / ảnh logic)", "PASS", "Nguồn không có hình.", details, group="C")
    if expected > 0 and produced == 0:
        status = "FAIL"
    elif produced < expected or missing_blocks:
        status = "WARN"
    else:
        status = "PASS"
    return LayerResult(3, "C. Hình (DrawingML / ảnh logic)", status,
                       f"Nguồn {expected} hình; Word có {groups} nhóm DrawingML + {len(fig_pictures)} ảnh hình"
                       + (f"; THIẾU {expected - produced}" if produced < expected else "")
                       + (f"; {len(missing_blocks)} khối chưa đối chiếu được" if missing_blocks else "") + ".",
                       details, group="C")


# ───────────────────────────── Layer 4 — Bảng ─────────────────────────────
def _source_tables(tex: str) -> dict[str, int]:
    body, _ = _mask_comments_verbatim(_body(tex))
    no_tikz = re.sub(r"\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}", " ", body, flags=re.S)
    mcq = len(re.findall(r"\\(?:choice[ABCD]?|boncot|haicot|motcot|bonpa)(?![A-Za-z])", no_tikz))
    tf = len(re.findall(r"\\choiceTF(?![A-Za-z])", no_tikz))
    data = 0
    mcq_tabx = 0
    for m in re.finditer(r"\\begin\{(tabular\*?|tabularx|longtable)\}(.*?)\\end\{\1\}", no_tikz, re.S):
        inner = m.group(2)
        if m.group(1) == "tabularx" and len(re.findall(r"(?:^|[&\s{])(?:\\\w+\{)*[A-D]\.", inner)) == 4:
            mcq_tabx += 1
        else:
            data += 1
    tikz = re.findall(r"\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}", body, re.S)
    try:
        from .math_tables import is_tkz_tab_block
        bbt = sum(1 for b in tikz if is_tkz_tab_block(b))
    except Exception:
        bbt = sum(1 for b in tikz if "\\tkzTabInit" in b)
    return {"mcq": mcq + mcq_tabx, "true_false": tf, "variation": bbt, "data": data}


def _table_roles(root: ET.Element) -> tuple[Counter, list[dict[str, Any]]]:
    roles: Counter = Counter()
    info = []
    for tbl in root.iter(f"{{{NS['w']}}}tbl"):
        desc = tbl.find(f"{{{NS['w']}}}tblPr/{{{NS['w']}}}tblDescription")
        val = desc.get(f"{{{NS['w']}}}val") if desc is not None else ""
        role = val.split("=", 1)[1].split(";")[0] if val.startswith("l2w:role=") else None
        texts = [("".join(t.text or "" for t in tc.iter(f"{{{NS['w']}}}t"))
                  + "".join(t.text or "" for t in tc.iter(f"{{{NS['m']}}}t"))).strip()
                 for tc in tbl.iter(f"{{{NS['w']}}}tc")]
        has_obj = tbl.find(f".//{{{NS['w']}}}drawing") is not None or tbl.find(f".//{{{NS['m']}}}oMath") is not None
        non_empty = [t for t in texts if t]
        if role is None:
            if len(non_empty) == 4 and all(non_empty[i].startswith(l) for i, l in enumerate(("A.", "B.", "C.", "D."))):
                role = "mcq"
            elif "VarArrow_" in ET.tostring(tbl, encoding="unicode"):
                role = "variation"
            else:
                role = "data"
            src = "heuristic"
        else:
            src = "metadata"
        roles[role] += 1
        info.append({"role": role, "role_source": src, "cells": len(texts), "non_empty_cells": len(non_empty),
                     "empty": not non_empty and not has_obj})
    return roles, info


def run_layer4_tables(tex_content: str, root: ET.Element, figure_images_for_bbt: int = 0) -> LayerResult:
    exp = _source_tables(tex_content)
    roles, info = _table_roles(root)
    produced = {"mcq": roles.get("mcq", 0) + roles.get("layout_mcq_pdf", 0), "true_false": roles.get("true_false", 0),
                "variation": roles.get("variation", 0) + figure_images_for_bbt, "data": roles.get("data", 0)}
    short = {k: exp[k] - produced[k] for k in exp if produced[k] < exp[k]}
    empty = [i for i, t in enumerate(info) if t["empty"]]
    mcq_bad = [i for i, t in enumerate(info) if t["role"] == "mcq" and t["non_empty_cells"] < 4]
    details = {"expected_by_role": exp, "produced_by_role": produced, "docx_tables_by_role": dict(roles),
               "short_by_role": short, "empty_tables": empty, "mcq_tables_missing_options": mcq_bad,
               "tables": info[:200],
               "note": "So theo từng vai trò; bảng biến thiên giữ bằng ảnh TeX (table_as_image) được tính là có "
                       "biểu diễn nhưng không phải bảng native."}
    status = "PASS"
    if short or empty or mcq_bad:
        status = "WARN"
    if sum(exp.values()) > 0 and sum(produced.values()) == 0:
        status = "FAIL"
    summ = ", ".join(f"{k} {produced[k]}/{exp[k]}" for k in exp if exp[k] or produced[k]) or "không có bảng"
    return LayerResult(4, "D. Bảng & bố cục (theo vai trò)", status,
                       f"{summ}" + (f"; {len(empty)} bảng rỗng" if empty else "")
                       + (f"; {len(mcq_bad)} bảng phương án thiếu ô" if mcq_bad else "") + ".", details, group="D")


# ───────────────────────────── Layer 5 — Hiển thị ─────────────────────────────
def run_layer5_visual(visual_info: dict[str, Any] | None) -> LayerResult:
    vi = visual_info or {}
    if vi.get("enabled") and vi.get("measured"):
        reasons = []
        if vi.get("page_count_mismatch") and vi.get("mode") == "pdf":
            reasons.append(f"số trang khác: nguồn {vi.get('reference_pages')} / Word {vi.get('word_pages')}")
        tr = vi.get("text_recall")
        if tr is not None and tr < VISUAL_TEXT_RECALL_PASS:
            reasons.append(f"chữ hiển thị giữ {tr:.1%}")
        status = "WARN" if reasons else "PASS"
        return LayerResult(5, "E. Hiển thị (render Word ↔ PDF nguồn)", status,
                           (f"Trang nguồn {vi.get('reference_pages')} / Word {vi.get('word_pages')}; "
                            f"chữ hiển thị giữ {tr:.1%}" if tr is not None else "Đã render")
                           + (f"; IoU mực (heuristic bố cục) {vi['ink_iou_mean']:.2f}" if vi.get("ink_iou_mean") is not None else "")
                           + ("; " + "; ".join(reasons) if reasons else "") + ".", vi, group="E")
    reason = vi.get("reason") or "Chưa bật --qa-pdf hoặc thiếu công cụ render."
    return LayerResult(5, "E. Hiển thị (render Word ↔ PDF nguồn)", "NOT_MEASURED", reason,
                       {**vi, "similarity_percent": None}, group="E")


def sanitize_latex_for_compilation(tex: str) -> str:
    """Bỏ dòng trống trong \\[...\\] và tùy chọn axis (gây lỗi TeX)."""

    def fix_axis(m: re.Match[str]) -> str:
        cleaned = [l for l in m.group(1).splitlines() if l.strip()]
        return "\\begin{axis}[\n" + "\n".join(cleaned) + "\n]"

    tex = re.sub(r"\\begin\{axis\}\s*\[(.*?)\]", fix_axis, tex, flags=re.S)

    def fix_display_math(m: re.Match[str]) -> str:
        cleaned = [l for l in m.group(1).splitlines() if l.strip()]
        return "\\[\n" + "\n".join(cleaned) + "\n\\]"

    return re.sub(r"\\\[(.*?)\\\]", fix_display_math, tex, flags=re.S)


def run_all_layers(
    tex_content: str,
    docx_path: Path,
    visual_info: dict[str, Any] | None = None,
    *,
    manifest: dict[str, Any] | None = None,
    normalized_tex: str | None = None,
) -> list[dict]:
    """Chạy 5 nhóm kiểm định trên tệp .docx CUỐI CÙNG. Giữ chữ ký cũ (tex, docx, visual_info).

    `normalized_tex` (tuỳ chọn): văn bản sau chuẩn hoá cấu trúc của pipeline — chỉ dùng để không tính nhãn do
    pipeline sinh là "lặp"; độ phủ nội dung luôn đo trên `tex_content`."""
    docx_path = Path(docx_path)
    if not docx_path.exists():
        return [LayerResult(0, "Gói Word", "FAIL", "Không tìm thấy file .docx.", group="package").to_dict()]
    try:
        root, media = _docx_roots(docx_path)
    except Exception as exc:
        return [LayerResult(0, "Gói Word", "FAIL", f"Không đọc được .docx: {type(exc).__name__}: {exc}",
                            group="package").to_dict()]
    inventory = None
    try:
        from .ooxml_tools import docx_inventory
        inventory = docx_inventory(docx_path)
    except Exception:
        pass
    imgs = _images(root)
    math_images = sum(1 for i in imgs if i["descr"].startswith("l2w:math"))
    bbt_images = sum(1 for i in imgs if i["descr"].startswith("l2w:bbt"))
    results = [
        run_layer1_text(tex_content, root, inventory, normalized_tex=normalized_tex),
        run_layer2_math(tex_content, root, math_images),
        run_layer3_illustrations(tex_content, root, media, manifest),
        run_layer4_tables(tex_content, root, bbt_images),
        run_layer5_visual(visual_info),
    ]
    return [r.to_dict() for r in results]


# ───────────────────────────── kiểm định nhánh PDF ─────────────────────────────
def run_pdf_layers(manifest: dict[str, Any], docx_path: Path, visual_info: dict[str, Any] | None = None) -> list[dict]:
    """Kiểm định đầu ra nhánh PDF theo manifest khối (block_id, method, status)."""
    root, media = _docx_roots(Path(docx_path))
    try:
        from .ooxml_tools import docx_inventory
        inventory = docx_inventory(Path(docx_path))
    except Exception:
        inventory = None
    blocks = manifest.get("blocks", [])
    ref_text = "\n".join(p.get("reference_text", "") for p in manifest.get("pages", []))
    w_text = _w_text(root)
    m_text = "".join(_omath_texts(root))
    ref_words = _words(ref_text)
    doc_words = _words(w_text) + [w for om in _omath_texts(root) for w in _words(om)]
    d = _content_check(ref_words, doc_words, verb_cmds=None, docx_text=w_text, inventory=inventory)
    ref_alnum = Counter(ch for ch in unicodedata.normalize("NFC", ref_text).lower() if ch.isalnum())
    doc_alnum = Counter(ch for ch in unicodedata.normalize("NFC", w_text + m_text).lower() if ch.isalnum())
    char_recall = sum((ref_alnum & doc_alnum).values()) / max(1, sum(ref_alnum.values()))
    invented = sum((doc_alnum - ref_alnum).values())
    d.update({"char_recall": round(char_recall, 4), "chars_not_in_source": invented,
              "note": ("Tham chiếu = chữ hiển thị của chính PDF (pdfplumber), trừ vùng đã giữ bằng ảnh. "
                       "token_recall tính trên TỪ (≥2 chữ cái); char_recall tính mọi chữ/số kể cả trong công thức "
                       "(OMML). Chữ tham chiếu theo thứ tự đường cơ sở nên order_ratio chỉ là tín hiệu.")})
    st1, why = _content_status(d)
    if sum(ref_alnum.values()) == 0:
        # không có chữ hiển thị nào để đối chiếu (trang scan/ảnh giữ nguyên ảnh) → KHÔNG đo được, không phải FAIL/PASS
        st1, why = "NOT_MEASURED", ("PDF không có lớp chữ hiển thị ngoài vùng ảnh để đối chiếu (trang scan/ảnh) — "
                                    "nội dung nằm trong ảnh trang, chưa kiểm được bằng chữ")
    elif char_recall < TEXT_RECALL_FAIL and sum(ref_alnum.values()) >= 20:
        st1, why = "FAIL", f"chỉ giữ {char_recall:.0%} chữ/số"
    elif char_recall < TEXT_RECALL_PASS and st1 == "PASS":
        st1, why = "WARN", f"thiếu {1 - char_recall:.1%} chữ/số (kể cả trong công thức)"
    if invented > max(3, 0.02 * sum(ref_alnum.values())) and st1 == "PASS":
        st1, why = "WARN", f"{invented} chữ/số trong Word không có trong PDF"
    summ = (f"{why}." if st1 == "NOT_MEASURED" else
            f"Giữ {d['token_recall']:.1%} từ và {char_recall:.1%} chữ/số của PDF (ngoài vùng ảnh); {why}.")
    l1 = LayerResult(1, "A. Nội dung (chữ PDF ↔ Word, dấu tiếng Việt, rò rỉ)", st1, summ, d, group="A")
    math_blocks = [b for b in blocks if b.get("type") == "math"]
    math_native = [b for b in math_blocks if b.get("method") == "math_omml"]
    math_img = [b for b in math_blocks if b.get("method") in ("vector_image", "raster_image")]
    math_unres = [b for b in math_blocks if b.get("method") == "unresolved"]
    inline_math = sum(len(b.get("inline_math") or []) for b in blocks if b.get("type") != "math")
    omath = len(_omath_texts(root))
    exp_omml = len(math_native) + inline_math
    imgs = _images(root)
    descrs = [i["descr"] for i in imgs if i["kind"] == "picture"]
    img_found = sum(1 for b in math_img if any(b["block_id"] in x for x in descrs))
    d2 = {"math_blocks": len(math_blocks), "math_omml_blocks": len(math_native), "inline_math_omml": inline_math,
          "expected_omml": exp_omml, "omml_in_docx": omath, "math_as_image": len(math_img),
          "math_images_found": img_found, "unresolved": len(math_unres),
          "metric": "Số OMML phải bằng số khối/đoạn toán đã nhận dạng native; khối toán giữ bằng ảnh phải có ảnh "
                    "mang đúng block_id. Nhận dạng toán từ PDF là heuristic (xem confidence từng khối)."}
    if omath < exp_omml or img_found < len(math_img) or math_unres:
        st2 = "FAIL" if exp_omml and omath < exp_omml * MATH_MATCH_FAIL else "WARN"
    else:
        st2 = "PASS"
    l2 = LayerResult(2, "B. Toán (OMML / ảnh vùng)", st2,
                     f"{omath}/{exp_omml} OMML, {img_found}/{len(math_img)} ảnh công thức (math_as_image)"
                     + (f", {len(math_unres)} unresolved" if math_unres else "") + ".", d2, group="B")
    fig_blocks = [b for b in blocks if b.get("type") in ("figure", "image", "page")]
    found = []
    missing = []
    for b in fig_blocks:
        ok = any(b["block_id"] in x for x in descrs)
        (found if ok else missing).append(b["block_id"])
    d3 = {"figure_blocks": len(fig_blocks), "found_in_docx": len(found), "missing_blocks": missing,
          "logical_images": len(descrs), "media_parts": len(media),
          "by_method": dict(Counter(b.get("method") for b in fig_blocks)),
          "note": "Mỗi khối hình/ảnh trang phải có đúng một ảnh logic mang block_id (SVG+PNG = 1)."}
    st3 = "PASS" if not missing else ("FAIL" if not found and fig_blocks else "WARN")
    l3 = LayerResult(3, "C. Hình (ảnh vùng gốc / DrawingML)", st3,
                     f"{len(found)}/{len(fig_blocks)} khối hình có trong Word" + (f"; thiếu {len(missing)}" if missing else "")
                     + ".", d3, group="C")
    tbl_blocks = [b for b in blocks if b.get("type") == "table"]
    tbl_native = [b for b in tbl_blocks if b.get("method") == "table_native"]
    tbl_img = [b for b in tbl_blocks if b.get("method") in ("vector_image", "raster_image")]
    roles, tinfo = _table_roles(root)
    tbl_img_found = sum(1 for b in tbl_img if any(b["block_id"] in x for x in descrs))
    native_found = 0
    for tbl in root.iter(f"{{{NS['w']}}}tbl"):
        desc = tbl.find(f"{{{NS['w']}}}tblPr/{{{NS['w']}}}tblDescription")
        val = desc.get(f"{{{NS['w']}}}val") if desc is not None else ""
        if any(f"block={b['block_id']}" in (val or "") for b in tbl_native):
            native_found += 1
    empty = [i for i, t in enumerate(tinfo) if t["empty"]]
    d4 = {"table_blocks": len(tbl_blocks), "native": len(tbl_native), "native_found": native_found,
          "as_image": len(tbl_img), "as_image_found": tbl_img_found, "docx_tables_by_role": dict(roles),
          "empty_tables": empty}
    st4 = "PASS" if native_found == len(tbl_native) and tbl_img_found == len(tbl_img) and not empty else "WARN"
    l4 = LayerResult(4, "D. Bảng & bố cục", st4,
                     f"{native_found}/{len(tbl_native)} bảng native, {tbl_img_found}/{len(tbl_img)} bảng giữ ảnh"
                     + (f"; {len(empty)} bảng rỗng" if empty else "") + ".", d4, group="D")
    l5 = run_layer5_visual(visual_info)
    return [r.to_dict() for r in (l1, l2, l3, l4, l5)]


# ───────────────────────────── tổng hợp trạng thái ─────────────────────────────
def summarize_quality(layers: list[dict], *, conversion_completed: bool, package_ok: bool = True,
                      required_layers: tuple[int, ...] = (1, 2, 3, 4, 5), unresolved_blocks: int = 0,
                      blocking_issues: list[str] | None = None, review_issues: list[str] | None = None) -> dict[str, Any]:
    """passed / review_required / failed — SKIP và NOT_MEASURED KHÔNG được coi là đạt."""
    not_performed = [f"L{l['layer']} {l['name']}: {l['summary']}" for l in layers
                     if l["status"] in ("SKIP", "NOT_MEASURED") and l["layer"] in required_layers]
    fails = [f"L{l['layer']}: {l['summary']}" for l in layers if l["status"] == "FAIL"]
    warns = [f"L{l['layer']}: {l['summary']}" for l in layers if l["status"] == "WARN"]
    reasons: list[str] = []
    if not conversion_completed or not package_ok:
        status = "failed"
        reasons.append("không tạo được gói Word hợp lệ" if conversion_completed else "chuyển đổi không hoàn tất")
    elif blocking_issues or fails:
        status = "failed" if blocking_issues else "review_required"
        reasons += list(blocking_issues or []) + fails
        if fails and not blocking_issues:
            reasons.insert(0, "có lớp kiểm định FAIL (thiếu nội dung có biểu diễn thay thế → cần kiểm tra)")
    elif warns or not_performed or unresolved_blocks or review_issues:
        status = "review_required"
        reasons += warns + [f"chưa thực hiện: {x}" for x in not_performed]
        if unresolved_blocks:
            reasons.append(f"{unresolved_blocks} khối unresolved")
        reasons += list(review_issues or [])
    else:
        status = "passed"
    return {"quality_status": status, "conversion_completed": conversion_completed, "reasons": reasons[:50],
            "not_performed": not_performed, "layer_status": {f"L{l['layer']}": l["status"] for l in layers},
            "scope_note": "‘passed’ chỉ có nghĩa đạt các tiêu chí tự động đã nêu cho tệp này, không phải đúng 100%."}


# ───────────────────────────── so sánh hiển thị ─────────────────────────────
def _render_pdf_pages(pdf: Path, out_dir: Path, prefix: str, dpi: int = 100, max_pages: int = 200) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    pages: list[Path] = []
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(pdf))
        try:
            for i in range(min(len(doc), max_pages)):
                img = doc[i].render(scale=dpi / 72.0).to_pil().convert("L")
                p = out_dir / f"{prefix}-{i + 1:04d}.png"
                img.save(p)
                pages.append(p)
        finally:
            doc.close()
        return pages
    except Exception:
        tool = shutil.which("pdftocairo")
        if not tool:
            return []
        subprocess.run([tool, "-png", "-gray", "-r", str(dpi), str(pdf), str(out_dir / prefix)],
                       capture_output=True, timeout=300)
        return sorted(out_dir.glob(prefix + "-*.png"))


def _pdf_text_tokens(pdf: Path, max_pages: int = 300) -> tuple[list[str], list[list[str]]]:
    try:
        import pdfplumber
    except Exception:
        return [], []
    per_page = []
    with pdfplumber.open(str(pdf)) as d:
        for i, page in enumerate(d.pages[:max_pages]):
            per_page.append(_norm_tokens(page.extract_text(x_tolerance=1.0) or ""))
    return [t for p in per_page for t in p], per_page


def _is_word(t: str) -> bool:
    return len(t) >= 2 and sum(ch.isalpha() for ch in t) >= 2 and not any(ch.isdigit() for ch in t)


def _words(text: str) -> list[str]:
    """Từ = dãy chữ cái liên tiếp (≥2), tách ở chữ số/ký hiệu — ổn định giữa PDF và Word/OMML."""
    text = unicodedata.normalize("NFC", text)
    return [t.lower() for t in re.findall(r"[^\W\d_]+", text) if len(t) >= 2]


_OMML_RENDER_CACHE: dict[str, Any] = {}


def renderer_supports_omml(work_dir: Path) -> dict[str, Any]:
    """Thử render một công thức OMML bằng LibreOffice. Bản LibreOffice thiếu module Math (vd. gói
    libreoffice-math chưa cài) sẽ BỎ công thức khi xuất PDF → PDF QA không hiện công thức."""
    if "result" in _OMML_RENDER_CACHE:
        return _OMML_RENDER_CACHE["result"]
    res: dict[str, Any] = {"checked": False}
    try:
        import docx as _docx
        from .omml_math import latex_to_omml_element
        work_dir.mkdir(parents=True, exist_ok=True)
        d = _docx.Document()
        p = d.add_paragraph("Probe ")
        p._p.append(latex_to_omml_element("zq+wk=7"))
        probe = work_dir / "omml_probe.docx"
        d.save(str(probe))
        from .pipeline import render_docx_to_pdf
        rr = render_docx_to_pdf(probe, work_dir / "omml_probe_out", timeout=120)
        if rr.get("ok"):
            import pdfplumber
            with pdfplumber.open(rr["pdf"]) as pd:
                txt = (pd.pages[0].extract_text() or "").replace(" ", "")
            res = {"checked": True, "supported": ("zq" in txt and "wk" in txt), "renderer": rr.get("engine")}
        else:
            res = {"checked": False, "reason": rr.get("error")}
    except Exception as exc:
        res = {"checked": False, "reason": f"{type(exc).__name__}: {exc}"}
    _OMML_RENDER_CACHE["result"] = res
    return res


def omml_words_in_docx(docx_path: Path) -> list[str]:
    root, _m = _docx_roots(Path(docx_path))
    return [w for t in _omath_texts(root) for w in _words(t)]


def _ink_iou(a_path: Path, b_path: Path) -> float | None:
    if not HAS_VISION:
        return None
    a = Image.open(a_path).convert("L")
    b = Image.open(b_path).convert("L").resize(a.size)
    a = a.point(lambda v: 255 if v < 200 else 0).filter(ImageFilter.MaxFilter(5))
    b = b.point(lambda v: 255 if v < 200 else 0).filter(ImageFilter.MaxFilter(5))
    aa = np.asarray(a) > 0
    bb = np.asarray(b) > 0
    union = np.logical_or(aa, bb).sum()
    if union == 0:
        return 1.0
    return float(np.logical_and(aa, bb).sum() / union)


def compare_rendered_pdfs(reference_pdf: Path, word_pdf: Path, work_dir: Path, *, mode: str,
                          exclude_words: Counter | None = None) -> tuple[dict, list]:
    """So PDF render từ Word với PDF tham chiếu. Chỉ số: số trang, độ phủ chữ hiển thị (token, theo thứ tự)
    và IoU mực theo trang (heuristic bố cục, KHÔNG phải độ chính xác nội dung)."""
    build = work_dir / f"visual_{uuid.uuid4().hex[:8]}"   # thư mục mới mỗi lần → không lẫn ảnh cũ
    ref_png = _render_pdf_pages(reference_pdf, build, "ref")
    word_png = _render_pdf_pages(word_pdf, build, "word")
    ref_all, ref_pages = _pdf_text_tokens(reference_pdf)
    word_all, word_pages = _pdf_text_tokens(word_pdf)
    rc, wc = Counter(ref_all), Counter(word_all)
    recall_all = (sum((rc & wc).values()) / max(1, sum(rc.values()))) if ref_all else None
    # Từ (≥2 chữ cái): ít bị ảnh hưởng bởi cách renderer vẽ công thức (LibreOffice thiếu module Math sẽ bỏ OMML)
    ref_words = [w for t in ref_all for w in _words(t)]
    word_words = [w for t in word_all for w in _words(t)]
    excluded = 0
    if exclude_words:
        # từ chỉ nằm trong OMML mà renderer QA không vẽ → không đo được, KHÔNG tính là mất
        budget = Counter(exclude_words)
        kept = []
        for w in ref_words:
            if budget.get(w, 0) > 0 and w not in word_words:
                budget[w] -= 1
                excluded += 1
                continue
            kept.append(w)
        ref_words = kept
    rw, ww = Counter(ref_words), Counter(word_words)
    recall = (sum((rw & ww).values()) / max(1, sum(rw.values()))) if ref_words else recall_all
    order = difflib.SequenceMatcher(None, ref_words, word_words, autojunk=False).ratio() \
        if ref_words and len(ref_words) * len(word_words) < 9_000_000 else None
    missing_words = [t for t, _ in (rw - ww).most_common(30)]
    pages: list[dict] = []
    ious = []
    n = max(len(ref_png), len(word_png))
    for i in range(n):
        rec: dict[str, Any] = {"page": i + 1}
        if i < len(ref_png):
            rec["reference_png"] = str(ref_png[i])
        if i < len(word_png):
            rec["word_png"] = str(word_png[i])
        if i < len(ref_png) and i < len(word_png):
            iou = _ink_iou(ref_png[i], word_png[i])
            rec["ink_iou"] = round(iou, 4) if iou is not None else None
            if iou is not None:
                ious.append(iou)
            if i < len(ref_pages) and i < len(word_pages) and ref_pages[i]:
                c1, c2 = Counter(ref_pages[i]), Counter(word_pages[i])
                rec["page_text_recall"] = round(sum((c1 & c2).values()) / max(1, sum(c1.values())), 4)
        else:
            rec["missing_in"] = "word" if i >= len(word_png) else "reference"
        pages.append(rec)
    info = {
        "enabled": True, "measured": True, "mode": mode,
        "reference_pdf": str(reference_pdf), "word_pdf": str(word_pdf),
        "reference_pages": len(ref_png), "word_pages": len(word_png),
        "page_count_mismatch": len(ref_png) != len(word_png),
        "pages_missing_in_word": [p["page"] for p in pages if p.get("missing_in") == "word"],
        "extra_pages_in_word": [p["page"] for p in pages if p.get("missing_in") == "reference"],
        "text_recall": round(recall, 4) if recall is not None else None,
        "text_recall_all_tokens": round(recall_all, 4) if recall_all is not None else None,
        "text_order_ratio": round(order, 4) if order is not None else None,
        "missing_words_sample": missing_words,
        "words_not_measurable_omml_not_rendered": excluded,
        "ink_iou_mean": round(sum(ious) / len(ious), 4) if ious else None,
        "similarity_percent": round(100 * sum(ious) / len(ious), 2) if ious else None,
        "metric": ("text_recall: TỪ (≥2 chữ cái) của PDF tham chiếu có trong PDF render từ Word (đa tập); "
                   "text_recall_all_tokens: tính cả số/ký hiệu (bị ảnh hưởng nếu renderer không vẽ OMML); "
                   "ink_iou: IoU vùng mực sau giãn 5px ở 100 dpi, theo trang — chỉ là tín hiệu bố cục; "
                   "similarity_percent = 100×ink_iou_mean (giữ khóa cũ, KHÔNG phải độ chính xác)."),
        "render_engine": "LibreOffice (khác Microsoft Word ở ngắt dòng/phông)",
    }
    if mode == "tex":
        info["note"] = "Nguồn TeX được dàn lại trang trong Word (reflow): số trang có thể khác; xét chữ theo thứ tự."
    return info, pages


def generate_visual_comparison(
    workdir: Path,
    tex_file: Path,
    word_pdf_path: Path,
    *,
    reference_pdf: Path | None = None,
    compile_reference: bool = True,
    mode: str = "tex",
    omml_words: list[str] | None = None,
) -> tuple[dict[str, Any], list[dict]]:
    """(API cũ, đã sửa) So PDF xuất từ Word với PDF tham chiếu THẬT.

    * PDF đầu vào: truyền `reference_pdf` = chính PDF gốc.
    * TeX: biên dịch nguồn gốc (shell-escape tắt, timeout) nếu được phép; lỗi → NOT_MEASURED.
    * Không có tham chiếu → không bịa điểm: enabled=False, similarity_percent=None."""
    workdir = Path(workdir)
    info: dict[str, Any] = {"enabled": False, "measured": False, "similarity_percent": None, "mode": mode}
    if not Path(word_pdf_path).exists():
        info["reason"] = "Không có PDF render từ Word (thiếu Microsoft Word/LibreOffice hoặc render lỗi)."
        return info, []
    ref = Path(reference_pdf) if reference_pdf else None
    if ref is None and compile_reference and tex_file is not None and Path(tex_file).exists():
        try:
            from .tex_compile import compile_reference_pdf
            res = compile_reference_pdf(Path(tex_file), workdir / f"ref_build_{uuid.uuid4().hex[:6]}")
            if res.get("ok"):
                ref = Path(res["pdf"])
                info["reference_compile"] = {k: v for k, v in res.items() if k != "log"}
            else:
                info["reference_compile"] = {k: v for k, v in res.items() if k != "log"}
        except Exception as exc:
            info["reference_compile"] = {"ok": False, "error": f"{type(exc).__name__}: {exc}"}
    if ref is None or not ref.exists():
        word_pages = len(_render_pdf_pages(Path(word_pdf_path), workdir / f"visual_{uuid.uuid4().hex[:8]}", "word"))
        info.update({"word_pages": word_pages,
                     "reason": "Không có nguồn so sánh hợp lệ (PDF gốc / PDF biên dịch từ TeX) → NOT_MEASURED. "
                               "Chỉ có ảnh Word thì không đo được độ giống."
                               + (f" Biên dịch: {info.get('reference_compile', {}).get('error', '')}"
                                  if info.get("reference_compile") else "")})
        return info, []
    excl = None
    if omml_words is not None:
        excl = Counter(omml_words)
    cmp_info, pages = compare_rendered_pdfs(ref, Path(word_pdf_path), workdir, mode=mode, exclude_words=excl)
    info.update(cmp_info)
    return info, pages
