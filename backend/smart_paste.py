"""Smart Paste pipeline (`backend/smart_paste.py`) — CHỈ dùng cho nhánh "Dán mã LaTeX".

Luồng:  mã dán → latex_preprocessor (mở rộng macro + chuẩn hoá cấu trúc lạ)
              → core_converter.convert_latex_file_unified (engine gốc, KHÔNG sửa)
              → apply_layout_markers (áp căn lề, màu, viền bảng chuẩn LaTeX, chú thích chân trang,
                cột, khung, tab, ngắt trang, header/footer...) lên .docx
Nếu bất kỳ bước tiền/hậu xử lý nào lỗi → tự động quay về đúng engine gốc (không bao giờ tệ hơn trước).
Nhánh tải file (.pdf/.tex/.zip) không gọi module này nên không bị ảnh hưởng.
"""
from __future__ import annotations

import copy
import inspect
import json
import os
import re
import sys
import traceback
import zipfile
from pathlib import Path
from typing import Any

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_UNDERLINE
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor, Twips
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .latex_preprocessor import (
    IDX_BASE, LIT_REVERSE, MK_DOTFILL, MK_END, MK_FOOTNOTE, MK_HFILL, MK_NUMPAGES, MK_PAGE, MK_PARA,
    MK_RULEFILL, MK_SPAN, MK_TAB, MK_TABLE, PreprocessResult, compute_page_geometry, preprocess_latex,
)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

PPR_SEQ = ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl",
           "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
           "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
           "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
           "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
           "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"]
RPR_SEQ = ["w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike",
           "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish",
           "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight",
           "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
           "w:eastAsianLayout", "w:specVanish", "w:oMath"]
TCPR_SEQ = ["w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd", "w:noWrap",
            "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"]
TBLPR_SEQ = ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
             "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
             "w:tblLayout", "w:tblCellMar", "w:tblLook"]
SECTPR_SEQ = ["w:headerReference", "w:footerReference", "w:footnotePr", "w:endnotePr", "w:type", "w:pgSz",
              "w:pgMar", "w:paperSrc", "w:pgBorders", "w:lnNumType", "w:pgNumType", "w:cols", "w:formProt",
              "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
              "w:printerSettings", "w:sectPrChange"]

HIGHLIGHTS = {"yellow": WD_COLOR_INDEX.YELLOW, "green": WD_COLOR_INDEX.BRIGHT_GREEN, "cyan": WD_COLOR_INDEX.TURQUOISE,
              "pink": WD_COLOR_INDEX.PINK, "red": WD_COLOR_INDEX.RED, "blue": WD_COLOR_INDEX.BLUE,
              "gray": WD_COLOR_INDEX.GRAY_25}
UNDERLINES = {"single": WD_UNDERLINE.SINGLE, "double": WD_UNDERLINE.DOUBLE, "wave": WD_UNDERLINE.WAVY,
              "dash": WD_UNDERLINE.DASH, "dotted": WD_UNDERLINE.DOTTED}
PUA_RE = re.compile("[-]")


def _insert_ordered(parent, child, seq: list[str]) -> None:
    tag = child.tag
    names = [qn(t) for t in seq]
    try:
        idx = names.index(tag)
    except ValueError:
        parent.append(child)
        return
    for existing in parent.findall(tag):
        parent.remove(existing)
    successors = set(names[idx + 1 :])
    for pos, el in enumerate(list(parent)):
        if el.tag in successors:
            parent.insert(pos, child)
            return
    parent.append(child)


def _get_or_add(parent, tag: str, seq: list[str]):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        _insert_ordered(parent, el, seq)
    return el


# =====================================================================================
# RUN FORMATTING
# =====================================================================================
def _merge_specs(stack: list[dict[str, Any]]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for sp in stack:
        out.update(sp)
    return out


def _apply_run_spec(r_el, spec: dict[str, Any]) -> None:
    if not spec:
        return
    run = Run(r_el, None)
    f = run.font
    if "b" in spec:
        f.bold = bool(spec["b"])
    if "i" in spec:
        f.italic = bool(spec["i"])
    if spec.get("c"):
        try:
            f.color.rgb = RGBColor.from_string(spec["c"])
        except ValueError:
            pass
    if spec.get("sz"):
        f.size = Pt(float(spec["sz"]))
    if spec.get("f"):
        f.name = spec["f"]
        rPr = r_el.get_or_add_rPr()
        rf = rPr.find(qn("w:rFonts"))
        if rf is not None:
            rf.set(qn("w:cs"), spec["f"])
            rf.set(qn("w:eastAsia"), spec["f"])
    if spec.get("sc"):
        f.small_caps = True
    if spec.get("caps"):
        f.all_caps = True
    if spec.get("s"):
        f.strike = True
    if spec.get("u"):
        f.underline = UNDERLINES.get(spec["u"], WD_UNDERLINE.SINGLE)
    if spec.get("va") == "sup":
        f.superscript = True
    elif spec.get("va") == "sub":
        f.subscript = True
    if spec.get("hl"):
        f.highlight_color = HIGHLIGHTS.get(spec["hl"], WD_COLOR_INDEX.YELLOW)
    rPr = r_el.get_or_add_rPr()
    if spec.get("shd"):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), spec["shd"])
        _insert_ordered(rPr, shd, RPR_SEQ)
    if spec.get("bdr"):
        bdr = OxmlElement("w:bdr")
        bdr.set(qn("w:val"), "single")
        bdr.set(qn("w:sz"), "4")
        bdr.set(qn("w:space"), "1")
        bdr.set(qn("w:color"), spec["bdr"] if isinstance(spec["bdr"], str) else "000000")
        _insert_ordered(rPr, bdr, RPR_SEQ)


def _apply_math_spec(math_el, spec: dict[str, Any]) -> None:
    if not spec:
        return
    for mr in math_el.iter(qn("m:r")):
        rPr = mr.find(qn("w:rPr"))
        if rPr is None:
            rPr = OxmlElement("w:rPr")
            t = mr.find(qn("m:t"))
            if t is not None:
                t.addprevious(rPr)
            else:
                mr.append(rPr)
        if spec.get("c"):
            c = OxmlElement("w:color")
            c.set(qn("w:val"), spec["c"])
            _insert_ordered(rPr, c, RPR_SEQ)
        if spec.get("sz"):
            v = str(int(round(float(spec["sz"]) * 2)))
            for tag in ("w:sz", "w:szCs"):
                e = OxmlElement(tag)
                e.set(qn("w:val"), v)
                _insert_ordered(rPr, e, RPR_SEQ)
        if spec.get("b"):
            _insert_ordered(rPr, OxmlElement("w:b"), RPR_SEQ)


# =====================================================================================
# PARAGRAPH PROCESSING
# =====================================================================================
class _Ctx:
    def __init__(self, pre: PreprocessResult, doc, textwidth: int) -> None:
        self.pre = pre
        self.doc = doc
        self.textwidth = textwidth
        self.footnote_refs: list[tuple[Any, int]] = []
        self.struct: list[tuple[Any, dict[str, Any]]] = []
        self.stats: dict[str, int] = {}

    def stat(self, k: str, n: int = 1) -> None:
        self.stats[k] = self.stats.get(k, 0) + n


def _new_run_like(src_rPr):
    r = OxmlElement("w:r")
    if src_rPr is not None:
        r.append(copy.deepcopy(src_rPr))
    return r


def _make_t(text: str):
    t = OxmlElement("w:t")
    t.text = text
    if text != text.strip() or "  " in text:
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return t


def _field_run(instr: str, placeholder: str, rPr=None):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = _new_run_like(rPr)
    r.append(_make_t(placeholder))
    fld.append(r)
    return fld


def _process_paragraph(p, ctx: _Ctx) -> None:
    has_marker = False
    for t in p.iter(qn("w:t"), qn("m:t")):
        if t.text and PUA_RE.search(t.text):
            has_marker = True
            break
    if not has_marker:
        return
    pre = ctx.pre
    stack: list[dict[str, Any]] = []
    para_specs: list[dict[str, Any]] = []
    fills: list[str] = []
    new_children: list[Any] = []

    for child in list(p):
        tag = child.tag
        if tag == qn("w:r"):
            texts = [t for t in child.findall(qn("w:t"))]
            needs_split = any(t.text and PUA_RE.search(t.text) for t in texts)
            if not needs_split:
                if stack and child.find(qn("w:drawing")) is None:
                    _apply_run_spec(child, _merge_specs(stack))
                new_children.append(child)
                continue
            rPr = child.find(qn("w:rPr"))
            cur = _new_run_like(rPr)
            cur_has = False
            out_runs: list[Any] = []

            def finalize() -> None:
                nonlocal cur, cur_has
                if cur_has:
                    _apply_run_spec(cur, _merge_specs(stack))
                    out_runs.append(cur)
                cur = _new_run_like(rPr)
                cur_has = False

            for sub in list(child):
                if sub.tag == qn("w:rPr"):
                    continue
                if sub.tag != qn("w:t"):
                    cur.append(copy.deepcopy(sub))
                    cur_has = True
                    continue
                txt = sub.text or ""
                buf: list[str] = []
                i = 0
                n = len(txt)

                def flush_buf() -> None:
                    nonlocal cur_has
                    if buf:
                        cur.append(_make_t("".join(buf)))
                        cur_has = True
                        buf.clear()

                while i < n:
                    ch = txt[i]
                    if ch == MK_SPAN and i + 1 < n:
                        flush_buf()
                        finalize()
                        idx = ord(txt[i + 1]) - IDX_BASE
                        stack.append(pre.span_specs[idx] if 0 <= idx < len(pre.span_specs) else {})
                        i += 2
                        continue
                    if ch == MK_END:
                        flush_buf()
                        finalize()
                        if stack:
                            stack.pop()
                        i += 1
                        continue
                    if ch == MK_PARA and i + 1 < n:
                        idx = ord(txt[i + 1]) - IDX_BASE
                        if 0 <= idx < len(pre.para_specs):
                            para_specs.append(pre.para_specs[idx])
                        i += 2
                        continue
                    if ch in (MK_TAB, MK_HFILL, MK_DOTFILL, MK_RULEFILL):
                        flush_buf()
                        cur.append(OxmlElement("w:tab"))
                        cur_has = True
                        if ch != MK_TAB:
                            fills.append(ch)
                        i += 1
                        continue
                    if ch == MK_FOOTNOTE and i + 1 < n:
                        flush_buf()
                        finalize()
                        idx = ord(txt[i + 1]) - IDX_BASE
                        fr = OxmlElement("w:r")
                        frPr = OxmlElement("w:rPr")
                        va = OxmlElement("w:vertAlign")
                        va.set(qn("w:val"), "superscript")
                        frPr.append(va)
                        fr.append(frPr)
                        ref = OxmlElement("w:footnoteReference")
                        ref.set(qn("w:id"), str(idx + 1))
                        fr.append(ref)
                        out_runs.append(fr)
                        ctx.footnote_refs.append((fr, idx))
                        i += 2
                        continue
                    if ch == MK_TABLE and i + 1 < n:
                        i += 2
                        continue
                    if ch in (MK_PAGE, MK_NUMPAGES):
                        flush_buf()
                        finalize()
                        out_runs.append(_field_run("PAGE" if ch == MK_PAGE else "NUMPAGES", "1", rPr))
                        i += 1
                        continue
                    if ch in LIT_REVERSE:
                        buf.append(LIT_REVERSE[ch])
                        i += 1
                        continue
                    if "" <= ch <= "":
                        i += 1
                        continue
                    buf.append(ch)
                    i += 1
                flush_buf()
            finalize()
            new_children.extend(out_runs)
            continue
        if tag in (qn("m:oMath"), qn("m:oMathPara")):
            for mt in child.iter(qn("m:t")):
                if mt.text and PUA_RE.search(mt.text):
                    mt.text = "".join(LIT_REVERSE.get(c, "" if "" <= c <= "" else c) for c in mt.text)
            if stack:
                _apply_math_spec(child, _merge_specs(stack))
            new_children.append(child)
            continue
        new_children.append(child)

    for c in list(p):
        p.remove(c)
    for c in new_children:
        p.append(c)
    # drop empty runs
    for r in p.findall(qn("w:r")):
        if len([c for c in r if c.tag != qn("w:rPr")]) == 0:
            p.remove(r)
        else:
            for t in r.findall(qn("w:t")):
                if t.text == "":
                    r.remove(t)
    spec: dict[str, Any] = {}
    for sp in para_specs:
        spec.update(sp)
    if spec or fills:
        _apply_para_spec(p, spec, fills, ctx)


def _apply_para_spec(p, spec: dict[str, Any], fills: list[str], ctx: _Ctx) -> None:
    para = Paragraph(p, None)
    pf = para.paragraph_format
    al = spec.get("align")
    if al:
        pf.alignment = {"c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT, "l": WD_ALIGN_PARAGRAPH.LEFT,
                        "j": WD_ALIGN_PARAGRAPH.JUSTIFY}.get(al, WD_ALIGN_PARAGRAPH.LEFT)
    left = int(spec.get("left") or 0)
    if left:
        pf.left_indent = Twips(left)
    if spec.get("hang"):
        pf.first_line_indent = Twips(-int(spec["hang"]))
    if spec.get("right"):
        pf.right_indent = Twips(int(spec["right"]))
    if "sb" in spec:
        pf.space_before = Pt(float(spec["sb"]))
    if "sa" in spec:
        pf.space_after = Pt(float(spec["sa"]))
    if spec.get("keep"):
        pf.keep_with_next = True
    pPr = p.get_or_add_pPr()
    if spec.get("heading"):
        lvl = OxmlElement("w:outlineLvl")
        lvl.set(qn("w:val"), str(max(0, min(8, int(spec["heading"]) - 1))))
        _insert_ordered(pPr, lvl, PPR_SEQ)
        ctx.stat("headings")
    box = spec.get("box")
    if box:
        pBdr = OxmlElement("w:pBdr")
        sides = box.get("sides", "tlbr")
        for side, key in (("top", "t"), ("left", "l"), ("bottom", "b"), ("right", "r")):
            if key not in sides:
                continue
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), box.get("style", "single"))
            e.set(qn("w:sz"), "6")
            e.set(qn("w:space"), "4")
            e.set(qn("w:color"), box.get("color", "000000"))
            pBdr.append(e)
        _insert_ordered(pPr, pBdr, PPR_SEQ)
        if not left:
            pf.left_indent = Twips(100)
            pf.right_indent = Twips(100)
    if spec.get("hr"):
        pBdr = OxmlElement("w:pBdr")
        e = OxmlElement("w:bottom")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "1")
        e.set(qn("w:color"), "000000")
        pBdr.append(e)
        _insert_ordered(pPr, pBdr, PPR_SEQ)
    if spec.get("shd"):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), spec["shd"])
        _insert_ordered(pPr, shd, PPR_SEQ)
    tabs = list(spec.get("tabs") or [])
    width = ctx.textwidth
    stops: list[tuple[int, str, str]] = [(int(t), "left", "none") for t in tabs]
    if fills:
        if spec.get("align") in ("c", "r", None):
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT  # a line with \\hfill spans the full text width
        nf = len(fills)
        avail = width
        for k, ch in enumerate(fills):
            leader = {"": "dot", "": "underscore"}.get(ch, "none")
            if nf == 1:
                stops.append((avail, "right", leader))
            elif nf == 2:
                stops.append(((avail // 2) if k == 0 else avail, "center" if k == 0 else "right", leader))
            else:
                pos = int(avail * (k + 1) / nf)
                stops.append((pos, "right" if k == nf - 1 else "center", leader))
    if stops:
        tabs_el = OxmlElement("w:tabs")
        for pos, kind, leader in sorted(stops):
            t = OxmlElement("w:tab")
            t.set(qn("w:val"), kind)
            if leader != "none":
                t.set(qn("w:leader"), leader)
            t.set(qn("w:pos"), str(max(0, pos)))
            tabs_el.append(t)
        _insert_ordered(pPr, tabs_el, PPR_SEQ)
    for key in ("pb", "vspace_only", "cols", "cols_end", "colbreak", "toc"):
        if key in spec:
            ctx.struct.append((p, spec))
            break


# =====================================================================================
# TABLES
# =====================================================================================
def _cell_text_len(tc) -> int:
    best = 0
    for p in tc.iter(qn("w:p")):
        line = 0
        for el in p.iter():
            if el.tag in (qn("w:t"), qn("m:t")) and el.text:
                line += len(el.text)
            elif el.tag == qn("w:br"):
                best = max(best, line)
                line = 0
        best = max(best, line)
    return best


def _set_tc_borders(tc, borders: dict[str, tuple[str, int] | None]) -> None:
    tcPr = tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:tcBorders"))
    diag = None
    if old is not None:
        diag = old.find(qn("w:tl2br"))
        tcPr.remove(old)
    tb = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        val = borders.get(side)
        if val:
            e.set(qn("w:val"), val[0])
            e.set(qn("w:sz"), str(val[1]))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), "000000")
        else:
            e.set(qn("w:val"), "nil")
        tb.append(e)
    if diag is not None:
        diag.set(qn("w:color"), "000000")
        tb.append(diag)
    _insert_ordered(tcPr, tb, TCPR_SEQ)


def _apply_table_spec(tbl, spec: dict[str, Any], ctx: _Ctx) -> None:
    ncols = int(spec.get("ncols") or 1)
    vlines = list(spec.get("vlines") or [0] * (ncols + 1))
    while len(vlines) < ncols + 1:
        vlines.append(0)
    rules = spec.get("rules") or []
    bottom = spec.get("bottom") or {}
    mc_bars = spec.get("mc_bars") or {}
    rows = tbl.findall(qn("w:tr"))
    nrows = len(rows)

    def line_for(cov: dict[str, Any], c0: int, c1: int) -> tuple[str, int] | None:
        for c in range(c0, c1 + 1):
            v = cov.get(str(c))
            if v:
                return (v[0], int(v[1]))
        return None

    # table-level properties
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tb = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "nil")
        tb.append(e)
    _insert_ordered(tblPr, tb, TBLPR_SEQ)

    for r_idx, tr in enumerate(rows):
        cov_top = rules[r_idx] if r_idx < len(rules) else {}
        cov_bot = (rules[r_idx + 1] if r_idx + 1 < len(rules) else {}) if r_idx + 1 < nrows else bottom
        bounds = [vlines[k] if k < len(vlines) else 0 for k in range(ncols + 1)]
        c = 0
        tcs = tr.findall(qn("w:tc"))
        spans: list[tuple[Any, int, int]] = []
        for tc in tcs:
            tcPr = tc.get_or_add_tcPr()
            gs = tcPr.find(qn("w:gridSpan"))
            span = int(gs.get(qn("w:val"))) if gs is not None else 1
            spans.append((tc, c, c + span - 1))
            key = f"{r_idx},{c}"
            if key in mc_bars:
                lb, rb = mc_bars[key]
                if c == 0:
                    bounds[0] = lb
                if c + span <= ncols:
                    bounds[min(ncols, c + span)] = rb
            c += span
        for tc, c0, c1 in spans:
            def vb(k: int) -> tuple[str, int] | None:
                k = max(0, min(ncols, k))
                cnt = bounds[k]
                if cnt <= 0:
                    return None
                return ("double", 4) if cnt >= 2 else ("single", 4)
            _set_tc_borders(tc, {
                "top": line_for(cov_top, c0, c1),
                "bottom": line_for(cov_bot, c0, c1),
                "left": vb(c0),
                "right": vb(c1 + 1),
            })
            tcPr = tc.get_or_add_tcPr()
            for shd in tcPr.findall(qn("w:shd")):
                tcPr.remove(shd)
            fill = (spec.get("cell_shd") or {}).get(f"{r_idx},{c0}") or (spec.get("row_shd") or {}).get(str(r_idx))
            if fill:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), fill)
                _insert_ordered(tcPr, shd, TCPR_SEQ)
            # header row: remove engine's automatic bold (explicit bold comes back via run spans)
            if spec.get("header_plain") and r_idx == 0:
                for r in tc.iter(qn("w:r")):
                    rPr = r.find(qn("w:rPr"))
                    if rPr is not None:
                        for b in rPr.findall(qn("w:b")) + rPr.findall(qn("w:bCs")):
                            rPr.remove(b)
            fs = spec.get("font_size")
            if fs:
                v = str(int(round(float(fs) * 2)))
                for r in tc.iter(qn("w:r")):
                    if r.find(qn("w:t")) is None:
                        continue
                    rPr = r.find(qn("w:rPr"))
                    if rPr is None:
                        rPr = OxmlElement("w:rPr")
                        r.insert(0, rPr)
                    for tag in ("w:sz", "w:szCs"):
                        e = OxmlElement(tag)
                        e.set(qn("w:val"), v)
                        _insert_ordered(rPr, e, RPR_SEQ)
    # ----- widths -----
    grid = tbl.find(qn("w:tblGrid"))
    gcols = grid.findall(qn("w:gridCol")) if grid is not None else []
    n = len(gcols) or ncols
    fixed = list(spec.get("widths") or [])
    xflags = list(spec.get("x") or [])
    while len(fixed) < n:
        fixed.append(None)
    while len(xflags) < n:
        xflags.append(False)
    content = [0] * n
    for tr in rows:
        c = 0
        for tc in tr.findall(qn("w:tc")):
            gs = tc.tcPr.find(qn("w:gridSpan")) if tc.tcPr is not None else None
            span = int(gs.get(qn("w:val"))) if gs is not None else 1
            if span == 1 and c < n:
                content[c] = max(content[c], _cell_text_len(tc))
            c += span
    max_w = int(spec.get("max_width") or ctx.textwidth)
    widths: list[int] = []
    for k in range(n):
        if fixed[k]:
            widths.append(int(fixed[k]) + 240)
        elif xflags[k]:
            widths.append(0)
        else:
            widths.append(max(560, min(max_w, content[k] * 140 + 420)))
    target = spec.get("target_width")
    if any(xflags[:n]):
        tw = int(target or max_w)
        rest = tw - sum(w for w in widths)
        nx = sum(1 for x in xflags[:n] if x)
        each = max(700, rest // max(1, nx))
        widths = [each if xflags[k] else widths[k] for k in range(n)]
    total = sum(widths)
    if total > max_w:
        factor = max_w / total
        widths = [max(300, int(w * factor)) for w in widths]
    elif target and total < target:
        extra = (int(target) - total) // max(1, n)
        widths = [w + extra for w in widths]
    for k, gc in enumerate(gcols):
        gc.set(qn("w:w"), str(widths[k]))
    for tr in rows:
        c = 0
        for tc in tr.findall(qn("w:tc")):
            tcPr = tc.get_or_add_tcPr()
            gs = tcPr.find(qn("w:gridSpan"))
            span = int(gs.get(qn("w:val"))) if gs is not None else 1
            w = sum(widths[c : c + span])
            tcW = OxmlElement("w:tcW")
            tcW.set(qn("w:w"), str(w))
            tcW.set(qn("w:type"), "dxa")
            _insert_ordered(tcPr, tcW, TCPR_SEQ)
            c += span
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    _insert_ordered(tblPr, tblW, TBLPR_SEQ)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    _insert_ordered(tblPr, lay, TBLPR_SEQ)
    ctx.stat("tables_latex_borders")


# =====================================================================================
# STRUCTURAL PARAGRAPHS
# =====================================================================================
def _next_block_sibling(el):
    nxt = el.getnext()
    while nxt is not None and nxt.tag not in (qn("w:p"), qn("w:tbl")):
        nxt = nxt.getnext()
    return nxt


def _clear_runs(p) -> None:
    for c in list(p):
        if c.tag != qn("w:pPr"):
            p.remove(c)


def _tiny_paragraph(p) -> None:
    pf = Paragraph(p, None).paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(1)
    pPr = p.get_or_add_pPr()
    rPr = pPr.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        _insert_ordered(pPr, rPr, PPR_SEQ)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    _insert_ordered(rPr, sz, RPR_SEQ)


def _make_sectpr(template, cols: int, continuous: bool):
    sp = copy.deepcopy(template)
    for t in sp.findall(qn("w:type")):
        sp.remove(t)
    if continuous:
        t = OxmlElement("w:type")
        t.set(qn("w:val"), "continuous")
        _insert_ordered(sp, t, SECTPR_SEQ)
    for c in sp.findall(qn("w:cols")):
        sp.remove(c)
    ce = OxmlElement("w:cols")
    ce.set(qn("w:space"), "425")
    if cols > 1:
        ce.set(qn("w:num"), str(cols))
    _insert_ordered(sp, ce, SECTPR_SEQ)
    return sp


def _apply_structure(doc, ctx: _Ctx) -> None:
    body = doc.element.body
    body_sectPr = body.find(qn("w:sectPr"))
    col_state = 2 if ctx.pre.page.get("twocolumn") else 1
    first_section = True
    sections_made = 0
    for p, spec in ctx.struct:
        if p.getparent() is None:
            continue
        if spec.get("toc"):
            _insert_toc(p, ctx)
            continue
        if spec.get("colbreak"):
            _clear_runs(p)
            r = OxmlElement("w:r")
            br = OxmlElement("w:br")
            br.set(qn("w:type"), "column")
            r.append(br)
            p.append(r)
            _tiny_paragraph(p)
            continue
        if spec.get("pb"):
            nxt = _next_block_sibling(p)
            if nxt is not None and nxt.tag == qn("w:p") and p.getparent() is body:
                Paragraph(nxt, None).paragraph_format.page_break_before = True
                p.getparent().remove(p)
            else:
                _clear_runs(p)
                r = OxmlElement("w:r")
                br = OxmlElement("w:br")
                br.set(qn("w:type"), "page")
                r.append(br)
                p.append(r)
            ctx.stat("page_breaks")
            continue
        if spec.get("vspace_only"):
            nxt = _next_block_sibling(p)
            if nxt is not None and nxt.tag == qn("w:p"):
                pf = Paragraph(nxt, None).paragraph_format
                cur = pf.space_before.pt if pf.space_before is not None else 0
                pf.space_before = Pt(cur + float(spec.get("sb", 0)))
                p.getparent().remove(p)
            else:
                _clear_runs(p)
            continue
        if ("cols" in spec or "cols_end" in spec) and body_sectPr is not None and p.getparent() is body:
            new_cols = int(spec["cols"]) if "cols" in spec else 1
            _clear_runs(p)
            _tiny_paragraph(p)
            pPr = p.get_or_add_pPr()
            sp = _make_sectpr(body_sectPr, col_state, continuous=not first_section)
            _insert_ordered(pPr, sp, PPR_SEQ)
            first_section = False
            col_state = new_cols
            sections_made += 1
            ctx.stat("column_sections")
    if sections_made and body_sectPr is not None:
        new_last = _make_sectpr(body_sectPr, col_state, continuous=True)
        body.replace(body_sectPr, new_last)
    elif col_state > 1 and body_sectPr is not None:
        body.replace(body_sectPr, _make_sectpr(body_sectPr, col_state, continuous=False))


def _insert_toc(p, ctx: _Ctx) -> None:
    _clear_runs(p)

    def fld(kind: str):
        r = OxmlElement("w:r")
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), kind)
        if kind == "begin":
            f.set(qn("w:dirty"), "true")
        r.append(f)
        return r

    p.append(fld("begin"))
    r = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    it.text = ' TOC \\o "1-3" \\h \\z \\u '
    r.append(it)
    p.append(r)
    p.append(fld("separate"))
    headings = ctx.pre.page.get("headings") or []
    last = p
    if not headings:
        rr = OxmlElement("w:r")
        rr.append(_make_t("Nhấn F9 (hoặc chuột phải → Update Field) để cập nhật mục lục."))
        p.append(rr)
        p.append(fld("end"))
        return
    first = True
    for lvl, text in headings:
        if first:
            target = p
            first = False
        else:
            target = OxmlElement("w:p")
            last.addnext(target)
        pf = Paragraph(target, None).paragraph_format
        pf.left_indent = Twips(max(0, (int(lvl) - 1)) * 340)
        pf.space_after = Pt(2)
        tabs = OxmlElement("w:tabs")
        t = OxmlElement("w:tab")
        t.set(qn("w:val"), "right")
        t.set(qn("w:leader"), "dot")
        t.set(qn("w:pos"), str(ctx.textwidth))
        tabs.append(t)
        _insert_ordered(target.get_or_add_pPr(), tabs, PPR_SEQ)
        rr = OxmlElement("w:r")
        if int(lvl) <= 1:
            rPr = OxmlElement("w:rPr")
            rPr.append(OxmlElement("w:b"))
            rr.append(rPr)
        rr.append(_make_t(text))
        target.append(rr)
        last = target
    last.append(fld("end"))
    ctx.stat("toc")


# =====================================================================================
# FOOTNOTES, HEADER/FOOTER, PAGE
# =====================================================================================
def _add_footnotes(doc, ctx: _Ctx) -> None:
    fns = ctx.pre.footnotes
    if not fns or not ctx.footnote_refs:
        return
    from .omml_math import append_rich_latex_text_to_paragraph

    xml = (
        f'<w:footnotes {nsdecls("w", "m", "r")}>'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
        '<w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
        '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        '</w:footnotes>'
    )
    root = parse_xml(xml)
    for idx, body_tex in enumerate(fns):
        fn = OxmlElement("w:footnote")
        fn.set(qn("w:id"), str(idx + 1))
        p = OxmlElement("w:p")
        fn.append(p)
        pf = Paragraph(p, None).paragraph_format
        pf.space_after = Pt(0)
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        va = OxmlElement("w:vertAlign")
        va.set(qn("w:val"), "superscript")
        rPr.append(va)
        r.append(rPr)
        r.append(OxmlElement("w:footnoteRef"))
        p.append(r)
        sp = OxmlElement("w:r")
        sp.append(_make_t(" "))
        p.append(sp)
        para = Paragraph(p, None)
        try:
            append_rich_latex_text_to_paragraph(para, body_tex, font_size_pt=10.0)
        except Exception:
            para.add_run(PUA_RE.sub("", body_tex))
        _process_paragraph(p, ctx)
        root.append(fn)
    part = XmlPart(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, root, doc.part.package)
    doc.part.relate_to(part, RT.FOOTNOTES)
    settings = doc.settings.element
    fpr = settings.find(qn("w:footnotePr"))
    if fpr is None:
        fpr = parse_xml(f'<w:footnotePr {nsdecls("w")}><w:footnote w:id="-1"/><w:footnote w:id="0"/></w:footnotePr>')
        successors = {qn(t) for t in ("w:endnotePr", "w:compat", "w:docVars", "w:rsids", "w:attachedSchema",
                                       "w:themeFontLang", "w:clrSchemeMapping", "w:doNotIncludeSubdocsInStats",
                                       "w:doNotAutoCompressPictures", "w:forceUpgrade", "w:captions",
                                       "w:readModeInkLockDown", "w:smartTagType", "w:schemaLibrary",
                                       "w:shapeDefaults", "w:doNotEmbedSmartTags", "w:decimalSymbol",
                                       "w:listSeparator")} | {qn("m:mathPr")}
        for pos, el in enumerate(list(settings)):
            if el.tag in successors:
                settings.insert(pos, fpr)
                break
        else:
            settings.append(fpr)
    ctx.stat("footnotes", len(fns))


def _style_tab_positions(p_obj) -> list[int]:
    out: list[int] = []
    try:
        st = p_obj.style
        while st is not None:
            ppr = st.element.pPr
            if ppr is not None:
                tabs = ppr.find(qn("w:tabs"))
                if tabs is not None:
                    for t in tabs.findall(qn("w:tab")):
                        if t.get(qn("w:val")) != "clear":
                            out.append(int(t.get(qn("w:pos"))))
            st = st.base_style
    except Exception:
        pass
    return out


def _fill_hf_paragraph(p_obj, parts: dict[str, str], width: int, rule: bool) -> None:
    p = p_obj._p
    _clear_runs(p)
    pf = p_obj.paragraph_format
    pPr = p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    for pos_clear in _style_tab_positions(p_obj):
        t = OxmlElement("w:tab")
        t.set(qn("w:val"), "clear")
        t.set(qn("w:pos"), str(pos_clear))
        tabs.append(t)
    for kind, pos in (("center", width // 2), ("right", width)):
        t = OxmlElement("w:tab")
        t.set(qn("w:val"), kind)
        t.set(qn("w:pos"), str(pos))
        tabs.append(t)
    _insert_ordered(pPr, tabs, PPR_SEQ)
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)

    def add_text(txt: str) -> None:
        segs = re.split("([])", txt)
        for seg in segs:
            if not seg:
                continue
            if seg in (MK_PAGE, MK_NUMPAGES):
                p.append(_field_run("PAGE" if seg == MK_PAGE else "NUMPAGES", "1"))
            else:
                r = OxmlElement("w:r")
                r.append(_make_t(seg))
                p.append(r)

    first = True
    for pos in ("L", "C", "R"):
        if not first:
            r = OxmlElement("w:r")
            r.append(OxmlElement("w:tab"))
            p.append(r)
        first = False
        if parts.get(pos):
            add_text(parts[pos])
    for r in p.findall(qn("w:r")):
        rPr = r.find(qn("w:rPr"))
        if rPr is None:
            rPr = OxmlElement("w:rPr")
            r.insert(0, rPr)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "22")
        _insert_ordered(rPr, sz, RPR_SEQ)
    if rule:
        pBdr = OxmlElement("w:pBdr")
        e = OxmlElement("w:bottom")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "1")
        e.set(qn("w:color"), "000000")
        pBdr.append(e)
        _insert_ordered(pPr, pBdr, PPR_SEQ)


def _apply_page(doc, ctx: _Ctx) -> None:
    page = ctx.pre.page
    geo = compute_page_geometry(page)
    for section in doc.sections:
        section.page_width = Twips(geo["width"])
        section.page_height = Twips(geo["height"])
        section.left_margin = Twips(geo["left"])
        section.right_margin = Twips(geo["right"])
        section.top_margin = Twips(geo["top"])
        section.bottom_margin = Twips(geo["bottom"])
    width = geo["width"] - geo["left"] - geo["right"]
    style = page.get("pagestyle", "plain")
    section = doc.sections[0]
    if style in ("empty",):
        return
    header = dict(page.get("header") or {})
    footer = dict(page.get("footer") or {})
    if style in ("plain", "headings", "myheadings") and not footer:
        footer = {"C": MK_PAGE}
    if style == "fancy" and not footer and not page.get("footer_cleared"):
        footer = {"C": MK_PAGE}
    if header:
        _fill_hf_paragraph(section.header.paragraphs[0], header, width,
                           rule=style == "fancy" and str(page.get("headrule", "0.4pt")).strip() not in ("0pt", "0", "0mm"))
        ctx.stat("header")
    if footer:
        _fill_hf_paragraph(section.footer.paragraphs[0], footer, width, rule=False)
        ctx.stat("footer")
    if page.get("first_empty"):
        section.different_first_page_header_footer = True


# =====================================================================================
# OOXML ORDER REPAIR
# =====================================================================================
_ORDER_SEQS = {
    "w:pPr": PPR_SEQ, "w:rPr": RPR_SEQ, "w:tcPr": TCPR_SEQ, "w:tblPr": TBLPR_SEQ,
    "w:tcBorders": ["w:top", "w:start", "w:left", "w:bottom", "w:end", "w:right", "w:insideH", "w:insideV",
                    "w:tl2br", "w:tr2bl"],
    "w:tblBorders": ["w:top", "w:start", "w:left", "w:bottom", "w:end", "w:right", "w:insideH", "w:insideV"],
    "w:pBdr": ["w:top", "w:left", "w:bottom", "w:right", "w:between", "w:bar"],
    "w:tcMar": ["w:top", "w:start", "w:left", "w:bottom", "w:end", "w:right"],
}


def _normalize_xml_order(root) -> int:
    fixed = 0
    seqs = {qn(k): [qn(t) for t in v] for k, v in _ORDER_SEQS.items()}
    for el in list(root.iter(*seqs.keys())):
        seq = seqs[el.tag]
        kids = [c for c in el if isinstance(c.tag, str)]
        if not kids or any(c.tag not in seq for c in kids):
            continue
        # keep the LAST occurrence of duplicated tags (later writers intend to override)
        last: dict[str, Any] = {}
        for c in kids:
            last[c.tag] = c
        wanted = sorted(last.values(), key=lambda c: seq.index(c.tag))
        if wanted != kids:
            for c in kids:
                el.remove(c)
            for c in wanted:
                el.append(c)
            fixed += 1
    return fixed


# =====================================================================================
# MAIN POST-PASS
# =====================================================================================
def apply_layout_markers(docx_path: Path, pre: PreprocessResult) -> dict[str, int]:
    doc = docx.Document(str(docx_path))
    geo = compute_page_geometry(pre.page)
    ctx = _Ctx(pre, doc, geo["width"] - geo["left"] - geo["right"])
    body = doc.element.body
    # 1) tables first (header-bold reset must happen before run spans re-apply explicit bold)
    for t in list(body.iter(qn("w:t"))):
        txt = t.text or ""
        pos = txt.find(MK_TABLE)
        if pos == -1 or pos + 1 >= len(txt):
            continue
        idx = ord(txt[pos + 1]) - IDX_BASE
        t.text = txt[:pos] + txt[pos + 2 :]
        tbl = t
        while tbl is not None and tbl.tag != qn("w:tbl"):
            tbl = tbl.getparent()
        if tbl is not None and 0 <= idx < len(pre.table_specs):
            try:
                _apply_table_spec(tbl, pre.table_specs[idx], ctx)
            except Exception as exc:  # never break the document
                ctx.stat("table_spec_errors")
                pre.warnings.append(f"Không áp được viền bảng #{idx}: {exc}")
    # 2) every paragraph (body, tables, text boxes)
    for p in list(body.iter(qn("w:p"))):
        _process_paragraph(p, ctx)
    # 3) page geometry + header/footer (before cloning sectPr for columns)
    _apply_page(doc, ctx)
    # 4) structural paragraphs (page breaks, spacing, columns, TOC)
    _apply_structure(doc, ctx)
    # 5) footnotes
    _add_footnotes(doc, ctx)
    for sec in doc.sections:
        for hf in (sec.header, sec.footer):
            try:
                if not hf.is_linked_to_previous:
                    _normalize_xml_order(hf._element)
            except Exception:
                pass
    # 6) repair OOXML child order / duplicates (Word is strict about CT_PPr, CT_TcPr ... sequences)
    _normalize_xml_order(body)
    # 7) final scrub of any leftover private markers
    for t in body.iter(qn("w:t"), qn("m:t")):
        if t.text and PUA_RE.search(t.text):
            t.text = "".join(LIT_REVERSE.get(c, "" if "" <= c <= "" else c) for c in t.text)
    doc.save(str(docx_path))
    return ctx.stats


# =====================================================================================
# PIPELINE (paste branch only)
# =====================================================================================
STANDALONE_HEAD = (
    "\\documentclass[12pt,a4paper]{article}\n"
    "\\usepackage{amsmath,amssymb,graphicx,tikz,tkz-tab}\n"
    "\\begin{document}\n"
)
STANDALONE_TAIL = "\n\\end{document}\n"
_VERB_STRIP_RE = re.compile(
    r"\\begin\{(verbatim\*?|Verbatim|lstlisting|minted|comment)\}.*?\\end\{\1\}|\\verb\*?(.).*?\2|%[^\n]*", re.S)


def _log(msg: str) -> None:
    try:
        print(f"[smart_paste] {msg}", file=sys.stderr, flush=True)
    except Exception:
        pass


def _call_engine(fn, tex: Path, output: Path, **kwargs: Any) -> dict[str, Any]:
    """Call the user's convert_latex_file_unified with only the keyword args its signature accepts."""
    try:
        params = inspect.signature(fn).parameters
        if not any(p.kind == p.VAR_KEYWORD for p in params.values()):
            kwargs = {k: v for k, v in kwargs.items() if k in params}
    except (TypeError, ValueError):
        pass
    return fn(tex, output, **kwargs)


def _docx_plain_text(path: Path) -> str:
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="replace")
        return "".join(re.findall(r"<(?:w|m):t(?:\s[^>]*)?>([^<]*)</(?:w|m):t>", xml))
    except Exception:
        return ""


def _leftover_raw_commands(src: str, processed: str, docx_text: str) -> list[str]:
    """Commands that the preprocessor rewrote but that still appear literally in Word → engine ignored our file."""
    clean = _VERB_STRIP_RE.sub("", src)
    orig = set(re.findall(r"\\([A-Za-z@]+)", clean))
    kept = set(re.findall(r"\\([A-Za-z@]+)", processed))
    cands = sorted(n for n in orig - kept if len(n) >= 2)
    return [n for n in cands if re.search(r"\\" + re.escape(n) + r"(?![A-Za-z])", docx_text)]


# ─────────── LỚP BẢO HIỂM PHỦ 100% (backend/universal_fallback.py) ───────────
# Nhập mềm: thiếu module hoặc module lỗi thì mọi thứ chạy y như trước khi có lớp này.
try:
    from . import universal_fallback as _universal
except Exception as _uf_exc:  # pragma: no cover
    _universal = None
    _log(f"không nạp được universal_fallback ({_uf_exc}) → chạy như cũ")


def _universal_protect(src: str, base_dir: Path, workdir: Path | None = None):
    """Thay khối lạ bằng chỗ đánh dấu + khởi động dựng ảnh ở luồng nền."""
    if _universal is None:
        return src, None
    try:
        return _universal.protect(src, base_dir=base_dir, workdir=workdir)
    except Exception as exc:
        _log(f"universal.protect lỗi → bỏ qua: {type(exc).__name__}: {exc}")
        return src, None


def _universal_restore(output: Path, ctx) -> dict[str, Any]:
    """Đổi chỗ đánh dấu trong .docx thành ảnh vector do TeX dựng."""
    if _universal is None or ctx is None:
        return {}
    try:
        return _universal.restore(output, ctx)
    except Exception as exc:
        _log(f"universal.restore lỗi → bỏ qua: {type(exc).__name__}: {exc}")
        return {}


def _universal_abort(ctx) -> None:
    """Thu dọn khi pipeline rơi về engine gốc."""
    try:
        if ctx is not None and getattr(ctx, "renderer", None) is not None:
            ctx.renderer.close()
    except Exception:
        pass


def convert_pasted_latex_unified(
    tex: Path,
    output: Path,
    *,
    engine: str = "xelatex",
    dpi: int = 400,
    reference: Path | None = None,
    strict: bool = False,
    qa_pdf: bool = False,
    mode: int = 2,
    show_solutions: bool = True,
) -> dict[str, Any]:
    """Drop-in replacement of `convert_latex_file_unified` for pasted LaTeX code.

    The preprocessed source is written (wrapped as a standalone document) to the SAME path the engine
    would read, so the engine always converts the preprocessed text whatever its internal file logic is;
    the original pasted source is restored right after (validators / previews still see the original).
    """
    from .core_converter import convert_docx_to_pdf, convert_latex_file_unified

    tex = Path(tex).resolve()
    output = Path(output).resolve()
    engine_kwargs = {"engine": engine, "dpi": dpi, "reference": reference, "strict": strict, "mode": mode}
    if os.environ.get("L2W_SMART_PASTE", "1").strip() in ("0", "false", "off", "no"):
        return _call_engine(convert_latex_file_unified, tex, output, qa_pdf=qa_pdf, **engine_kwargs)

    original_bytes = tex.read_bytes()
    src = original_bytes.decode("utf-8-sig", errors="replace")
    pre = None
    # Bước 0 — khối nào engine chưa dựng được thì giao thẳng cho TeX dựng.
    uf_src, uf_ctx = _universal_protect(src, tex.parent, workdir=tex.parent / "_l2w_truth")
    try:
        pre = preprocess_latex(uf_src, base_dir=tex.parent, show_solutions=show_solutions)
        tex.write_text(STANDALONE_HEAD + pre.text + STANDALONE_TAIL, encoding="utf-8")
        if os.environ.get("L2W_KEEP_SMART_TEX"):
            tex.with_name(tex.stem + ".l2w_smart.tex").write_text(STANDALONE_HEAD + pre.text + STANDALONE_TAIL,
                                                                   encoding="utf-8")
        try:
            report = _call_engine(convert_latex_file_unified, tex, output, qa_pdf=False, **engine_kwargs)
        finally:
            tex.write_bytes(original_bytes)  # always give the original source back
        layout_stats = apply_layout_markers(output, pre)
        universal_stats = _universal_restore(output, uf_ctx)
    except Exception as exc:  # fallback: exactly the original engine behaviour
        _universal_abort(uf_ctx)
        if tex.read_bytes() != original_bytes:
            tex.write_bytes(original_bytes)
        _log(f"LỖI → dùng engine gốc: {type(exc).__name__}: {exc}")
        _log(traceback.format_exc())
        report = _call_engine(convert_latex_file_unified, tex, output, qa_pdf=qa_pdf, **engine_kwargs)
        report["smart_paste"] = {"status": "fallback_original_engine", "error": f"{type(exc).__name__}: {exc}",
                                 "trace": traceback.format_exc()[-3000:]}
        report["smart_paste_status"] = "LỖI → engine gốc"
        report.setdefault("warnings", []).append("Tiền xử lý thông minh lỗi, đã dùng engine gốc: " + str(exc))
        output.with_suffix(".report.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
        return report
    leftovers = _leftover_raw_commands(src, pre.text, _docx_plain_text(output))
    report["input"] = str(tex)
    report["smart_paste"] = {"status": "ok" if not leftovers else "partially_applied", **pre.summary(),
                             "layout": layout_stats, "raw_commands_left_in_word": leftovers,
                             "universal_fallback": {**(uf_ctx.summary() if uf_ctx is not None else
                                                       {"bat": False, "ly_do": "chua nap duoc module"}),
                                                    "ket_qua": universal_stats}}
    report["smart_paste_status"] = "OK" if not leftovers else "CÒN LỆNH THÔ: " + ", ".join(leftovers[:5])
    if leftovers:
        _log("Engine vẫn in nguyên văn các lệnh: " + ", ".join(leftovers))
    else:
        st = pre.stats
        _log(f"OK — {st.get('macros_expanded', 0)} macro mở rộng, {st.get('questions', 0) + st.get('exam_questions', 0)} câu hỏi, "
             f"{st.get('choice_tables', 0)} bảng trắc nghiệm, {len(pre.unknown_commands)} lệnh lạ tự xử lý")
    if pre.warnings:
        report["warnings"] = list(report.get("warnings", [])) + [w for w in pre.warnings[:50]]
        report["status"] = "review_required"
    if qa_pdf:
        render_dir = output.parent / (output.stem + "_qa")
        render_dir.mkdir(exist_ok=True)
        target_pdf = render_dir / (output.stem + ".pdf")
        if convert_docx_to_pdf(output, target_pdf) and target_pdf.exists():
            report["qa_rendered_pdf"] = str(target_pdf)
    output.with_suffix(".report.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    return report


