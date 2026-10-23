"""`backend/tex_blocks.py` — chuẩn bị nguồn TeX cho engine (nhánh tải tệp) và khôi phục khối sau khi dựng.

Dùng khi pipeline chạy engine `convert_latex_file_unified` trực tiếp (không có tiền xử lý thông minh
`latex_preprocessor`/`tex_truth` của nhánh dán mã).  Mọi biến đổi chạy trên BẢN PHẲNG trong workspace
của job (không ghi tệp người dùng) và được ghi vào báo cáo:

1. chế độ học sinh: bỏ dấu đáp án có quy ước rõ (`\\True`, `\\choiceA..D`, cột đáp án `\\choiceTF`, `\\loigiai`);
2. nở macro người dùng ĐƠN GIẢN định nghĩa ở preamble (đúng định nghĩa của họ, kể cả `\\vv`…);
   macro có điều kiện / lệnh nội bộ TeX → không nở, ghi lý do;
3. danh sách itemize/enumerate/description → đoạn có nhãn (1., a), •…) thay vì mất số thứ tự;
4. môi trường lạ → bóc vỏ giữ nội dung (định lý/ex có tiêu đề); lệnh lạ → bỏ tên, giữ đối số;
5. hình: `\\includegraphics` → chỗ đánh dấu + ảnh đã chuyển định dạng (PNG/JPEG/PDF→SVG+PNG/SVG);
   TikZ → thử DrawingML của engine, đối chiếu với ảnh TeX thật; không đạt → ảnh vector TeX;
   bảng biến thiên tkz-tab ngoài phạm vi native → ảnh TeX (table_as_image);
   `\\immini{…}{hình không phải TikZ}` → dựng lại bảng 2 cột sau khi engine chạy (engine cũ tự vẽ một đồ thị
   mặc định cho cột phải không phải TikZ — bịa hình).
Sau khi engine dựng xong, `inject_engine_blocks` thay chỗ đánh dấu bằng ảnh, dựng lại bố cục immini và
kiểm tra từng chỗ đánh dấu: chỗ nào mất → khối `unresolved` (không bao giờ im lặng).
"""
from __future__ import annotations

import io
import json
import re
import shutil
import subprocess
import tempfile
import time
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from .source_resolver import FlattenResult, find_matching_brace, mask_tex

__all__ = ["EngineBlock", "PreparedTex", "prepare_tex_for_engine", "inject_engine_blocks", "expand_user_macros",
           "convert_lists", "IMAGE_EXTS"]

EMU_PER_PT = 12700
IMAGE_EXTS = [".pdf", ".png", ".jpg", ".jpeg", ".eps", ".svg", ".gif", ".bmp", ".tif", ".tiff",
              ".PDF", ".PNG", ".JPG", ".JPEG", ".EPS", ".SVG"]
LAYOUT_PREFIX = "ZLZLAYZ"


@dataclass
class EngineBlock:
    block_id: str
    kind: str                      # includegraphics | figure | bbt | immini_layout
    source: dict[str, Any]
    tex: str
    index: int | None = None       # số thứ tự chỗ đánh dấu ZLZIMGZ#####ZEND
    method: str = "pending"        # drawingml | table_native | vector_image | raster_image | unresolved
    status: str = "pending"        # ok | fallback | unresolved
    reason: str = ""
    warnings: list[str] = field(default_factory=list)
    asset: dict[str, Any] | None = None
    max_width_frac: float = 1.0
    inline: bool = False
    verify: dict[str, Any] | None = None
    parent: str | None = None

    def to_dict(self) -> dict[str, Any]:
        d = asdict(self)
        d["tex"] = self.tex[:600]
        if d.get("asset"):
            d["asset"] = {k: v for k, v in d["asset"].items() if k not in ("data", "png_data")}
        return d


@dataclass
class PreparedTex:
    text: str
    validation_text: str
    blocks: list[EngineBlock] = field(default_factory=list)
    transforms: dict[str, Any] = field(default_factory=dict)
    issues: list[dict[str, Any]] = field(default_factory=list)
    assets_dir: Path | None = None
    verbatim: dict[int, dict[str, Any]] = field(default_factory=dict)
    #: văn bản sau chuẩn hoá cấu trúc (có nhãn do pipeline sinh: "Câu n.", "Bảng n.", "Lời giải."…) — chỉ dùng để
    #: validator KHÔNG tính các nhãn này là "lặp nội dung"; độ phủ vẫn đo trên validation_text.
    normalized_text: str = ""


def _issue(prep: PreparedTex, sev: str, code: str, msg: str, **kw: Any) -> None:
    prep.issues.append({"severity": sev, "code": code, "message": msg, **kw})


# ───────────────────────────── tiện ích vị trí ─────────────────────────────
def _body_span(text: str, masked: str) -> tuple[int, int]:
    mb = re.search(r"\\begin\s*\{document\}", masked)
    if not mb:
        return 0, len(text)
    me = re.search(r"\\end\s*\{document\}", masked[mb.end():])
    return mb.end(), (mb.end() + me.start()) if me else len(text)


_MATH_ENV_NAMES = ("equation", "equation*", "align", "align*", "gather", "gather*", "multline", "multline*",
                   "eqnarray", "eqnarray*", "flalign", "flalign*", "alignat", "alignat*", "displaymath", "math")


def _math_spans(masked: str, a: int, b: int) -> list[tuple[int, int]]:
    spans: list[tuple[int, int]] = []
    i = a
    env_re = re.compile(r"\\begin\{(" + "|".join(re.escape(e) for e in _MATH_ENV_NAMES) + r")\}")
    while i < b:
        c = masked[i]
        if c == "\\":
            if masked.startswith("\\[", i):
                e = masked.find("\\]", i + 2)
                if 0 < e <= b:
                    spans.append((i, e + 2))
                    i = e + 2
                    continue
            if masked.startswith("\\(", i):
                e = masked.find("\\)", i + 2)
                if 0 < e <= b:
                    spans.append((i, e + 2))
                    i = e + 2
                    continue
            m = env_re.match(masked, i)
            if m:
                tok = "\\end{" + m.group(1) + "}"
                e = masked.find(tok, m.end())
                if 0 < e <= b:
                    spans.append((i, e + len(tok)))
                    i = e + len(tok)
                    continue
            i += 2
            continue
        if c == "$":
            if masked.startswith("$$", i):
                e = masked.find("$$", i + 2)
                if 0 < e <= b:
                    spans.append((i, e + 2))
                    i = e + 2
                    continue
            j = i + 1
            while j < b:
                if masked[j] == "\\":
                    j += 2
                    continue
                if masked[j] == "$":
                    break
                j += 1
            if j < b:
                spans.append((i, j + 1))
                i = j + 1
                continue
        i += 1
    return spans


def _env_spans(masked: str, names: tuple[str, ...] | str, a: int = 0, b: int | None = None) -> list[tuple[int, int, str]]:
    """Môi trường ngoài cùng (có tính lồng cùng tên) trong [a, b)."""
    b = len(masked) if b is None else b
    if isinstance(names, str):
        pat = names
    else:
        pat = "|".join(re.escape(n) for n in names)
    out: list[tuple[int, int, str]] = []
    pos = a
    beg_re = re.compile(r"\\begin\s*\{(" + pat + r")\}")
    while True:
        m = beg_re.search(masked, pos, b)
        if not m:
            break
        name = m.group(1)
        depth = 1
        scan = m.end()
        tb, te = "\\begin{" + name + "}", "\\end{" + name + "}"
        end = -1
        while True:
            nb = masked.find(tb, scan, b)
            ne = masked.find(te, scan, b)
            if ne < 0:
                break
            if 0 <= nb < ne:
                depth += 1
                scan = nb + len(tb)
                continue
            depth -= 1
            if depth == 0:
                end = ne + len(te)
                break
            scan = ne + len(te)
        if end < 0:
            pos = m.end()
            continue
        out.append((m.start(), end, name))
        pos = end
    return out


def _inside(spans: list[tuple[int, ...]], i: int) -> bool:
    return any(s[0] <= i < s[1] for s in spans)


def _read_arg(text: str, pos: int) -> tuple[str | None, int]:
    """Đối số TeX không phân cách: {nhóm} | \\lệnh | một ký tự. Trả (nội dung, vị trí sau)."""
    k = pos
    while k < len(text) and text[k] in " \t\n":
        k += 1
    if k >= len(text):
        return None, pos
    if text[k] == "{":
        e = find_matching_brace(text, k)
        if e < 0:
            return None, pos
        return text[k + 1: e], e + 1
    if text[k] == "\\":
        m = re.match(r"\\(?:[A-Za-z@]+|.)", text[k:])
        return (m.group(0), k + m.end()) if m else (None, pos)
    if text[k] in "}":
        return None, pos
    return text[k], k + 1


def _read_opt(text: str, pos: int) -> tuple[str | None, int]:
    k = pos
    while k < len(text) and text[k] in " \t":
        k += 1
    if k < len(text) and text[k] == "[":
        depth = 0
        j = k
        while j < len(text):
            if text[j] == "{":
                depth += 1
            elif text[j] == "}":
                depth -= 1
            elif text[j] == "]" and depth == 0:
                return text[k + 1: j], j + 1
            j += 1
    return None, pos


def _apply(text: str, repls: list[tuple[int, int, str]]) -> str:
    out = []
    last = 0
    for a, b, s in sorted(repls, key=lambda r: r[0]):
        if a < last:
            continue
        out.append(text[last:a])
        out.append(s)
        last = b
    out.append(text[last:])
    return "".join(out)


# ───────────────────────────── 1. chế độ học sinh ─────────────────────────────
def strip_answers_for_student(text: str, show_solutions: bool) -> tuple[str, dict[str, int]]:
    masked = mask_tex(text)
    a, b = _body_span(text, masked)
    stats = {"true_markers": 0, "choice_letter_macros": 0, "choiceTF_answers": 0, "solutions_hidden": 0}
    repls: list[tuple[int, int, str]] = []
    for m in re.finditer(r"\\True(?![A-Za-z@])", masked[a:b]):
        repls.append((a + m.start(), a + m.end(), ""))
        stats["true_markers"] += 1
    for m in re.finditer(r"\\choice([ABCD])(?![A-Za-z@])", masked[a:b]):
        repls.append((a + m.start(), a + m.end(), "\\choice"))
        stats["choice_letter_macros"] += 1
    for m in re.finditer(r"\\choiceTF(?![A-Za-z@])", masked[a:b]):
        pos = a + m.end()
        args: list[tuple[int, int]] = []
        for _ in range(8):
            k = pos
            while k < len(text) and text[k] in " \t\n":
                k += 1
            if k < len(text) and text[k] == "{":
                e = find_matching_brace(text, k)
                if e < 0:
                    break
                args.append((k, e))
                pos = e + 1
            else:
                break
        if len(args) == 8:
            for idx in (1, 3, 5, 7):
                s, e = args[idx]
                repls.append((s + 1, e, ""))
            stats["choiceTF_answers"] += 4
    if not show_solutions:
        for m in re.finditer(r"\\loigiai\s*\{", masked[a:b]):
            ob = a + m.end() - 1
            e = find_matching_brace(text, ob)
            if e > 0:
                repls.append((a + m.start(), e + 1, ""))
                stats["solutions_hidden"] += 1
        for s, e, _n in _env_spans(masked, ("loigiai", "huongdan", "dapan"), a, b):
            repls.append((s, e, ""))
            stats["solutions_hidden"] += 1
    return _apply(text, repls), stats


# ───────────────────────────── 2. nở macro người dùng ─────────────────────────────
_UNSAFE_BODY = re.compile(r"\\(?:if[A-Za-z@]*|fi|else|or|csname|endcsname|expandafter|noexpand|def|gdef|edef|xdef|let|"
                          r"global|catcode|begingroup|endgroup|newcommand|renewcommand|providecommand|the|advance|"
                          r"setcounter|addtocounter|stepcounter|refstepcounter|NewDocumentCommand|futurelet|"
                          r"afterassignment|aftergroup|relax|unexpanded|detokenize|string|makeatletter|makeatother|"
                          r"@[A-Za-z@]*)(?![A-Za-z])|[A-Za-z]@|@[A-Za-z]")
_DEF_RE = re.compile(r"\\(newcommand|renewcommand|providecommand|DeclareRobustCommand)\s*\*?\s*(?:\{\s*\\([A-Za-z@]+)\s*\}|\\([A-Za-z@]+))")
_TEXDEF_RE = re.compile(r"\\def\s*\\([A-Za-z@]+)([^{}\n%]*?)\s*\{")
_MATHOP_RE = re.compile(r"\\DeclareMathOperator(\*?)\s*\{\s*\\([A-Za-z@]+)\s*\}\s*\{")
_XPARSE_RE = re.compile(r"\\(NewDocumentCommand|RenewDocumentCommand|ProvideDocumentCommand|DeclareDocumentCommand)"
                        r"\s*(?:\{\s*\\([A-Za-z@]+)\s*\}|\\([A-Za-z@]+))\s*\{")
_LET_RE = re.compile(r"\\let\s*\\([A-Za-z@]+)\s*=?\s*\\([A-Za-z@]+)(?![A-Za-z@])")
_PAIRED_RE = re.compile(r"\\DeclarePairedDelimiter\s*\{?\s*\\([A-Za-z@]+)\s*\}?\s*\{")
_ALIAS_BLOCKED = {"relax", "undefined", "empty", "par", "fi", "else", "or", "begingroup", "endgroup", "bgroup", "egroup"}

#: Định nghĩa của người dùng THẮNG khuôn của engine (đúng như TeX). Ngoại lệ: \immini (pipeline dựng bố cục
#: riêng) và \True (thành ngữ đánh dấu đáp án — có điều kiện TeX, do macro_fidelity xử lý).
_ENGINE_TEMPLATES = {"immini", "True"}


@dataclass
class _Macro:
    name: str
    nargs: int
    body: str
    opt_default: str | None
    #: \def có tham số phân cách: (chữ đứng trước #1, [chữ phân cách sau #i hoặc None]) — None nếu không phân cách
    delimited: tuple[str, list[str | None]] | None = None
    #: \DeclarePairedDelimiter: nhận thêm dạng * và [\big] trước đối số
    paired: bool = False
    #: xparse (\NewDocumentCommand): danh sách đặc tả đối số ("s" | "m" | "o" | ("O", mặc định))
    xspec: list | None = None


def _parse_xparse_spec(spec: str) -> list | None:
    """Tập con đặc tả xparse: s, m, o, O{mặc định}; bỏ qua tiền tố + (long) và ! (không bỏ khoảng trắng).
    Đặc tả khác (t, r, d, e, v, …) → None (không nở, ghi báo cáo)."""
    out: list = []
    i, n = 0, len(spec)
    while i < n:
        c = spec[i]
        if c in " \t\n+!":
            i += 1
            continue
        if c in "smo":
            out.append(c)
            i += 1
            continue
        if c == "O":
            j = i + 1
            while j < n and spec[j] in " \t":
                j += 1
            if j < n and spec[j] == "{":
                e = find_matching_brace(spec, j)
                if e < 0:
                    return None
                out.append(("O", spec[j + 1: e]))
                i = e + 1
                continue
            return None
        return None
    return out if len(out) <= 9 else None


_XP_NOVALUE = "-NoValue-"
_XP_TRUE = "\\BooleanTrue"
_XP_FALSE = "\\BooleanFalse"


def _xparse_conditionals(body: str) -> str:
    r"""Rút gọn \IfBooleanTF/\IfValueTF/\IfNoValueTF (và dạng T/F) sau khi đã thay đối số."""
    pat = re.compile(r"\\If(Boolean|Value|NoValue)(TF|T|F)(?![A-Za-z@])")
    for _ in range(200):
        m = pat.search(body)
        if not m:
            break
        kind, form = m.group(1), m.group(2)
        pos = m.end()
        val, pos = _read_arg(body, pos)
        if val is None:
            break
        branches = []
        for _b in range(2 if form == "TF" else 1):
            br, pos = _read_arg(body, pos)
            if br is None:
                return body
            branches.append(br)
        v = val.strip()
        if kind == "Boolean":
            truth = v == _XP_TRUE
        elif kind == "Value":
            truth = v != _XP_NOVALUE
        else:
            truth = v == _XP_NOVALUE
        if form == "TF":
            rep = branches[0] if truth else branches[1]
        elif form == "T":
            rep = branches[0] if truth else ""
        else:
            rep = branches[0] if not truth else ""
        body = body[:m.start()] + rep + body[pos:]
    return body


def _parse_def_params(params: str) -> tuple[str, list[str | None]] | None:
    """`#1;#2;` → ("", [";", ";"]); `#1#2` → ("", [None, None]); `(#1,#2)` → ("(", [",", ")"]).
    Trả None nếu số thứ tự tham số không liên tiếp 1..n."""
    parts = re.split(r"#(\d)", params)
    lead = parts[0]
    nums = parts[1::2]
    seps = parts[2::2]
    if nums != [str(i) for i in range(1, len(nums) + 1)]:
        return None
    delims: list[str | None] = []
    for sep in seps:
        sep = sep.strip() if sep.strip() else sep
        delims.append(sep if sep else None)
    return lead.strip(), delims


def _read_delimited(text: str, pos: int, delim: str) -> tuple[str | None, int]:
    """Đọc tới `delim` đầu tiên ở độ sâu ngoặc 0 (đúng quy tắc tham số phân cách của TeX, bỏ một cặp ngoặc bao)."""
    depth = 0
    k = pos
    while k < len(text):
        c = text[k]
        if c == "\\":
            k += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth < 0:
                return None, pos
        elif depth == 0 and text.startswith(delim, k):
            val = text[pos:k]
            st = val.strip()
            if st.startswith("{") and find_matching_brace(st, 0) == len(st) - 1:
                val = st[1:-1]
            return val, k + len(delim)
        if c == "\n" and text.startswith("\n\n", k) and depth == 0:
            return None, pos  # \par giữa tham số phân cách → TeX báo lỗi; không đoán
        k += 1
    return None, pos


def _parse_macro_defs(preamble: str) -> tuple[dict[str, _Macro], dict[str, str]]:
    masked = mask_tex(preamble)
    macros: dict[str, _Macro] = {}
    skipped: dict[str, str] = {}

    def depth_at(i: int) -> int:
        d = 0
        k = 0
        while k < i:
            c = masked[k]
            if c == "\\":
                k += 2
                continue
            if c == "{":
                d += 1
            elif c == "}":
                d -= 1
            k += 1
        return d

    for m in _DEF_RE.finditer(masked):
        name = m.group(2) or m.group(3)
        pos = m.end()
        nargs = 0
        opt = None
        o, p2 = _read_opt(preamble, pos)
        if o is not None:
            try:
                nargs = int(o.strip())
            except ValueError:
                skipped[name] = "số đối số không đọc được"
                continue
            pos = p2
            o2, p3 = _read_opt(preamble, pos)
            if o2 is not None:
                opt = o2
                pos = p3
        k = pos
        while k < len(preamble) and preamble[k] in " \t\n":
            k += 1
        if k >= len(preamble) or preamble[k] != "{":
            skipped[name] = "thân macro không có ngoặc nhọn"
            continue
        e = find_matching_brace(preamble, k)
        if e < 0:
            skipped[name] = "thân macro không đóng ngoặc"
            continue
        if depth_at(m.start()) != 0:
            skipped[name] = "định nghĩa trong nhóm (phạm vi cục bộ)"
            continue
        body = preamble[k + 1: e]
        if m.group(1) == "providecommand" and name in macros:
            continue
        macros[name] = _Macro(name, nargs, body, opt)
    for m in _TEXDEF_RE.finditer(masked):
        name = m.group(1)
        params = m.group(2)
        ob = m.end() - 1
        e = find_matching_brace(preamble, ob)
        if e < 0 or depth_at(m.start()) != 0:
            skipped[name] = "\\def không đọc được hoặc nằm trong nhóm"
            continue
        spec = _parse_def_params(params) if params else ("", [])
        if spec is None:
            skipped[name] = "\\def có thứ tự tham số không hợp lệ"
            continue
        lead, delims = spec
        if lead or any(d is not None for d in delims):
            if any(d is not None and ("\\" in d or "{" in d or "}" in d) for d in delims) or "\\" in lead:
                skipped[name] = "\\def có chuỗi phân cách là lệnh — không nở"
                continue
            macros[name] = _Macro(name, len(delims), preamble[ob + 1: e], None, delimited=(lead, delims))
        else:
            macros[name] = _Macro(name, len(delims), preamble[ob + 1: e], None)
    for m in _XPARSE_RE.finditer(masked):
        name = m.group(2) or m.group(3)
        ob = m.end() - 1
        e1 = find_matching_brace(preamble, ob)
        if e1 < 0:
            skipped[name] = "\\NewDocumentCommand không đọc được đặc tả"
            continue
        spec = _parse_xparse_spec(preamble[ob + 1: e1])
        k = e1 + 1
        while k < len(preamble) and preamble[k] in " \t\n":
            k += 1
        e2 = find_matching_brace(preamble, k) if k < len(preamble) and preamble[k] == "{" else -1
        if spec is None or e2 < 0 or depth_at(m.start()) != 0:
            skipped[name] = "đặc tả xparse ngoài tập hỗ trợ (s, m, o, O{…})"
            continue
        if m.group(1) == "ProvideDocumentCommand" and name in macros:
            continue
        macros[name] = _Macro(name, len(spec), preamble[k + 1: e2], None, xspec=spec)
    for m in _PAIRED_RE.finditer(masked):
        name = m.group(1)
        ob = m.end() - 1
        e1 = find_matching_brace(preamble, ob)
        if e1 < 0:
            continue
        k = e1 + 1
        while k < len(preamble) and preamble[k] in " \t\n":
            k += 1
        if k >= len(preamble) or preamble[k] != "{":
            continue
        e2 = find_matching_brace(preamble, k)
        if e2 < 0:
            continue
        left, right = preamble[ob + 1: e1].strip(), preamble[k + 1: e2].strip()
        macros[name] = _Macro(name, 1, f"\\left{left} #1 \\right{right}", None, paired=True)
    for m in _LET_RE.finditer(masked):
        name, target = m.group(1), m.group(2)
        if target in _ALIAS_BLOCKED or target.startswith("@") or "@" in name or depth_at(m.start()) != 0:
            skipped[name] = "\\let tới lệnh điều khiển/nội bộ — không nở"
            continue
        if target in macros:
            src = macros[target]
            macros[name] = _Macro(name, src.nargs, src.body, src.opt_default, src.delimited, src.paired, src.xspec)
        else:
            macros[name] = _Macro(name, 0, "\\" + target, None)
    for m in _MATHOP_RE.finditer(masked):
        name = m.group(2)
        ob = m.end() - 1
        e = find_matching_brace(preamble, ob)
        if e > 0:
            star = m.group(1)
            macros[name] = _Macro(name, 0, ("\\operatorname*{" if star else "\\operatorname{") + preamble[ob + 1: e] + "}", None)
    for name in list(macros):
        mc = macros[name]
        if name in _ENGINE_TEMPLATES:
            skipped[name] = "lệnh có khuôn dựng riêng của engine/đề thi"
            del macros[name]
        elif _UNSAFE_BODY.search(mc.body):
            skipped[name] = "thân có lệnh điều khiển/điều kiện TeX"
            del macros[name]
        elif re.search(r"\\" + re.escape(name) + r"(?![A-Za-z@])", mc.body):
            skipped[name] = "macro đệ quy"
            del macros[name]
        elif mc.nargs > 9:
            skipped[name] = "quá 9 đối số"
            del macros[name]
    return macros, skipped


def expand_user_macros(text: str, max_passes: int = 12, max_expansions: int = 50000) -> tuple[str, dict[str, Any]]:
    masked = mask_tex(text)
    a, b = _body_span(text, masked)
    preamble = text[:a]
    macros, skipped = _parse_macro_defs(preamble)
    body_masked = masked[a:b]
    for m in _DEF_RE.finditer(body_masked):
        name = m.group(2) or m.group(3)
        if name in macros:
            del macros[name]
            skipped[name] = "được định nghĩa lại trong thân tài liệu (phạm vi thay đổi) → không nở"
    stats: dict[str, Any] = {"expanded": {}, "skipped": skipped, "defined": sorted(macros)}
    if not macros:
        return text, stats
    names = sorted(macros, key=len, reverse=True)
    pat = re.compile(r"\\(" + "|".join(re.escape(n) for n in names) + r")(?![A-Za-z@])")
    body = text[a:b]
    total = 0
    for _ in range(max_passes):
        bm = mask_tex(body, keep_iffalse=True)
        repls: list[tuple[int, int, str]] = []
        for m in pat.finditer(bm):
            if repls and m.start() < repls[-1][1]:
                continue
            mc = macros[m.group(1)]
            pos = m.end()
            args: list[str] = []
            if mc.delimited is not None:
                lead, delims = mc.delimited
                k = pos
                while k < len(body) and body[k] in " \t":
                    k += 1
                if lead:
                    if not body.startswith(lead, k):
                        continue
                    k += len(lead)
                ok = True
                for d in delims:
                    if d is None:
                        arg, k2 = _read_arg(body, k)
                    else:
                        arg, k2 = _read_delimited(body, k, d)
                    if arg is None:
                        ok = False
                        break
                    args.append(arg)
                    k = k2
                if not ok:
                    continue
                rep = mc.body
                for i in range(len(args), 0, -1):
                    rep = rep.replace(f"#{i}", args[i - 1])
                repls.append((m.start(), k, rep.replace("##", "#")))
                stats["expanded"][mc.name] = stats["expanded"].get(mc.name, 0) + 1
                continue
            if mc.xspec is not None:
                k = pos
                ok = True
                for sp in mc.xspec:
                    if sp == "s":
                        j = k
                        while j < len(body) and body[j] in " \t":
                            j += 1
                        if j < len(body) and body[j] == "*":
                            args.append(_XP_TRUE)
                            k = j + 1
                        else:
                            args.append(_XP_FALSE)
                    elif sp == "m":
                        arg, k2 = _read_arg(body, k)
                        if arg is None:
                            ok = False
                            break
                        args.append(arg)
                        k = k2
                    else:
                        o, k2 = _read_opt(body, k)
                        if o is not None:
                            args.append(o)
                            k = k2
                        else:
                            args.append(sp[1] if isinstance(sp, tuple) else _XP_NOVALUE)
                if not ok:
                    continue
                rep = mc.body
                for i in range(len(args), 0, -1):
                    rep = rep.replace(f"#{i}", args[i - 1])
                rep = _xparse_conditionals(rep.replace("##", "#"))
                repls.append((m.start(), k, rep))
                stats["expanded"][mc.name] = stats["expanded"].get(mc.name, 0) + 1
                continue
            if mc.paired:
                k = pos
                if k < len(body) and body[k] == "*":
                    k += 1
                o, k2 = _read_opt(body, k)
                if o is not None:
                    k = k2
                pos = k
            if mc.opt_default is not None:
                o, p2 = _read_opt(body, pos)
                if o is not None:
                    args.append(o)
                    pos = p2
                else:
                    args.append(mc.opt_default)
            ok = True
            while len(args) < mc.nargs:
                arg, p2 = _read_arg(body, pos)
                if arg is None:
                    ok = False
                    break
                args.append(arg)
                pos = p2
            if not ok:
                continue
            rep = mc.body
            for i in range(mc.nargs, 0, -1):
                rep = rep.replace(f"#{i}", args[i - 1])
            rep = rep.replace("##", "#")
            # giữ ranh giới token: lệnh chữ theo sau bởi chữ cái cần khoảng trắng
            if mc.nargs == 0 and pos < len(body) and body[pos].isalpha() and re.search(r"\\[A-Za-z@]+$", rep):
                rep += " "
            repls.append((m.start(), pos, rep))
            stats["expanded"][mc.name] = stats["expanded"].get(mc.name, 0) + 1
        if not repls:
            break
        total += len(repls)
        body = _apply(body, repls)
        if total > max_expansions:
            stats["truncated"] = True
            break
    return text[:a] + body + text[b:], stats


# ───────────────────────────── 3. danh sách ─────────────────────────────
def _alph(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(97 + r) + s
    return s


def _roman(n: int) -> str:
    vals = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"), (50, "l"), (40, "xl"),
            (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out


def _enum_label(template: str | None, n: int, depth: int) -> str:
    if template is None:
        template = {1: "\\arabic*.", 2: "(\\alph*)", 3: "\\roman*.", 4: "\\Alph*."}.get(depth, "\\arabic*.")
    t = template
    t = t.replace("\\arabic*", str(n)).replace("\\alph*", _alph(n)).replace("\\Alph*", _alph(n).upper())
    t = t.replace("\\roman*", _roman(n)).replace("\\Roman*", _roman(n).upper())
    return t


def _shorthand_enum_template(opt: str) -> str | None:
    """Gói enumerate: [a)] [(i)] [A.] [1.] — ký tự đầu tiên a/A/i/I/1 là bộ đếm."""
    m = re.search(r"[aAiI1]", opt)
    if not m:
        return None
    ch = m.group(0)
    rep = {"a": "\\alph*", "A": "\\Alph*", "i": "\\roman*", "I": "\\Roman*", "1": "\\arabic*"}[ch]
    return opt[: m.start()] + rep + opt[m.end():]


def convert_lists(text: str) -> tuple[str, dict[str, int]]:
    stats = {"itemize": 0, "enumerate": 0, "description": 0, "items": 0}
    names = ("itemize", "enumerate", "description", "itemize*", "enumerate*", "tasks")
    guard = 0
    while guard < 500:
        guard += 1
        masked = mask_tex(text)
        a, b = _body_span(text, masked)
        spans = _env_spans(masked, names, a, b)
        # chọn môi trường trong cùng (không chứa danh sách khác)
        inner = None
        for s, e, n in spans:
            sub = _env_spans(masked, names, s + 1, e - 1)
            if not sub:
                inner = (s, e, n)
                break
        if inner is None:
            # có thể danh sách ngoài cùng chứa danh sách → tìm sâu hơn
            found = None
            for s, e, n in spans:
                stack = [(s, e, n)]
                while stack:
                    cs, ce, cn = stack.pop()
                    sub = _env_spans(masked, names, cs + 1, ce - 1)
                    if not sub:
                        found = (cs, ce, cn)
                        break
                    stack.extend(sub)
                if found:
                    break
            inner = found
        if inner is None:
            break
        s, e, name = inner
        depth = 1 + sum(1 for ss, ee, _ in _all_list_spans(masked, names, a, b) if ss < s and e <= ee)
        begin_m = re.match(r"\\begin\s*\{" + re.escape(name) + r"\}", text[s:])
        pos = s + begin_m.end()
        opt, pos2 = _read_opt(text, pos)
        if name == "tasks":
            # tasks: \begin{tasks}(n) ... \task ...
            mcol = re.match(r"\s*\(\d+\)", text[pos2:])
            if mcol:
                pos2 += mcol.end()
        content = text[pos2: e - len("\\end{" + name + "}")]
        tmpl = None
        start = 1
        if opt:
            ml = re.search(r"label\s*=\s*(\{(?:[^{}]|\{[^{}]*\})*\}|[^,\]]+)", opt)
            if ml:
                tmpl = ml.group(1).strip()
                if tmpl.startswith("{") and tmpl.endswith("}"):
                    tmpl = tmpl[1:-1]
            elif name.startswith("enumerate") and "=" not in opt:
                tmpl = _shorthand_enum_template(opt)
            ms = re.search(r"start\s*=\s*(\d+)", opt)
            if ms:
                start = int(ms.group(1))
        item_kw = "task" if name == "tasks" else "item"
        cm = mask_tex(content)
        idxs = [m for m in re.finditer(r"\\" + item_kw + r"(?![A-Za-z@])", cm) if _brace_depth(cm, m.start()) == 0]
        pieces = []
        if idxs and content[: idxs[0].start()].strip():
            pieces.append(content[: idxs[0].start()].strip())
        for k, m in enumerate(idxs):
            seg_end = idxs[k + 1].start() if k + 1 < len(idxs) else len(content)
            seg = content[m.end(): seg_end]
            custom, p3 = _read_opt(seg, 0)
            body = seg[p3:].strip() if custom is not None else seg.strip()
            n = start + k
            if custom is not None:
                label = custom
            elif name.startswith("enumerate") or name == "tasks":
                label = _enum_label(tmpl if name.startswith("enumerate") else (tmpl or "\\alph*)"), n, depth)
            elif name.startswith("itemize"):
                label = tmpl or {1: "•", 2: "–", 3: "∗"}.get(depth, "•")
            else:
                label = "\\textbf{" + (custom or "") + "}"
            if name == "description" and custom is not None:
                label = "\\textbf{" + custom + "}"
            pieces.append(f"{label} {body}")
            stats["items"] += 1
        key = "enumerate" if name == "tasks" else name.rstrip("*")
        stats[key] = stats.get(key, 0) + 1
        repl = "\n\n" + "\n\n".join(pieces) + "\n\n"
        text = text[:s] + repl + text[e:]
    return text, stats


def _all_list_spans(masked: str, names, a: int, b: int) -> list[tuple[int, int, str]]:
    out = []
    stack = list(_env_spans(masked, names, a, b))
    while stack:
        s, e, n = stack.pop()
        out.append((s, e, n))
        stack.extend(_env_spans(masked, names, s + 1, e - 1))
    return out


def _brace_depth(s: str, pos: int) -> int:
    d = 0
    i = 0
    while i < pos:
        c = s[i]
        if c == "\\":
            i += 2
            continue
        if c == "{":
            d += 1
        elif c == "}":
            d -= 1
        i += 1
    return d


# ───────────────────────────── 4. môi trường & lệnh lạ ─────────────────────────────
#: môi trường engine tự xử lý hoặc cần giữ nguyên cho bước sau
_KEEP_ENVS = {"document", "center", "figure", "figure*", "table", "table*", "tikzpicture", "axis", "scope",
              "tabular", "tabular*", "tabularx", "longtable", "array", "cases", "dcases", "matrix", "pmatrix",
              "bmatrix", "Bmatrix", "vmatrix", "Vmatrix", "smallmatrix", "aligned", "alignedat", "gathered", "split",
              "verbatim", "verbatim*", "Verbatim", "lstlisting", "minted", "comment", *_MATH_ENV_NAMES,
              "itemize", "enumerate", "description", "tasks", "filecontents", "filecontents*", "pgfonlayer",
              "semilogxaxis", "semilogyaxis", "loglogaxis", "polaraxis", "groupplot", "tikzcd", "pspicture",
              "circuitikz", "picture"}
_ENV_ARITY = {"minipage": (1, True), "multicols": (1, False), "multicols*": (1, False), "wrapfigure": (2, True),
              "wraptable": (2, True), "adjustbox": (1, False), "tcolorbox": (0, True), "mdframed": (0, True),
              "framed": (0, False), "spacing": (1, False), "adjustwidth": (2, False), "thebibliography": (1, False),
              "flushleft": (0, False), "flushright": (0, False), "quote": (0, False), "quotation": (0, False),
              "abstract": (0, False), "small": (0, False), "footnotesize": (0, False), "samepage": (0, False)}
_TITLED_ENVS = {"ex": "Câu", "bt": "Bài", "baitap": "Bài", "vd": "Ví dụ", "vidu": "Ví dụ", "dl": "Định lý",
                "dn": "Định nghĩa", "nx": "Nhận xét", "chuy": "Chú ý", "tinhchat": "Tính chất", "hequa": "Hệ quả",
                "bode": "Bổ đề", "theorem": "Theorem", "lemma": "Lemma", "corollary": "Corollary",
                "proposition": "Proposition", "definition": "Definition", "example": "Example", "remark": "Remark"}
_UNNUMBERED_TITLES = {"proof": "Chứng minh.", "loigiai": "Lời giải.", "huongdan": "Hướng dẫn.", "dapan": "Đáp án."}

#: lệnh engine (core_converter/omml_math) đã hiểu trong văn bản — không bóc
_ENGINE_TEXT_CMDS = {"textbf", "textit", "underline", "emph", "texttt", "section", "subsection", "subsubsection",
                     "paragraph", "caption", "includegraphics", "immini", "choice", "choiceA", "choiceB", "choiceC",
                     "choiceD", "choiceTF", "boncot", "haicot", "motcot", "maketitle", "title", "author", "date",
                     "newline", "par", "noindent", "centering", "hspace", "hfill", "vfill", "dots", "ldots",
                     "cdots", "item", "begin", "end", "Huge", "huge", "LARGE", "Large", "large", "normalsize", "small",
                     "footnotesize", "scriptsize", "tiny", "bfseries", "itshape", "bf", "it", "em", "textbackslash",
                     "hline", "cline", "multicolumn", "multirow", "diagbox", "toprule", "midrule", "bottomrule",
                     "textwidth", "linewidth", "tkzTabInit", "tkzTabLine", "tkzTabVar",
                     "heva", "hoac", "today", "clearpage", "newpage", "pagebreak", "linebreak",
                     "True", "smallskip", "medskip", "bigskip", "indent", "raggedright", "raggedleft",
                     "url", "href", "LaTeX", "TeX"}
#: lệnh chỉ trình bày (khoảng cách/màu nền ô/kiểu trang) → bỏ cả tên và đối số (không có chữ hiển thị)
_DROP_CMDS = {"vspace", "setstretch", "linespread", "rowcolor", "cellcolor", "columncolor", "arrayrulecolor",
              "addvspace", "enlargethispage", "thispagestyle", "pagestyle", "pagenumbering", "setlength",
              "addtolength", "setcounter", "addtocounter", "vskip", "hskip"}
#: lệnh không đối số → chữ hiển thị tương đương
_ZERO_ARG_TEXT = {"dotfill": " ……………… ", "hrulefill": " ________________ ", "ldots": "…", "dots": "…",
                  "quad": "\u2003", "qquad": "\u2003\u2003", "enspace": "\u2002", "thinspace": "\u2009",
                  "textbackslash": "\\textbackslash ", "LaTeX": "LaTeX", "TeX": "TeX", "today": ""}


#: lệnh định dạng giữ nội dung: tên → (số đối số bắt buộc, chỉ số đối số là nội dung, khuôn thay "@")
_CONTENT_CMDS: dict[str, tuple[int, int, str]] = {
    "textcolor": (2, 1, "@"), "colorbox": (2, 1, "@"), "fcolorbox": (3, 2, "@"), "textsc": (1, 0, "@"),
    "textsf": (1, 0, "@"), "textrm": (1, 0, "@"), "textsl": (1, 0, "@"), "textup": (1, 0, "@"),
    "textmd": (1, 0, "@"), "textnormal": (1, 0, "@"), "uline": (1, 0, "\\underline{@}"),
    "uuline": (1, 0, "\\underline{@}"), "uwave": (1, 0, "\\underline{@}"), "sout": (1, 0, "@"),
    "xout": (1, 0, "@"), "hl": (1, 0, "@"), "mbox": (1, 0, "@"), "fbox": (1, 0, "@"), "framebox": (1, 0, "@"),
    "makebox": (1, 0, "@"), "parbox": (2, 1, "@"), "raisebox": (2, 1, "@"), "ensuremath": (1, 0, "$@$"),
    "textsuperscript": (1, 0, "$^{\\text{@}}$"), "textsubscript": (1, 0, "$_{\\text{@}}$"),
    "footnote": (1, 0, " (@)"), "scalebox": (2, 1, "@"), "resizebox": (3, 2, "@"), "rotatebox": (2, 1, "@"),
    "ce": (1, 0, "$\\ce{@}$"), "num": (1, 0, "$@$"),
    "MakeUppercase": (1, 0, "@"), "MakeLowercase": (1, 0, "@"), "emph": (1, 0, "\\textit{@}"),
    "loigiai": (1, 0, "\n\n\\textbf{Lời giải.} @\n\n"), "label": (1, 0, ""),
    "ref": (1, 0, "[@]"), "eqref": (1, 0, "([@])"), "pageref": (1, 0, "[@]"),
}
#: lệnh sổ đăng ký ghi NATIVE (với tiền xử lý thông minh) nhưng engine thuần không hiểu → vẫn bóc vỏ
_STRIP_EVEN_IF_REGISTERED = {"textsc", "textsf", "textrm", "textsl", "colorbox", "fcolorbox", "textcolor", "hl",
                             "cite", "pageref", "footnotetext"}


def _exam_conventions(body: str, user_defined: set[str]) -> tuple[str, dict[str, int]]:
    """Quy ước đề thi (ex_test) khi người dùng KHÔNG tự định nghĩa: \\cau → "Câu n.", \\bonpa → \\boncot (4 cột),
    \\choiceTF 4 đối số (đánh dấu \\True) → dạng 8 đối số của engine; đáp án chỉ lấy từ \\True."""
    st = {"cau": 0, "bonpa": 0, "choiceTF4": 0}
    bm = mask_tex(body)
    repls: list[tuple[int, int, str]] = []
    counter = 0
    if "cau" not in user_defined and "Cau" not in user_defined:
        for m in re.finditer(r"\\[cC]au(?![A-Za-z@])", bm):
            arg, pos = _read_arg(body, m.end())
            if arg is not None and arg.strip().isdigit():
                n = arg.strip()
                counter = int(n)
            else:
                counter += 1
                n = str(counter)
                pos = m.end()
            repls.append((m.start(), pos, f"\n\n\\textbf{{Câu {n}.}} "))
            st["cau"] += 1
    if "bonpa" not in user_defined:
        for m in re.finditer(r"\\bonpa(?![A-Za-z@])", bm):
            repls.append((m.start(), m.end(), "\\boncot"))
            st["bonpa"] += 1
    if "choiceTF" not in user_defined:
        for m in re.finditer(r"\\choiceTF(?![A-Za-z@])", bm):
            pos = m.end()
            args = []
            spans = []
            for _ in range(8):
                k = pos
                while k < len(body) and body[k] in " \t\n":
                    k += 1
                if k < len(body) and body[k] == "{":
                    e = find_matching_brace(body, k)
                    if e < 0:
                        break
                    args.append(body[k + 1: e])
                    spans.append((k, e))
                    pos = e + 1
                else:
                    break
            if len(args) == 4:
                parts = []
                for a in args:
                    has_true = re.search(r"\\True(?![A-Za-z@])", a) is not None
                    stmt = re.sub(r"\\True(?![A-Za-z@])", "", a).strip()
                    any_true = any(re.search(r"\\True(?![A-Za-z@])", x) for x in args)
                    ans = ("Đ" if has_true else "S") if any_true else ""
                    parts.append("{" + stmt + "}{" + ans + "}")
                repls.append((m.start(), pos, "\\choiceTF" + "".join(parts)))
                st["choiceTF4"] += 1
    return _apply(body, repls), st


_SI_UNIT_MACROS = {"metre": "m", "meter": "m", "second": "s", "kilogram": "kg", "gram": "g", "newton": "N",
                   "joule": "J", "watt": "W", "volt": "V", "ampere": "A", "ohm": "Ω", "kelvin": "K",
                   "celsius": "°C", "degreeCelsius": "°C", "mole": "mol", "litre": "L", "liter": "L", "hertz": "Hz",
                   "pascal": "Pa", "coulomb": "C", "farad": "F", "tesla": "T", "henry": "H", "hour": "h",
                   "minute": "min", "percent": "%", "degree": "°", "centi": "c", "milli": "m", "kilo": "k",
                   "mega": "M", "micro": "μ", "nano": "n", "giga": "G"}


def _si_unit_latex(unit: str) -> str:
    r"""Đơn vị siunitx (dạng chữ "m/s^2" hoặc macro \metre\per\second\squared) → LaTeX toán chữ đứng."""
    u = unit.strip()
    if "\\" in u:
        toks = re.findall(r"\\([A-Za-z]+)|([^\\\s]+)", u)
        parts: list[str] = []
        pending_per = False
        for cmd, lit in toks:
            if cmd == "per":
                pending_per = True
                continue
            if cmd in ("squared", "cubed"):
                if parts:
                    parts[-1] += "^{2}" if cmd == "squared" else "^{3}"
                continue
            if cmd in ("square", "cubic"):
                parts.append("")
                continue
            sym = _SI_UNIT_MACROS.get(cmd, cmd) if cmd else lit
            if parts and parts[-1] in ("c", "m", "k", "M", "μ", "n", "G") and cmd in _SI_UNIT_MACROS \
                    and cmd not in ("centi", "milli", "kilo", "mega", "micro", "nano", "giga") and not pending_per:
                parts[-1] += sym
                continue
            parts.append(("/" if pending_per else "") + sym)
            pending_per = False
        u = "\\,".join(p for p in parts if p)
        u = u.replace("\\,/", "/")
    out = re.sub(r"([A-Za-zΩμ°%]+)", lambda mm: "\\mathrm{" + mm.group(1) + "}", u)
    out = out.replace(".", "\\cdot ").replace("~", "\\,")
    return out


def _table_captions(body: str) -> tuple[str, int]:
    r"""\caption trong môi trường table/table* → đoạn "Bảng n. …" (đánh số theo thứ tự xuất hiện; \caption*
    không đánh số). Chú thích hình để engine xử lý cùng khối hình."""
    bm = mask_tex(body)
    repls: list[tuple[int, int, str]] = []
    n = 0
    for s0, e0, _name in _env_spans(bm, r"table\*?", 0, len(bm)):
        for m in re.finditer(r"\\caption(\*?)(?![A-Za-z@])", bm[s0:e0]):
            st = s0 + m.start()
            pos = s0 + m.end()
            _o, p2 = _read_opt(body, pos)
            if _o is not None:
                pos = p2
            arg, pos = _read_arg(body, pos)
            if arg is None:
                continue
            if m.group(1):
                head = ""
            else:
                n += 1
                head = f"\\textbf{{Bảng {n}.}} "
            repls.append((st, pos, "\n\n\\begin{center}" + head + arg + "\\end{center}\n\n"))
    return _apply(body, repls), n


_NEWIF_RE = re.compile(r"\\newif\s*\\if([A-Za-z@]+)")


def resolve_newif(text: str) -> tuple[str, dict[str, Any]]:
    r"""Điều kiện \newif\ifX của người dùng: giá trị lấy từ \Xtrue/\Xfalse trong PREAMBLE (mặc định false như TeX);
    trong thân, khối \ifX … [\else …] \fi được thay bằng nhánh đúng. Bật/tắt lại trong thân → không đoán, giữ nguyên
    và báo cáo."""
    masked = mask_tex(text)
    a, b = _body_span(text, masked)
    pre_m = masked[:a]
    names = [m.group(1) for m in _NEWIF_RE.finditer(pre_m) if "@" not in m.group(1)]
    stats: dict[str, Any] = {"resolved": {}, "skipped": {}}
    if not names:
        return text, stats
    values: dict[str, bool] = {}
    for nm in names:
        val = False
        for m in re.finditer(r"\\" + re.escape(nm) + r"(true|false)(?![A-Za-z@])", pre_m):
            val = m.group(1) == "true"
        values[nm] = val
    body = text[a:b]
    bm = masked[a:b]
    for nm in names:
        if re.search(r"\\" + re.escape(nm) + r"(?:true|false)(?![A-Za-z@])", bm):
            stats["skipped"][nm] = "giá trị đổi trong thân tài liệu"
            del values[nm]
    if not values:
        return text, stats
    tok_re = re.compile(r"\\(if[A-Za-z@]*|else|fi|or)(?![A-Za-z@])")
    for _ in range(50):
        bm = mask_tex(body)
        hit = None
        for m in re.finditer(r"\\if(" + "|".join(re.escape(n) for n in values) + r")(?![A-Za-z@])", bm):
            hit = m
            break
        if hit is None:
            break
        depth = 0
        else_pos = None
        end = None
        for t in tok_re.finditer(bm, hit.end()):
            w = t.group(1)
            if w.startswith("if"):
                depth += 1
            elif w == "fi":
                if depth == 0:
                    end = t
                    break
                depth -= 1
            elif w == "else" and depth == 0 and else_pos is None:
                else_pos = t
        if end is None:
            stats["skipped"][hit.group(1)] = "thiếu \\fi"
            values.pop(hit.group(1), None)
            if not values:
                break
            continue
        nm = hit.group(1)
        if else_pos is not None:
            yes, no = body[hit.end(): else_pos.start()], body[else_pos.end(): end.start()]
        else:
            yes, no = body[hit.end(): end.start()], ""
        keep = yes if values[nm] else no
        body = body[:hit.start()] + keep + body[end.end():]
        stats["resolved"][nm] = stats["resolved"].get(nm, 0) + 1
    return text[:a] + body + text[b:], stats


def normalize_structure(text: str, *, show_solutions: bool = True) -> tuple[str, dict[str, Any]]:
    masked = mask_tex(text)
    a, b = _body_span(text, masked)
    pre = text[:a]
    theorem_titles: dict[str, str] = {}
    for m in re.finditer(r"\\newtheorem\*?\s*\{([^}]+)\}\s*(?:\[[^\]]*\])?\s*\{([^}]+)\}", mask_tex(pre)):
        theorem_titles[m.group(1).strip()] = m.group(2).strip()
    newenvs: dict[str, tuple[int, str, str]] = {}
    for m in re.finditer(r"\\(?:re)?newenvironment\s*\{([^}]+)\}", mask_tex(pre)):
        name = m.group(1).strip()
        pos = m.end()
        o, p2 = _read_opt(pre, pos)
        n = int(o) if o and o.strip().isdigit() else 0
        if o is not None:
            pos = p2
            o2, p3 = _read_opt(pre, pos)
            if o2 is not None:
                pos = p3
                n = -1  # có đối số tuỳ chọn → không tự nở
        g1, pos = _read_arg(pre, pos)
        g2, pos = _read_arg(pre, pos)
        if g1 is None or g2 is None or n < 0 or _UNSAFE_BODY.search(g1 + g2):
            continue
        newenvs[name] = (n, g1, g2)
    stats: dict[str, Any] = {"unwrapped": {}, "titled": {}, "user_envs_expanded": {}, "unknown_commands": {},
                             "kept_args_of_unknown": 0}
    counters: dict[str, int] = {}
    body = text[a:b]
    bm = mask_tex(body)
    math = _math_spans(bm, 0, len(bm))
    fig = [(s, e) for s, e, _n in _env_spans(bm, r"tikzpicture|axis|pspicture|circuitikz|tikzcd", 0, len(bm))]
    tabs = [(s, e) for s, e, _n in _env_spans(bm, r"tabular\*?|tabularx|longtable", 0, len(bm))]
    repls: list[tuple[int, int, str]] = []
    for m in re.finditer(r"\\(begin|end)\s*\{([A-Za-z@*]+)\}", bm):
        kind, name = m.group(1), m.group(2)
        if name in _KEEP_ENVS or _inside(math, m.start()) or _inside(fig, m.start()):
            continue
        pos = m.end()
        if kind == "begin":
            if name in newenvs:
                n, beg, _end = newenvs[name]
                args = []
                for _ in range(n):
                    arg, pos = _read_arg(body, pos)
                    args.append(arg or "")
                rep = beg
                for i in range(n, 0, -1):
                    rep = rep.replace(f"#{i}", args[i - 1])
                repls.append((m.start(), pos, " " + rep + " "))
                stats["user_envs_expanded"][name] = stats["user_envs_expanded"].get(name, 0) + 1
                continue
            opt = None
            arity, has_opt_first = _ENV_ARITY.get(name, (0, True))
            o, p2 = _read_opt(body, pos)
            if o is not None:
                opt = o
                pos = p2
            for _ in range(arity):
                _arg, pos = _read_arg(body, pos)
            title = theorem_titles.get(name) or _TITLED_ENVS.get(name)
            if title:
                counters[name] = counters.get(name, 0) + 1
                head = f"\\textbf{{{title} {counters[name]}.}}" + (f" ({opt})" if opt else "") + " "
                repls.append((m.start(), pos, "\n\n" + head))
                stats["titled"][name] = stats["titled"].get(name, 0) + 1
            elif name in _UNNUMBERED_TITLES:
                if name in ("loigiai", "huongdan", "dapan") and not show_solutions:
                    continue
                repls.append((m.start(), pos, "\n\n\\textbf{" + _UNNUMBERED_TITLES[name] + "} "))
                stats["titled"][name] = stats["titled"].get(name, 0) + 1
            else:
                t_opt = ""
                if name == "tcolorbox" and opt:
                    mt = re.search(r"title\s*=\s*(\{[^{}]*\}|[^,\]]+)", opt)
                    if mt:
                        t_opt = "\\textbf{" + mt.group(1).strip("{}") + "}\n\n"
                repls.append((m.start(), pos, "\n\n" + t_opt))
                stats["unwrapped"][name] = stats["unwrapped"].get(name, 0) + 1
        else:
            if name in newenvs:
                repls.append((m.start(), m.end(), " " + newenvs[name][2] + " "))
                continue
            repls.append((m.start(), m.end(), "\n\n"))
    body = _apply(body, repls)
    body, cap_stats = _table_captions(body)
    if cap_stats:
        stats["table_captions"] = cap_stats
    body, exam_stats = _exam_conventions(body, user_defined=set(_parse_macro_defs(pre)[0]))
    stats["exam_conventions"] = exam_stats
    # lệnh lạ trong văn bản: bỏ tên, giữ đối số
    try:
        from .structure_registry import tier_of_command
    except Exception:  # pragma: no cover
        tier_of_command = lambda n: None  # noqa: E731
    bm = mask_tex(body)
    math = _math_spans(bm, 0, len(bm))
    fig = [(s, e) for s, e, _n in _env_spans(bm, r"tikzpicture|axis|pspicture|circuitikz|tikzcd", 0, len(bm))]
    tabspec = [(m.start(), m.end()) for m in re.finditer(r"\\begin\{(?:tabular\*?|tabularx|longtable)\}(?:\{[^{}]*\})?\{(?:[^{}]|\{[^{}]*\})*\}", bm)]
    repls = []
    for m in re.finditer(r"\\([A-Za-z]+)\*?", bm):
        name = m.group(1)
        if (_inside(math, m.start()) or _inside(fig, m.start()) or _inside(tabspec, m.start())
                or name in _ENGINE_TEXT_CMDS):
            continue
        if repls and m.start() < repls[-1][1]:
            continue
        if name in _ZERO_ARG_TEXT:
            repls.append((m.start(), m.end(), _ZERO_ARG_TEXT[name]))
            continue
        if name in _DROP_CMDS:
            pos = m.end()
            o, p2 = _read_opt(body, pos)
            if o is not None:
                pos = p2
            n_args = 0 if name in ("vskip", "hskip") else (2 if name in ("setlength", "addtolength", "setcounter",
                                                                         "addtocounter") else 1)
            for _ in range(n_args):
                arg, p2 = _read_arg(body, pos)
                if arg is None:
                    break
                pos = p2
            repls.append((m.start(), pos, " "))
            stats["dropped_presentation"] = stats.get("dropped_presentation", {})
            stats["dropped_presentation"][name] = stats["dropped_presentation"].get(name, 0) + 1
            continue
        if name in ("SI", "si", "qty", "unit"):
            pos = m.end()
            o, p2 = _read_opt(body, pos)
            if o is not None:
                pos = p2
            args_si: list[str] = []
            for _ in range(2 if name in ("SI", "qty") else 1):
                arg, pos2 = _read_arg(body, pos)
                if arg is None:
                    break
                args_si.append(arg)
                pos = pos2
            if len(args_si) == (2 if name in ("SI", "qty") else 1):
                unit = _si_unit_latex(args_si[-1])
                rep = f"${args_si[0]}\\,{unit}$" if len(args_si) == 2 else f"${unit}$"
                repls.append((m.start(), pos, rep))
                stats["content_commands"] = stats.get("content_commands", {})
                stats["content_commands"][name] = stats["content_commands"].get(name, 0) + 1
            continue
        if name in _CONTENT_CMDS:
            n_args, content_idx, wrap = _CONTENT_CMDS[name]
            pos = m.end()
            args: list[str] = []
            ok = True
            for _ in range(n_args):
                o, p2 = _read_opt(body, pos)
                if o is not None:
                    pos = p2
                arg, pos = _read_arg(body, pos)
                if arg is None:
                    ok = False
                    break
                args.append(arg)
            if not ok:
                continue
            content = args[content_idx]
            repls.append((m.start(), pos, wrap.replace("@", content)))
            stats["content_commands"] = stats.get("content_commands", {})
            stats["content_commands"][name] = stats["content_commands"].get(name, 0) + 1
            continue
        tier = tier_of_command(name)
        if tier in ("NATIVE", "PICTURE") and name not in _STRIP_EVEN_IF_REGISTERED:
            continue
        pos = m.end()
        args: list[str] = []
        while True:
            o, p2 = _read_opt(body, pos)
            if o is None:
                break
            pos = p2
        while True:
            k = pos
            while k < len(body) and body[k] in " \t":
                k += 1
            if k < len(body) and body[k] == "{":
                e = find_matching_brace(body, k)
                if e < 0:
                    break
                args.append(body[k + 1: e])
                pos = e + 1
            else:
                break
        if tier == "IGNORE":
            repls.append((m.start(), pos, " "))
            continue
        stats["unknown_commands"][name] = stats["unknown_commands"].get(name, 0) + 1
        stats["kept_args_of_unknown"] += len(args)
        repls.append((m.start(), pos, " " + " ".join(args) + " " if args else " "))
    body = _apply(body, repls)
    return text[:a] + body + text[b:], stats


# ───────────────────────────── 5. hình & ảnh ─────────────────────────────
_HARD_TIKZ_RE = re.compile(
    r"\\(?:shade|shadedraw|pattern|tdplot\w*|foreach|pgfmathsetmacro|pgfmathparse|tkz(?!Tab)\w+|addplot3|"
    r"includegraphics|pic|path|clip|filldraw|draw\s*plot|spy|graph|matrix)\b|fill\s*between|name path|"
    r"intersections|canvas is|decorat|opacity|\(\$|\bpattern\b|3d\b|\bsmooth\b|\bpos\s*=|\bshading\b")


def _label_count_tikz(tex: str) -> int:
    n = 0
    for m in re.finditer(r"\bnode\b[^{;]*\{", tex):
        ob = m.end() - 1
        e = find_matching_brace(tex, ob)
        if e > 0 and tex[ob + 1: e].strip():
            n += 1
    return n


def _sniff_image(data: bytes) -> str | None:
    if data.startswith(b"\x89PNG"):
        return "png"
    if data.startswith(b"\xff\xd8\xff"):
        return "jpg"
    if data.startswith(b"GIF8"):
        return "gif"
    if data.startswith(b"BM"):
        return "bmp"
    if data[:4] in (b"II*\x00", b"MM\x00*"):
        return "tiff"
    if b"%PDF-" in data[:1024]:
        return "pdf"
    if data.startswith(b"%!PS") or data.startswith(b"\xc5\xd0\xd3\xc6"):
        return "eps"
    head = data[:2048].lstrip().lower()
    if head.startswith(b"<?xml") or head.startswith(b"<svg") or b"<svg" in head:
        return "svg"
    return None


def _ci_find(p: Path) -> Path | None:
    if p.is_file():
        return p
    if p.parent.is_dir():
        low = p.name.casefold()
        for c in p.parent.iterdir():
            if c.name.casefold() == low and c.is_file():
                return c
    return None


def resolve_graphic(name: str, root: Path, bases: list[str], graphicspath: list[str]) -> tuple[Path | None, str]:
    name = name.strip().strip('"')
    if not name or name.startswith("/") or re.match(r"^[A-Za-z]:", name):
        return None, "đường dẫn tuyệt đối không được phép"
    rr = root.resolve()
    dirs = list(dict.fromkeys(bases + [g for g in graphicspath] + [""]))
    for d in dirs:
        base = (rr / d).resolve()
        cand0 = (base / name).resolve()
        if rr not in cand0.parents and cand0 != rr:
            return None, "đường dẫn ra ngoài gói"
        tries = [cand0] if Path(name).suffix else []
        tries += [cand0.with_name(cand0.name + e) for e in IMAGE_EXTS]
        if Path(name).suffix and Path(name).suffix.lower() not in [e.lower() for e in IMAGE_EXTS]:
            tries += [cand0.with_name(cand0.name + e) for e in IMAGE_EXTS]
        for c in tries:
            f = _ci_find(c)
            if f is not None:
                how = "ok" if f == c else "ok_case_insensitive"
                return f, how
    return None, "không tìm thấy"


def _pdf_page_to_png(pdf_bytes: bytes, page: int, dpi: int) -> tuple[bytes, float, float]:
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(pdf_bytes)
    try:
        pg = doc[max(0, min(page, len(doc) - 1))]
        w, h = pg.get_size()
        img = pg.render(scale=dpi / 72.0).to_pil()
        buf = io.BytesIO()
        img.save(buf, "PNG")
        return buf.getvalue(), float(w), float(h)
    finally:
        doc.close()


def _pdf_page_to_svg(pdf_path: Path, page: int, out_dir: Path) -> bytes | None:
    tool = shutil.which("pdftocairo")
    if not tool:
        return None
    out = out_dir / f"p{page + 1}_{int(time.time() * 1e6)}.svg"
    try:
        subprocess.run([tool, "-svg", "-f", str(page + 1), "-l", str(page + 1), str(pdf_path), str(out)],
                       capture_output=True, timeout=120)
    except subprocess.TimeoutExpired:
        return None
    return out.read_bytes() if out.exists() and out.stat().st_size > 0 else None


def pdf_with_cropbox_as_mediabox(pdf_path: Path, page: int, out_path: Path) -> Path:
    """Một trang, MediaBox := CropBox (pdftocairo -svg vẽ theo MediaBox; TeX/graphicx dùng CropBox)."""
    import pikepdf
    with pikepdf.open(str(pdf_path)) as src:
        dst = pikepdf.new()
        dst.pages.append(src.pages[page])
        pg = dst.pages[0]
        try:
            cb = pg.cropbox
            pg.mediabox = cb
        except Exception:
            pass
        dst.save(str(out_path))
    return out_path


def _svg_size_pt(data: bytes) -> tuple[float, float] | None:
    txt = data[:4000].decode("utf-8", "ignore")
    m = re.search(r"<svg[^>]*>", txt, re.S)
    if not m:
        return None
    tag = m.group(0)

    def val(attr: str) -> float | None:
        mm = re.search(attr + r'\s*=\s*"([0-9.]+)\s*(pt|px|mm|cm|in)?"', tag)
        if not mm:
            return None
        v = float(mm.group(1))
        unit = mm.group(2) or "px"
        return v * {"pt": 1, "px": 0.75, "mm": 72 / 25.4, "cm": 72 / 2.54, "in": 72}[unit]

    w, h = val("width"), val("height")
    if w and h:
        return w, h
    vb = re.search(r'viewBox\s*=\s*"[-0-9.]+\s+[-0-9.]+\s+([0-9.]+)\s+([0-9.]+)"', tag)
    if vb:
        return float(vb.group(1)) * 0.75, float(vb.group(2)) * 0.75
    return None


def _svg_png_fallback(svg_path: Path, out_dir: Path, width_px: int) -> bytes | None:
    from .capabilities import find_tool
    lo = find_tool("soffice", "libreoffice")
    if not lo:
        return None
    out_dir.mkdir(parents=True, exist_ok=True)
    flt = 'png:draw_png_Export:{"PixelWidth":{"type":"long","value":"%d"}}' % max(64, width_px)
    try:
        subprocess.run([lo, "--headless", "-env:UserInstallation=" + (out_dir / "lo_profile").resolve().as_uri(),
                        "--convert-to", flt, "--outdir", str(out_dir), str(svg_path)],
                       capture_output=True, timeout=120)
    except subprocess.TimeoutExpired:
        return None
    p = out_dir / (svg_path.stem + ".png")
    return p.read_bytes() if p.exists() else None


def convert_graphic(path: Path, opts: str, dpi: int, work: Path) -> dict[str, Any]:
    """Chuyển tệp ảnh nguồn thành dữ liệu nhúng Word. Nhận dạng theo NỘI DUNG, không theo đuôi."""
    data = path.read_bytes()
    kind = _sniff_image(data)
    page = 0
    mp = re.search(r"page\s*=\s*(\d+)", opts or "")
    if mp:
        page = int(mp.group(1)) - 1
    res: dict[str, Any] = {"source": str(path), "detected": kind, "warnings": []}
    if kind in ("png", "jpg"):
        from PIL import Image
        im = Image.open(io.BytesIO(data))
        dpi_x = (im.info.get("dpi") or (96, 96))[0] or 96
        res.update({"data": data, "ext": kind, "png_data": b"", "width_pt": im.width * 72.0 / dpi_x,
                    "height_pt": im.height * 72.0 / dpi_x, "method": "raster_image"})
        return res
    if kind in ("gif", "bmp", "tiff"):
        from PIL import Image
        im = Image.open(io.BytesIO(data))
        buf = io.BytesIO()
        im.save(buf, "PNG")
        dpi_x = (im.info.get("dpi") or (96, 96))[0] or 96
        res.update({"data": buf.getvalue(), "ext": "png", "png_data": b"", "width_pt": im.width * 72.0 / dpi_x,
                    "height_pt": im.height * 72.0 / dpi_x, "method": "raster_image"})
        res["warnings"].append(f"đã chuyển {kind.upper()} → PNG")
        return res
    if kind == "pdf":
        tmp = work / f"g_{path.stem[:40]}_{int(time.time() * 1e6)}.pdf"
        try:
            pdf_with_cropbox_as_mediabox(path, page, tmp)
            src_pdf = tmp
            pg_index = 0
        except Exception as exc:
            res["warnings"].append(f"không chuẩn hoá CropBox ({type(exc).__name__}); dùng trang gốc")
            src_pdf, pg_index = path, page
        png, w, h = _pdf_page_to_png(src_pdf.read_bytes(), pg_index, dpi)
        svg = _pdf_page_to_svg(src_pdf, pg_index, work)
        if svg:
            res.update({"data": svg, "ext": "svg", "png_data": png, "width_pt": w, "height_pt": h,
                        "method": "vector_image"})
        else:
            res.update({"data": png, "ext": "png", "png_data": b"", "width_pt": w, "height_pt": h,
                        "method": "raster_image"})
            res["warnings"].append("thiếu pdftocairo → hình PDF chỉ nhúng PNG (không vector)")
        return res
    if kind == "svg":
        size = _svg_size_pt(data) or (200.0, 150.0)
        png = _svg_png_fallback(path, work / "svgpng", int(size[0] / 72 * dpi))
        if not png:
            res["warnings"].append("không tạo được PNG dự phòng cho SVG (thiếu LibreOffice) → Word cũ có thể không hiện")
        res.update({"data": data, "ext": "svg", "png_data": png or b"", "width_pt": size[0], "height_pt": size[1],
                    "method": "vector_image"})
        return res
    if kind == "eps":
        gs = shutil.which("gs") or shutil.which("gswin64c") or shutil.which("gswin32c")
        if not gs:
            res["error"] = "ảnh EPS cần Ghostscript (gs) để chuyển đổi — máy chưa có"
            return res
        out_pdf = work / f"eps_{path.stem[:40]}.pdf"
        try:
            subprocess.run([gs, "-dSAFER", "-dBATCH", "-dNOPAUSE", "-dEPSCrop", "-sDEVICE=pdfwrite",
                            f"-sOutputFile={out_pdf}", str(path)], capture_output=True, timeout=120)
        except subprocess.TimeoutExpired:
            res["error"] = "Ghostscript quá thời gian"
            return res
        if not out_pdf.exists():
            res["error"] = "Ghostscript không chuyển được EPS"
            return res
        sub = convert_graphic(out_pdf, "", dpi, work)
        sub["warnings"] = res["warnings"] + ["EPS → PDF bằng Ghostscript"] + sub.get("warnings", [])
        return sub
    res["error"] = "không nhận dạng được định dạng ảnh theo nội dung"
    return res


_LEN_RE = re.compile(r"([0-9.]+)\s*(\\textwidth|\\linewidth|\\columnwidth|\\textheight|cm|mm|in|pt|em|ex)?")


def graphic_size(opts: str, natural_w: float, natural_h: float, text_w_pt: float) -> tuple[float, float, list[str]]:
    warns: list[str] = []
    w = h = None
    scale = 1.0

    def length(v: str) -> float | None:
        m = _LEN_RE.match(v.strip())
        if not m:
            return None
        x = float(m.group(1))
        u = m.group(2) or "pt"
        return x * {"\\textwidth": text_w_pt, "\\linewidth": text_w_pt, "\\columnwidth": text_w_pt,
                    "\\textheight": text_w_pt * 1.4, "cm": 72 / 2.54, "mm": 72 / 25.4, "in": 72.0, "pt": 1.0,
                    "em": 12.0, "ex": 6.0}[u]

    for kv in [p.strip() for p in (opts or "").split(",") if p.strip()]:
        if "=" not in kv:
            if kv not in ("keepaspectratio", "clip"):
                warns.append(f"tuỳ chọn ảnh chưa hỗ trợ: {kv}")
            continue
        k, v = [s.strip() for s in kv.split("=", 1)]
        if k == "width":
            w = length(v)
        elif k == "height":
            h = length(v)
        elif k == "scale":
            try:
                scale = float(v)
            except ValueError:
                warns.append(f"scale không đọc được: {v}")
        elif k in ("angle", "trim", "viewport", "origin"):
            warns.append(f"tuỳ chọn ảnh '{k}' chưa áp dụng (ảnh giữ nguyên hướng/khung)")
        elif k == "page":
            pass
    if w and h and "keepaspectratio" in (opts or ""):
        r = min(w / natural_w, h / natural_h)
        return natural_w * r, natural_h * r, warns
    if w and not h:
        return w, natural_h * w / natural_w, warns
    if h and not w:
        return natural_w * h / natural_h, h, warns
    if w and h:
        return w, h, warns
    return natural_w * scale, natural_h * scale, warns


# ───────────────────────────── snapshot TeX ─────────────────────────────
def _snapshot_blocks(blocks: list[EngineBlock], preamble: str, main_dir: Path, work: Path, *, dpi: int,
                     engine_pref: str, timeout: int = 240) -> None:
    from .tex_compile import choose_engine, run_tex
    eng = choose_engine(preamble, engine_pref)
    if not eng:
        for b in blocks:
            b.status, b.method, b.reason = "unresolved", "unresolved", "máy không có TeX để chụp ảnh khối"
        return
    work.mkdir(parents=True, exist_ok=True)
    pre = re.sub(r"\\begin\s*\{document\}.*", "", preamble, flags=re.S)
    if "\\documentclass" not in pre:
        pre = "\\documentclass{article}\n\\usepackage{amsmath,amssymb,tikz,pgfplots,tkz-tab,graphicx}\n" + pre
    head = pre + "\n\\usepackage[active,tightpage]{preview}\n\\setlength\\PreviewBorder{2pt}\n\\begin{document}\n"

    def compile_group(group: list[EngineBlock], tag: str) -> dict | None:
        body = "\n".join("\\begin{preview}\n" + b.tex + "\n\\end{preview}\n" for b in group)
        name = f"_l2w_snap_{tag}.tex"
        tex_path = main_dir / name
        tex_path.write_text(head + body + "\n\\end{document}\n", encoding="utf-8")
        try:
            res = run_tex(tex_path, work / f"build_{tag}", engine=eng, timeout=timeout)
        finally:
            try:
                tex_path.unlink()
            except OSError:
                pass
        if res.get("ok"):
            import pypdfium2 as pdfium
            pdf = Path(res["pdf"])
            n = len(pdfium.PdfDocument(str(pdf)))
            if n == len(group):
                return {"pdf": pdf, "res": res}
            res["error"] = f"số trang ảnh {n} ≠ số khối {len(group)}"
        return {"error": res.get("error"), "errors": res.get("errors", [])}

    out = compile_group(blocks, "all") if blocks else None
    plan: list[tuple[list[EngineBlock], dict | None]] = []
    if out and out.get("pdf"):
        plan.append((blocks, out))
    else:
        for i, b in enumerate(blocks):
            plan.append(([b], compile_group([b], f"b{i}")))
    for group, res in plan:
        if not res or not res.get("pdf"):
            for b in group:
                b.status, b.method = "unresolved", "unresolved"
                b.reason = "TeX không dựng được khối: " + str((res or {}).get("error") or "")
                b.warnings += [str(e) for e in (res or {}).get("errors", [])[:3]]
            continue
        pdf = res["pdf"]
        for i, b in enumerate(group):
            png, w, h = _pdf_page_to_png(pdf.read_bytes(), i, dpi)
            svg = _pdf_page_to_svg(pdf, i, work)
            b.asset = {"data": svg or png, "ext": "svg" if svg else "png", "png_data": png if svg else b"",
                       "width_pt": w, "height_pt": h, "engine": eng}
            b.method = "vector_image" if svg else "raster_image"
            b.status = "fallback" if b.kind in ("figure", "bbt") else "ok"


_HIDE_NODES = "\\tikzset{every node/.append style={opacity=0,text opacity=0}}"


def _geometry_truth(cands: list[EngineBlock], preamble: str, main_dir: Path, work: Path, *, dpi: int,
                    engine_pref: str, timeout: int = 240) -> None:
    """Ảnh TeX "chỉ hình học" (mọi node/nhãn opacity=0 — khung bao giữ nguyên) để đối chiếu với bản DrawingML khi
    trình render kiểm thử (LibreOffice) không vẽ hộp chữ trong nhóm hình. Lưu vào asset["geom_png"]."""
    from .tex_compile import choose_engine, run_tex
    eng = choose_engine(preamble, engine_pref)
    todo = [b for b in cands if b.asset and "\\begin{tikzpicture}" in b.tex]
    if not eng or not todo:
        return
    pre = re.sub(r"\\begin\s*\{document\}.*", "", preamble, flags=re.S)
    if "\\documentclass" not in pre:
        pre = "\\documentclass{article}\n\\usepackage{amsmath,amssymb,tikz,pgfplots,tkz-tab,graphicx}\n" + pre
    head = pre + "\n\\usepackage[active,tightpage]{preview}\n\\setlength\\PreviewBorder{2pt}\n\\begin{document}\n"
    body = []
    for b in todo:
        hidden = re.sub(r"(\\begin\{tikzpicture\}(?:\s*\[[^\]]*\])?)", lambda m: m.group(1) + _HIDE_NODES, b.tex,
                        count=1)
        body.append("\\begin{preview}\n" + hidden + "\n\\end{preview}\n")
    tex_path = main_dir / "_l2w_geom.tex"
    work.mkdir(parents=True, exist_ok=True)
    tex_path.write_text(head + "".join(body) + "\n\\end{document}\n", encoding="utf-8")
    try:
        res = run_tex(tex_path, work / "build_geom", engine=eng, timeout=timeout)
    finally:
        try:
            tex_path.unlink()
        except OSError:
            pass
    if not res.get("ok"):
        return
    pdf = Path(res["pdf"])
    import pypdfium2 as pdfium
    if len(pdfium.PdfDocument(str(pdf))) != len(todo):
        return
    data = pdf.read_bytes()
    for i, b in enumerate(todo):
        png, _w, _h = _pdf_page_to_png(data, i, dpi)
        b.asset["geom_png"] = png


def _verify_native_figures(cands: list[EngineBlock], work: Path, dpi: int) -> None:
    """Thử native: dựng DrawingML bằng engine thật (drawingml_engine), render bằng LibreOffice, đối chiếu với
    ảnh TeX của cùng khối. F1 mực ≥ 0.75 và đủ nhãn → giữ native (sửa được); ngược lại dùng ảnh TeX.
    Đây là phép đối chiếu heuristic, ghi lại trong `verify` của từng khối."""
    try:
        import docx as _docx
        from docx.enum.text import WD_BREAK
        from .drawingml_engine import convert_tikz_or_axis_to_drawingml
    except Exception as exc:
        for b in cands:
            b.verify = {"done": False, "reason": f"không nạp được drawingml_engine: {type(exc).__name__}"}
        return
    from .capabilities import find_tool
    lo = find_tool("soffice", "libreoffice")
    if not lo:
        for b in cands:
            b.verify = {"done": False, "reason": "thiếu LibreOffice để render thử DrawingML"}
        return
    from lxml import etree
    doc = _docx.Document()
    native_xml: dict[str, str] = {}
    for i, b in enumerate(cands):
        try:
            el = convert_tikz_or_axis_to_drawingml(b.tex, group_id=90000 + i)
            native_xml[b.block_id] = etree.tostring(el, encoding="unicode")
            p = doc.add_paragraph()
            p._p.append(el)
        except Exception as exc:
            b.verify = {"done": True, "native_ok": False, "reason": f"engine lỗi: {type(exc).__name__}: {exc}"}
            p = doc.add_paragraph("L2W-PROBE-ERROR")
        if i < len(cands) - 1:
            p.add_run().add_break(WD_BREAK.PAGE)
    probe = work / "probe_native.docx"
    work.mkdir(parents=True, exist_ok=True)
    doc.save(str(probe))
    try:
        subprocess.run([lo, "--headless", "-env:UserInstallation=" + (work / "lo_profile").resolve().as_uri(), "--convert-to", "pdf",
                        "--outdir", str(work), str(probe)], capture_output=True, timeout=180)
    except subprocess.TimeoutExpired:
        pass
    probe_pdf = work / "probe_native.pdf"
    if not probe_pdf.exists():
        for b in cands:
            b.verify = b.verify or {"done": False, "reason": "LibreOffice không render được bản thử"}
        return
    import numpy as np
    import pypdfium2 as pdfium
    from PIL import Image, ImageFilter
    pdoc = pdfium.PdfDocument(str(probe_pdf))

    def ink_crop(img: "Image.Image"):
        a = np.asarray(img.convert("L")) < 200
        ys, xs = np.where(a)
        if len(xs) == 0:
            return None
        return Image.fromarray(((a[ys.min(): ys.max() + 1, xs.min(): xs.max() + 1]) * 255).astype("uint8"))

    def f1(t_img, p_img) -> dict[str, float]:
        tw = 320
        th = max(8, int(t_img.height * tw / t_img.width))
        if abs((t_img.height / t_img.width) - (p_img.height / p_img.width)) > 0.35 * (t_img.height / t_img.width):
            return {"recall": 0.0, "precision": 0.0, "f1": 0.0, "aspect_mismatch": 1.0}
        t = t_img.resize((tw, th))
        p = p_img.resize((tw, th))
        td, pd = t.filter(ImageFilter.MaxFilter(7)), p.filter(ImageFilter.MaxFilter(7))
        ta, pa = np.asarray(t) > 0, np.asarray(p) > 0
        rec = float((ta & (np.asarray(pd) > 0)).sum() / max(1, ta.sum()))
        prec = float((pa & (np.asarray(td) > 0)).sum() / max(1, pa.sum()))
        return {"recall": round(rec, 3), "precision": round(prec, 3),
                "f1": round(2 * rec * prec / max(1e-9, rec + prec), 3)}

    for i, b in enumerate(cands):
        if b.verify and b.verify.get("done"):
            continue
        if i >= len(pdoc) or not b.asset:
            b.verify = {"done": False, "reason": "thiếu trang render hoặc ảnh TeX"}
            continue
        native_img = ink_crop(pdoc[i].render(scale=100 / 72).to_pil())
        truth = ink_crop(Image.open(io.BytesIO(b.asset["png_data"] or b.asset["data"])))
        if native_img is None or truth is None:
            b.verify = {"done": True, "native_ok": False, "reason": "bản native trống"}
            continue
        m_full = f1(truth, native_img)
        geom = ink_crop(Image.open(io.BytesIO(b.asset["geom_png"]))) if b.asset.get("geom_png") else None
        m = f1(geom, native_img) if geom is not None else m_full
        labels_src = _label_count_tikz(b.tex)
        labels_native = native_xml.get(b.block_id, "").count("<wps:txbx")
        # Nhãn: dùng bản "chỉ hình học" khi bản render native không có mực chữ (LibreOffice không vẽ hộp chữ trong
        # nhóm) → nhãn chỉ kiểm được SỐ LƯỢNG trong XML, vị trí chưa kiểm bằng hình.
        ok = m["f1"] >= 0.75 and labels_native >= labels_src
        b.verify = {"done": True, "native_ok": ok, **m, "f1_with_labels": m_full["f1"],
                    "compared_against": "tex_geometry_only" if geom is not None else "tex_full",
                    "labels_source": labels_src, "labels_native": labels_native,
                    "labels_check": "đếm hộp chữ trong XML (vị trí nhãn chưa kiểm bằng hình)",
                    "threshold_f1": 0.75, "method": "F1 mực (giãn 7px, rộng 320px) DrawingML↔ảnh TeX — heuristic",
                    "renderer": "LibreOffice"}


# ───────────────────────────── chuẩn bị ─────────────────────────────
def prepare_tex_for_engine(flat: FlattenResult, root: Path, work_dir: Path, *, main_rel: str, dpi: int = 400,
                           figure_policy: str = "auto", answer_mode: str = "source", show_solutions: bool = True,
                           tex_engine: str = "xelatex", structure: bool = True,
                           text_width_pt: float = 482.0) -> PreparedTex:
    text = flat.text
    prep = PreparedTex(text="", validation_text="")
    prep.assets_dir = work_dir / "_l2w_assets"
    prep.assets_dir.mkdir(parents=True, exist_ok=True)
    tr: dict[str, Any] = {}
    if answer_mode == "student":
        text, st = strip_answers_for_student(text, show_solutions)
        tr["student_mode"] = st
    elif not show_solutions:
        text, st = strip_answers_for_student_solutions_only(text)
        tr["solutions_hidden"] = st
    text, st = resolve_newif(text)
    tr["newif"] = st
    text, st = expand_user_macros(text)
    tr["user_macros"] = st
    try:
        from .macro_fidelity import rewrite_user_macros
        mf = rewrite_user_macros(text)
        if getattr(mf, "expanded", None):
            text = mf.text
        tr["macro_fidelity"] = mf.summary() if hasattr(mf, "summary") else {}
    except Exception as exc:
        tr["macro_fidelity"] = {"error": f"{type(exc).__name__}: {exc}"}
    # Văn bản đối chiếu = nguồn sau các phép biến đổi NGỮ NGHĨA (ẩn đáp án theo chế độ, điều kiện \newif, nở macro),
    # TRƯỚC các bước chuẩn hoá cấu trúc của chính pipeline — để validator bắt được nội dung bị bước đó làm rơi.
    prep.validation_text = text
    if structure:
        text, st = convert_lists(text)
        tr["lists"] = st
        text, st = normalize_structure(text, show_solutions=show_solutions)
        tr["structure"] = st
    prep.normalized_text = text
    text = _protect_verbatim(text, prep)

    masked = mask_tex(text)
    a, b = _body_span(text, masked)
    preamble = text[:a]
    math = _math_spans(masked, a, b)
    tikz_spans = _env_spans(masked, "tikzpicture", a, b)
    axis_alone = [s for s in _env_spans(masked, "axis", a, b) if not _inside(tikz_spans, s[0])]
    figs = sorted(tikz_spans + axis_alone)
    blocks: list[EngineBlock] = []
    counter = [0]

    def new_block(kind: str, s: int, e: int) -> EngineBlock:
        counter[0] += 1
        loc = flat.where(min(s, len(flat.text) - 1)) if flat.text else {"file": main_rel, "line": None}
        return EngineBlock(block_id=f"tex-{kind[:3]}-{counter[0]:04d}", kind=kind, source=loc, tex=text[s:e])

    try:
        from .math_tables import is_tkz_tab_block, tkz_tab_support
    except Exception:  # pragma: no cover
        is_tkz_tab_block = lambda t: "\\tkzTabInit" in t  # noqa: E731
        tkz_tab_support = lambda t: {"native": True, "reasons": []}  # noqa: E731
    fig_blocks: list[tuple[EngineBlock, int, int]] = []
    for s, e, _n in figs:
        blk = new_block("figure", s, e)
        if is_tkz_tab_block(blk.tex):
            blk.kind = "bbt"
            sup = tkz_tab_support(blk.tex)
            if sup["native"]:
                blk.method, blk.status = "table_native", "ok"
                blocks.append(blk)
                continue
            blk.reason = "; ".join(sup["reasons"]) or "bảng biến thiên ngoài phạm vi native"
            fig_blocks.append((blk, s, e))
            blocks.append(blk)
            continue
        if figure_policy == "native":
            blk.method, blk.status = "drawingml", "ok"
            blocks.append(blk)
            continue
        fig_blocks.append((blk, s, e))
        blocks.append(blk)

    # Ảnh TeX cho mọi ứng viên (cần để chụp HOẶC để đối chiếu native)
    snap = [blk for blk, _s, _e in fig_blocks]
    if snap:
        _snapshot_blocks(snap, preamble, (root / main_rel).parent, work_dir / "_l2w_snap", dpi=dpi,
                         engine_pref=tex_engine)
    if figure_policy == "auto":
        verifiable = [blk for blk in snap if blk.kind == "figure" and blk.asset]
        hard = [blk for blk in verifiable if _HARD_TIKZ_RE.search(blk.tex)]
        for blk in hard:
            blk.verify = {"done": False, "reason": "có cấu trúc TikZ chưa hỗ trợ native (danh sách heuristic) → ảnh TeX"}
        to_verify = [blk for blk in verifiable if blk not in hard]
        if to_verify:
            _geometry_truth(to_verify, preamble, (root / main_rel).parent, work_dir / "_l2w_snap", dpi=dpi,
                            engine_pref=tex_engine)
        _verify_native_figures(to_verify, work_dir / "_l2w_probe", dpi)
        shutil.rmtree(work_dir / "_l2w_probe" / "lo_profile", ignore_errors=True)
        for blk in snap:
            if blk.kind != "figure":
                continue
            if not blk.asset:
                # TeX không dựng được → giữ engine native (không tệ hơn trước), ghi rõ chưa đối chiếu
                blk.method, blk.status = "drawingml", "review"
                blk.reason = (blk.reason + "; " if blk.reason else "") + "không có ảnh TeX để đối chiếu → giữ DrawingML chưa kiểm chứng"
                continue
            v = blk.verify or {}
            if v.get("native_ok"):
                blk.method, blk.status, blk.asset = "drawingml", "ok", None
                blk.reason = ("DrawingML khớp ảnh TeX (heuristic F1 hình học "
                              f"{v.get('f1')}; nhãn {v.get('labels_native')}/{v.get('labels_source')} theo XML)")
                if v.get("compared_against") == "tex_geometry_only":
                    blk.warnings.append("vị trí nhãn chưa kiểm bằng hình (trình render kiểm thử không vẽ hộp chữ)")
            else:
                blk.reason = v.get("reason") or (f"DrawingML lệch ảnh TeX (F1={v.get('f1')}, nhãn "
                                                 f"{v.get('labels_native')}/{v.get('labels_source')})")
    elif figure_policy == "snapshot":
        for blk in snap:
            if blk.kind == "figure" and not blk.asset:
                blk.method, blk.status = "drawingml", "review"
                blk.reason = "TeX không dựng được → giữ DrawingML của engine"
    # Bảng biến thiên ngoài phạm vi native mà TeX cũng không dựng được → để engine (sẽ có thông báo hiển thị)
    for blk in snap:
        if blk.kind == "bbt" and not blk.asset:
            blk.method, blk.status = "unresolved", "unresolved"

    repls: list[tuple[int, int, str]] = []
    marker_n = [0]

    def marker_for(blk: EngineBlock, inline: bool) -> str:
        from .docx_media import make_img_marker
        marker_n[0] += 1
        blk.index = marker_n[0]
        blk.inline = inline
        m = make_img_marker(blk.index)
        return m if inline else "\n\n" + m + "\n\n"

    def is_inline(s: int, e: int) -> bool:
        ls = text.rfind("\n", 0, s) + 1
        le = text.find("\n", e)
        le = len(text) if le < 0 else le
        before = masked[ls:s].strip()
        after = masked[e:le].strip()
        return bool(before or after) and not re.fullmatch(r"(\\(?:centering|par|noindent)\s*)*", before + after)

    for blk, s, e in fig_blocks:
        if blk.method in ("vector_image", "raster_image") and blk.asset:
            repls.append((s, e, marker_for(blk, False)))

    # \includegraphics ngoài hình vẽ/toán
    gp = flat.graphicspath
    for m in re.finditer(r"\\includegraphics\s*(\*)?", masked[a:b]):
        s = a + m.start()
        if _inside(figs, s) or _inside(math, s):
            continue
        pos = a + m.end()
        opts, p2 = _read_opt(text, pos)
        if opts is not None:
            pos = p2
        # \includegraphics[..][..]{..} (dạng 2 tuỳ chọn cũ)
        opts2, p3 = _read_opt(text, pos)
        if opts2 is not None:
            pos = p3
        arg, pos_end = _read_arg(text, pos)
        if arg is None:
            continue
        blk = new_block("includegraphics", s, pos_end)
        src_file = blk.source.get("file") or main_rel
        bases = flat.base_dirs.get(src_file, [str(Path(main_rel).parent).replace(".", "")])
        path, how = resolve_graphic(arg, root, bases, gp)
        if path is None:
            blk.method, blk.status, blk.reason = "unresolved", "unresolved", f"ảnh '{arg}': {how}"
        else:
            if how != "ok":
                blk.warnings.append(f"ảnh '{arg}' → {path.name} ({how})")
            conv = convert_graphic(path, opts or "", dpi, prep.assets_dir)
            blk.warnings += conv.get("warnings", [])
            if conv.get("error"):
                blk.method, blk.status, blk.reason = "unresolved", "unresolved", f"ảnh '{arg}': {conv['error']}"
            else:
                w, h, warns = graphic_size(opts or "", conv["width_pt"], conv["height_pt"], text_width_pt)
                blk.warnings += warns
                conv["width_pt"], conv["height_pt"] = w, h
                blk.asset = conv
                blk.method, blk.status = conv["method"], "ok"
                blk.asset["file"] = str(path.relative_to(root)) if root in path.parents else path.name
        blocks.append(blk)
        repls.append((s, pos_end, marker_for(blk, is_inline(s, pos_end))))
    text = _apply(text, repls)

    # \immini{trái}{phải} — phải KHÔNG phải TikZ native → dựng lại bảng 2 cột sau engine
    masked = mask_tex(text)
    a, b = _body_span(text, masked)
    repls = []
    lay = 0
    for m in re.finditer(r"\\immini(?![A-Za-z@])", masked[a:b]):
        s = a + m.start()
        pos = a + m.end()
        _o, p2 = _read_opt(text, pos)
        if _o is not None:
            pos = p2
        left, pos = _read_arg(text, pos)
        right, pos = _read_arg(text, pos)
        if left is None or right is None:
            continue
        if re.search(r"\\begin\s*\{(?:tikzpicture|axis)\}", right):
            continue  # engine dựng immini + DrawingML như trước
        lay += 1
        blk = new_block("immini_layout", s, pos)
        blk.index = lay
        blk.method, blk.status = "layout", "pending"
        imgs_in_right = [x for x in blocks if x.index and f"ZLZIMGZ{x.index:05d}ZEND" in right]
        for x in imgs_in_right:
            x.max_width_frac = 0.40
            x.parent = blk.block_id
        blocks.append(blk)
        tag = f"{LAYOUT_PREFIX}{lay:05d}Z"
        repls.append((s, pos, f"\n\n{tag}BEG\n\n{left}\n\n{tag}MID\n\n{right}\n\n{tag}FIN\n\n"))
    text = _apply(text, repls)
    prep.text = text
    prep.blocks = blocks
    prep.transforms = tr
    for blk in blocks:
        if blk.status == "unresolved":
            _issue(prep, "warning", f"{blk.kind}_unresolved", f"{blk.block_id}: {blk.reason}", block_id=blk.block_id,
                   source=blk.source)
    return prep


VERB_PREFIX = "ZLZVERBZ"
_VERB_MARK_RE = re.compile(VERB_PREFIX + r"(\d{5})ZEND")


def _protect_verbatim(text: str, prep: PreparedTex) -> str:
    """verbatim/lstlisting/minted/Verbatim và \\verb: engine không hiểu (làm mất ngoặc, xử lý lệnh bên trong)
    → thay bằng chỗ đánh dấu; sau engine chèn lại NGUYÊN VĂN bằng phông đơn cách."""
    masked_full = mask_tex(text)  # verbatim đã bị che → dùng bản gốc để tìm
    a, b = _body_span(text, masked_full)
    repls: list[tuple[int, int, str]] = []
    n = 0
    for m in re.finditer(r"\\begin\s*\{(verbatim\*?|Verbatim\*?|lstlisting|minted)\}", text[a:b]):
        s0 = a + m.start()
        name = m.group(1)
        pos = a + m.end()
        if name == "minted":
            _o, pos = _read_opt(text, pos)[0], _read_opt(text, pos)[1]
            _lang, pos = _read_arg(text, pos)
        else:
            o, p2 = _read_opt(text, pos)
            if o is not None and name != "verbatim":
                pos = p2
        end_tok = "\\end{" + name + "}"
        e = text.find(end_tok, pos)
        if e < 0:
            continue
        content = text[pos:e].strip("\n")
        n += 1
        prep.verbatim[n] = {"text": content, "block": True, "env": name}
        repls.append((s0, e + len(end_tok), f"\n\n{VERB_PREFIX}{n:05d}ZEND\n\n"))
    for m in re.finditer(r"\\verb\*?([^A-Za-z\s*])", text[a:b]):
        s0 = a + m.start()
        if any(r0 <= s0 < r1 for r0, r1, _x in repls):
            continue
        if masked_full[s0] == " " and text[s0] == "\\":
            pass
        d = m.group(1)
        e = text.find(d, a + m.end())
        if e < 0 or "\n" in text[a + m.end(): e]:
            continue
        n += 1
        prep.verbatim[n] = {"text": text[a + m.end(): e], "block": False, "env": "verb"}
        repls.append((s0, e + 1, f"{VERB_PREFIX}{n:05d}ZEND"))
    # bỏ \verb nằm trong chú thích: chỉ nhận vị trí mà ký tự '\' không bị che vì chú thích
    out = []
    for r0, r1, rep in repls:
        line_start = text.rfind("\n", 0, r0) + 1
        seg = text[line_start:r0]
        if re.search(r"(?<!\\)%", seg):
            continue
        out.append((r0, r1, rep))
    return _apply(text, out)


def _inject_verbatim(doc, prep: PreparedTex) -> dict[str, int]:
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    st = {"verbatim_blocks": 0, "verbatim_inline": 0, "verbatim_missing": 0}
    found: set[int] = set()

    def run_xml(text: str) -> str:
        from xml.sax.saxutils import escape
        return (f'<w:r xmlns:w="{W}"><w:rPr><w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" '
                f'w:cs="Courier New"/><w:sz w:val="20"/></w:rPr><w:t xml:space="preserve">{escape(text)}</w:t></w:r>')

    for p in list(doc.element.body.iter(qn("w:p"))):
        ts = [t for t in p.iter(qn("w:t"))]
        full = "".join(t.text or "" for t in ts)
        ms = list(_VERB_MARK_RE.finditer(full))
        if not ms:
            continue
        if len(ms) == 1 and full.strip() == ms[0].group(0):
            idx = int(ms[0].group(1))
            info = prep.verbatim.get(idx)
            found.add(idx)
            if info is None:
                continue
            parent = p.getparent()
            pos = list(parent).index(p)
            for k, line in enumerate(info["text"].split("\n")):
                np_ = parse_xml(f'<w:p xmlns:w="{W}"><w:pPr><w:spacing w:before="0" w:after="0"/></w:pPr>'
                                + run_xml(line.replace("\t", "    ")) + "</w:p>")
                parent.insert(pos + k, np_)
            parent.remove(p)
            st["verbatim_blocks"] += 1
            continue
        # chỗ đánh dấu inline: thay trong từng w:t (chỗ đánh dấu chỉ gồm chữ/số nên không bị tách run)
        for t in ts:
            if not t.text or VERB_PREFIX not in t.text:
                continue
            r = t.getparent()
            pieces = _VERB_MARK_RE.split(t.text)
            # pieces: [chữ, idx, chữ, idx, ...]
            t.text = pieces[0]
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            anchor = r
            for j in range(1, len(pieces), 2):
                idx = int(pieces[j])
                found.add(idx)
                info = prep.verbatim.get(idx, {"text": ""})
                vr = parse_xml(run_xml(info["text"]))
                anchor.addnext(vr)
                anchor = vr
                tail = pieces[j + 1]
                if tail:
                    import copy
                    tr = copy.deepcopy(r)
                    for x in tr.findall(qn("w:t"))[1:]:
                        tr.remove(x)
                    tt = tr.find(qn("w:t"))
                    tt.text = tail
                    tt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    anchor.addnext(tr)
                    anchor = tr
                st["verbatim_inline"] += 1
    st["verbatim_missing"] = len([k for k in prep.verbatim if k not in found])
    return st


def strip_answers_for_student_solutions_only(text: str) -> tuple[str, dict[str, int]]:
    masked = mask_tex(text)
    a, b = _body_span(text, masked)
    repls = []
    n = 0
    for m in re.finditer(r"\\loigiai\s*\{", masked[a:b]):
        ob = a + m.end() - 1
        e = find_matching_brace(text, ob)
        if e > 0:
            repls.append((a + m.start(), e + 1, ""))
            n += 1
    for s_, e_, _n in _env_spans(masked, ("loigiai", "huongdan", "dapan"), a, b):
        repls.append((s_, e_, ""))
        n += 1
    return _apply(text, repls), {"solutions_hidden": n}


# ───────────────────────────── sau engine ─────────────────────────────
def inject_engine_blocks(docx_path: Path, prep: PreparedTex) -> dict[str, Any]:
    import docx as _docx
    from docx.oxml.ns import qn
    from .docx_media import IMG_MARKER_RE, MediaInjector
    doc = _docx.Document(str(docx_path))
    sec = doc.sections[0]
    text_w_pt = (int(sec.page_width) - int(sec.left_margin) - int(sec.right_margin)) / EMU_PER_PT
    by_index = {blk.index: blk for blk in prep.blocks if blk.index and blk.kind != "immini_layout"}
    inj = MediaInjector(doc)

    def resolver(idx: int) -> dict[str, Any] | None:
        blk = by_index.get(idx)
        if blk is None:
            return {"text": f"[L2W: CHỖ ĐÁNH DẤU ẢNH #{idx} KHÔNG RÕ NGUỒN]"}
        if not blk.asset:
            blk.status, blk.method = "unresolved", "unresolved"
            return {"text": f"[L2W: THIẾU HÌNH — {blk.reason or blk.tex[:80]}]"}
        a = blk.asset
        w, h = float(a["width_pt"]), float(a["height_pt"])
        cap = text_w_pt * blk.max_width_frac
        if w > cap:
            h *= cap / w
            w = cap
        kind = {"figure": "fig", "bbt": "bbt", "includegraphics": "img"}.get(blk.kind, "img")
        return {"data": a["data"], "ext": a["ext"], "png": a.get("png_data", b""),
                "cx": max(1, int(w * EMU_PER_PT)), "cy": max(1, int(h * EMU_PER_PT)),
                "name": f"L2W {blk.block_id}", "descr": f"l2w:{kind}:{blk.block_id}",
                "block": not blk.inline, "center": not blk.inline}

    inj.replace_markers(resolver, IMG_MARKER_RE)
    for idx, blk in by_index.items():
        if idx not in inj.found:
            blk.status, blk.method = "unresolved", "unresolved"
            blk.reason = (blk.reason + "; " if blk.reason else "") + "chỗ đánh dấu không còn trong Word (engine làm mất)"
    lay_stats = _rebuild_layouts(doc, prep)
    verb_stats = _inject_verbatim(doc, prep) if prep.verbatim else {}
    doc.save(str(docx_path))
    by_status: dict[str, int] = {}
    for blk in prep.blocks:
        by_status[blk.status] = by_status.get(blk.status, 0) + 1
    if verb_stats.get("verbatim_missing"):
        _issue(prep, "warning", "verbatim_lost",
               f"{verb_stats['verbatim_missing']} khối verbatim không còn chỗ đánh dấu trong Word (engine làm mất)")
    return {"media": dict(inj.stats), "layouts": lay_stats, "verbatim": verb_stats, "blocks_by_status": by_status}


def _rebuild_layouts(doc, prep: PreparedTex) -> dict[str, int]:
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    from .ooxml_tools import set_table_role
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    stats = {"rebuilt": 0, "failed": 0}
    sent_re = re.compile(LAYOUT_PREFIX + r"(\d{5})Z(BEG|MID|FIN)")
    found: dict[int, dict[str, Any]] = {}
    for p in list(doc.element.body.iter(qn("w:p"))):
        txt = "".join(t.text or "" for t in p.iter(qn("w:t")))
        m = sent_re.search(txt)
        if m:
            found.setdefault(int(m.group(1)), {})[m.group(2)] = p
    lay_blocks = {blk.index: blk for blk in prep.blocks if blk.kind == "immini_layout"}
    for idx, parts in found.items():
        blk = lay_blocks.get(idx)
        beg, mid, fin = parts.get("BEG"), parts.get("MID"), parts.get("FIN")
        ok = beg is not None and mid is not None and fin is not None and \
            beg.getparent() is mid.getparent() is fin.getparent()
        if ok:
            parent = beg.getparent()
            kids = list(parent)
            ib, im_, iff = kids.index(beg), kids.index(mid), kids.index(fin)
            ok = ib < im_ < iff
        if not ok:
            for p in parts.values():
                if p is not None and p.getparent() is not None:
                    p.getparent().remove(p)
            stats["failed"] += 1
            if blk:
                blk.status, blk.method = "fallback", "sequential"
                blk.reason = "không dựng lại được bảng 2 cột → nội dung giữ tuần tự (không mất)"
            continue
        left = kids[ib + 1: im_]
        right = kids[im_ + 1: iff]
        right_imgs = [x for x in prep.blocks if x.parent == (blk.block_id if blk else None) and x.asset]
        frac = 0.38
        if right_imgs:
            sec = doc.sections[0]
            tw = (int(sec.page_width) - int(sec.left_margin) - int(sec.right_margin)) / EMU_PER_PT
            wmax = max(min(float(x.asset["width_pt"]), tw * x.max_width_frac) for x in right_imgs)
            frac = min(0.5, max(0.25, wmax / tw + 0.04))
        lp = int(5000 * (1 - frac))
        tbl = parse_xml(
            f'<w:tbl xmlns:w="{W}"><w:tblPr><w:tblW w:w="5000" w:type="pct"/><w:jc w:val="center"/>'
            '<w:tblBorders><w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/><w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/><w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            '<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/></w:tblBorders>'
            '<w:tblLook w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="1" w:noVBand="1"/>'
            '</w:tblPr>'
            f'<w:tblGrid><w:gridCol w:w="{int(9638 * (1 - frac))}"/><w:gridCol w:w="{int(9638 * frac)}"/></w:tblGrid>'
            f'<w:tr><w:tc><w:tcPr><w:tcW w:w="{lp}" w:type="pct"/><w:vAlign w:val="top"/></w:tcPr></w:tc>'
            f'<w:tc><w:tcPr><w:tcW w:w="{5000 - lp}" w:type="pct"/><w:vAlign w:val="center"/></w:tcPr></w:tc></w:tr></w:tbl>')
        tcs = tbl.findall(f"{{{W}}}tr/{{{W}}}tc")
        for tc, elems in ((tcs[0], left), (tcs[1], right)):
            for el in elems:
                parent.remove(el)
                tc.append(el)
            if not elems or elems[-1].tag != qn("w:p"):
                tc.append(parse_xml(f'<w:p xmlns:w="{W}"/>'))
        set_table_role(tbl, "layout_immini", block=blk.block_id if blk else None)
        parent.insert(list(parent).index(beg), tbl)
        for p in (beg, mid, fin):
            parent.remove(p)
        stats["rebuilt"] += 1
        if blk:
            blk.status, blk.method = "ok", "layout_table"
    for idx, blk in lay_blocks.items():
        if idx not in found:
            blk.status = "fallback"
            blk.method = "unknown"
            blk.reason = "không thấy dấu bố cục trong Word (engine đã bỏ?)"
    return stats
