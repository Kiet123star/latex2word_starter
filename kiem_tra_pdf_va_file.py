#!/usr/bin/env python3
"""Kiểm thử nhánh PDF / .tex / .zip → Word QUA PIPELINE CÔNG KHAI (`backend.pipeline`).

    python kiem_tra_pdf_va_file.py                      # chạy mọi nhóm, kết quả vào ./ket_qua_kiem_tra/
    python kiem_tra_pdf_va_file.py --out D:/kq --qa     # bật render Word→PDF (cần LibreOffice/Word) cho ca QA
    python kiem_tra_pdf_va_file.py --nhom G2,G7         # chỉ chạy một số nhóm
    python kiem_tra_pdf_va_file.py --chay-test-cu       # chạy thêm kiem_tra_smart_paste.py, kiem_tra_universal.py

Mọi fixture được SINH LẠI từ mã (tests_l2w/tao_fixture.py) — không cần tải gì. Mỗi ca ghi:
  PASS        tiêu chí của ca (ghi rõ trong "checks") đều đúng trên máy này;
  FAIL        ít nhất một tiêu chí sai (xem "checks");
  SKIP        máy thiếu điều kiện (phông, TeX, LibreOffice, tesseract…) — KHÔNG tính là đạt;
  LIMITATION  hành vi đã biết chưa đạt mức lý tưởng; ca kiểm rằng pipeline BÁO CÁO TRUNG THỰC điều đó.
"PASS" chỉ có nghĩa với đúng bộ tiêu chí và đúng tập fixture này — không phải tuyên bố "chuyển đổi đúng 100%".

Kết quả: <out>/ket_qua_kiem_tra.json và <out>/BAO_CAO_KIEM_TRA.md (+ fixture, job, .docx để mở xem).
Mã thoát: 0 nếu không có FAIL; 1 nếu có FAIL.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import platform
import shutil
import subprocess
import sys
import threading
import time
import traceback
import unicodedata
import zipfile
from pathlib import Path
from typing import Any, Callable

REPO = Path(__file__).resolve().parent
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))
for _s in (sys.stdout, sys.stderr):  # Windows: in tiếng Việt ra console/tệp không lỗi mã hoá
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

from lxml import etree  # noqa: E402

from backend.pipeline import ConversionError, convert_bytes, convert_paste  # noqa: E402
from backend.ooxml_tools import docx_inventory, docx_text_stream, validate_package  # noqa: E402
from backend import validators as V  # noqa: E402
from tests_l2w import tao_fixture as F  # noqa: E402

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS = {"w": W_NS, "m": M_NS}


# ───────────────────────────── khung chạy ─────────────────────────────
class Skip(Exception):
    pass


class Limitation(Exception):
    pass


class Ctx:
    def __init__(self, out: Path, qa: bool) -> None:
        self.out = out
        self.fx = out / "fixtures"
        self.jobs = out / "jobs"
        self.qa = qa
        self.fx.mkdir(parents=True, exist_ok=True)
        self.jobs.mkdir(parents=True, exist_ok=True)
        self.docx_seen: list[tuple[str, str]] = []
        self._cache: dict[str, Path] = {}
        self.checks: list[dict[str, Any]] = []
        self.info: dict[str, Any] = {}

    # fixture có bộ nhớ đệm (sinh một lần)
    def fixture(self, key: str, fn: Callable[[], Path]) -> Path:
        if key not in self._cache:
            try:
                self._cache[key] = fn()
            except F.FixtureSkip as exc:
                raise Skip(str(exc)) from exc
        return self._cache[key]

    def write(self, name: str, text: str) -> Path:
        p = self.fx / name
        p.write_text(text, encoding="utf-8")
        return p

    def check(self, ok: bool, name: str, detail: Any = "") -> bool:
        self.checks.append({"ok": bool(ok), "name": name, "detail": _short(detail)})
        return bool(ok)

    def convert(self, filename: str, data: bytes, case_id: str, **opts: Any) -> dict[str, Any]:
        opts.setdefault("qa_pdf", False)
        opts.setdefault("dpi", 150)
        rep = convert_bytes(filename, data, jobs_root=self.jobs, **opts)
        self.docx_seen.append((case_id, rep["output"]["docx"]))
        return rep

    def convert_file(self, path: Path, case_id: str, **opts: Any) -> dict[str, Any]:
        return self.convert(path.name, path.read_bytes(), case_id, **opts)


def _short(x: Any, n: int = 400) -> Any:
    if isinstance(x, str):
        return x if len(x) <= n else x[:n] + "…"
    if isinstance(x, (list, tuple)):
        return [_short(i, 120) for i in list(x)[:20]]
    if isinstance(x, dict):
        return {k: _short(v, 160) for k, v in list(x.items())[:30]}
    return x


CASES: list[tuple[str, str, str, Callable[[Ctx], None]]] = []


def case(group: str, cid: str, title: str):
    def deco(fn: Callable[[Ctx], None]):
        CASES.append((group, cid, title, fn))
        return fn
    return deco


# ───────────────────────────── tiện ích đọc kết quả ─────────────────────────────
def doc_root(rep_or_path: Any) -> etree._Element:
    p = Path(rep_or_path["output"]["docx"]) if isinstance(rep_or_path, dict) else Path(rep_or_path)
    with zipfile.ZipFile(p) as z:
        return etree.fromstring(z.read("word/document.xml"))


def doc_text(rep: dict[str, Any], with_math: bool = True) -> str:
    out = []
    for para in docx_text_stream(doc_root(rep)):
        s = para["text"]
        if with_math and para.get("math"):
            s += " " + " ".join(para["math"])
        out.append(s)
    return unicodedata.normalize("NFC", "\n".join(out))


def inv(rep: dict[str, Any]) -> dict[str, Any]:
    return docx_inventory(Path(rep["output"]["docx"]))


def layer(rep: dict[str, Any], n: int) -> dict[str, Any]:
    for lay in rep.get("layers", []):
        if lay.get("layer") == n:
            return lay
    return {"status": "MISSING", "summary": ""}


def issue_codes(rep: dict[str, Any]) -> list[str]:
    return [i["code"] for i in rep.get("issues", [])]


def tables(rep: dict[str, Any], role: str | None = None) -> list[dict[str, Any]]:
    t = inv(rep)["tables"]
    return [x for x in t if role is None or x["role"] == role]


def omml_texts(root: etree._Element) -> list[str]:
    return ["".join(t.text or "" for t in om.iter(f"{{{M_NS}}}t")) for om in root.iter(f"{{{M_NS}}}oMath")]


def tool(name: str) -> bool:
    from backend.capabilities import find_tool
    return find_tool(name) is not None


def _force_rmtree(path: Path) -> None:
    """Xoá thư mục kết quả cũ — tệp đầu vào của job được đặt CHỈ ĐỌC (Windows không cho xoá nếu không gỡ cờ)."""
    import stat as _stat

    def _fix(func, p, *_a):
        try:
            os.chmod(p, _stat.S_IWRITE | _stat.S_IREAD)
            func(p)
        except Exception:
            pass
    if sys.version_info >= (3, 12):
        shutil.rmtree(path, onexc=_fix)
    else:  # pragma: no cover
        shutil.rmtree(path, onerror=_fix)


def expect_error(fn: Callable[[], Any]) -> ConversionError | None:
    try:
        fn()
    except ConversionError as exc:
        return exc
    return None


def zbytes(files: dict[str, str | bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in files.items():
            z.writestr(n, d.encode("utf-8") if isinstance(d, str) else d)
    return buf.getvalue()


# ═════════════════════════════ G1. cùng nguồn qua 3 đường ═════════════════════════════
@case("G1", "G1.1", "Cùng một nguồn qua dán mã / tệp .tex / gói .zip cho cùng kết quả")
def g1_same_source(ctx: Ctx) -> None:
    src = F.TEX_G1
    r_paste = convert_paste(src, jobs_root=ctx.jobs, qa_pdf=False, dpi=150)
    ctx.docx_seen.append(("G1.1", r_paste["output"]["docx"]))
    r_tex = ctx.convert("de.tex", src.encode("utf-8"), "G1.1")
    zp = ctx.fixture("zip_same", lambda: F.zip_same_source(ctx.fx))
    r_zip = ctx.convert_file(zp, "G1.1", main_file="de/main.tex")
    reps = {"paste": r_paste, "tex": r_tex, "zip": r_zip}
    for k, r in reps.items():
        ctx.check(r["conversion_completed"], f"{k}: tạo được .docx hợp lệ", r["package"].get("errors"))
    texts = {k: doc_text(r) for k, r in reps.items()}
    ctx.check(texts["paste"] == texts["tex"] == texts["zip"], "văn bản + công thức giống hệt nhau giữa 3 đường",
              {k: v[:160] for k, v in texts.items()})
    eq = {k: inv(r)["equation_nodes"] for k, r in reps.items()}
    ctx.check(len(set(eq.values())) == 1 and eq["tex"] > 0, "số công thức OMML bằng nhau", eq)
    roles = {k: inv(r)["tables_by_role"] for k, r in reps.items()}
    ctx.check(roles["paste"] == roles["tex"] == roles["zip"], "vai trò bảng giống nhau", roles)
    st = {k: [layer(r, n)["status"] for n in (1, 2, 3, 4)] for k, r in reps.items()}
    ctx.check(st["paste"] == st["tex"] == st["zip"], "trạng thái kiểm định L1–L4 giống nhau", st)
    ctx.check("Trường học, người thầy, ước mơ" in texts["tex"], "giữ đúng chữ có dấu", texts["tex"][-120:])
    ctx.info = {"layers": st["tex"], "equation_nodes": eq["tex"], "tables_by_role": roles["tex"],
                "preprocess": r_tex.get("preprocess", {}).get("mode")}


@case("G1", "G1.2", "QA hiển thị (render Word→PDF + so với PDF tham chiếu) cho .tex và .pdf")
def g1_qa(ctx: Ctx) -> None:
    if not ctx.qa:
        raise Skip("chạy với --qa để bật render Word→PDF (LibreOffice/Word)")
    from backend.capabilities import find_tool
    if not (find_tool("soffice", "libreoffice") or sys.platform == "win32"):
        raise Skip("máy không có LibreOffice (soffice) để render Word→PDF")
    r = ctx.convert("de_qa.tex", F.TEX_G1.encode("utf-8"), "G1.2", qa_pdf=True)
    l5 = layer(r, 5)
    ctx.check(l5["status"] in ("PASS", "WARN", "FAIL", "NOT_MEASURED"), "có kết quả nhóm E", l5["summary"])
    if l5["status"] == "NOT_MEASURED":
        ctx.check(r["quality_status"] != "passed", "chưa đo được → không được 'passed'", r["quality_status"])
    if "qa_renderer_drops_omml" in issue_codes(r):
        ctx.check(r["quality_status"] != "passed", "trình render bỏ OMML → báo cáo, không 'passed'", r["quality_status"])
    pdf = ctx.fixture("vn_pdf", lambda: F.pdf_vietnamese(ctx.fx))
    r2 = ctx.convert_file(pdf, "G1.2", qa_pdf=True)
    l5b = layer(r2, 5)
    ctx.check(l5b["status"] != "MISSING", "PDF: có nhóm E", l5b["summary"])
    ctx.info = {"tex_L5": l5["status"] + ": " + l5["summary"][:200], "pdf_L5": l5b["status"] + ": " + l5b["summary"][:200],
                "tex_issues": issue_codes(r)}


# ═════════════════════════════ G2. tiếng Việt ═════════════════════════════
@case("G2", "G2.1", "TeX: chữ dựng sẵn (NFC) và chữ tổ hợp (NFD) cho cùng văn bản Word")
def g2_tex_nfc_nfd(ctx: Ctx) -> None:
    nfc = F.TEX_VN_PRECOMPOSED
    head, body = nfc.split("\\begin{document}", 1)
    nfd = head + "\\begin{document}" + unicodedata.normalize("NFD", body)
    r1 = ctx.convert("vn_nfc.tex", nfc.encode("utf-8"), "G2.1")
    r2 = ctx.convert("vn_nfd.tex", nfd.encode("utf-8"), "G2.1")
    t1, t2 = doc_text(r1), doc_text(r2)
    ctx.check(t1 == t2, "văn bản Word (chuẩn hoá NFC) giống nhau", [t1[:120], t2[:120]])
    for ln in F.VN_LINES:
        ctx.check(ln in t1, f"có dòng: {ln[:40]}")
    ctx.check(layer(r1, 1)["details"].get("lost_diacritics", 0) == 0, "không mất dấu (L1)", layer(r1, 1)["summary"])


@case("G2", "G2.2", "PDF phông TrueType nhúng tập con (subset), chữ tiếng Việt dựng sẵn")
def g2_pdf_subset(ctx: Ctx) -> None:
    pdf = ctx.fixture("vn_pdf", lambda: F.pdf_vietnamese(ctx.fx))
    r = ctx.convert_file(pdf, "G2.2")
    t = doc_text(r)
    for ln in F.VN_LINES:
        ctx.check(ln in t, f"có dòng: {ln[:40]}", t[:200])
    ctx.check(layer(r, 1)["status"] == "PASS", "L1 PASS (chữ PDF ↔ Word)", layer(r, 1)["summary"])
    ctx.info = {"blocks": r["blocks_summary"]}


@case("G2", "G2.3", "PDF chữ tổ hợp (NFD: chữ cơ sở + dấu rời) → chữ dựng sẵn trong Word")
def g2_pdf_nfd(ctx: Ctx) -> None:
    pdf = ctx.fixture("vn_nfd_pdf", lambda: F.pdf_vietnamese(ctx.fx, nfd=True))
    r = ctx.convert_file(pdf, "G2.3")
    t = doc_text(r)
    for ln in F.VN_LINES:
        ctx.check(ln in t, f"có dòng (NFC): {ln[:40]}", t[:200])
    ctx.check(layer(r, 1)["status"] == "PASS", "L1 PASS", layer(r, 1)["summary"])
    ctx.check("text_mapping_suspect" not in issue_codes(r), "không báo nhầm ánh xạ sai", issue_codes(r))


@case("G2", "G2.4", "PDF thiếu ToUnicode (glyph không ánh xạ) → giữ ảnh glyph + cảnh báo, không bịa chữ")
def g2_pdf_no_tounicode(ctx: Ctx) -> None:
    pdf = ctx.fixture("vn_removed", lambda: F.pdf_broken_tounicode(ctx.fx, mode="removed"))
    r = ctx.convert_file(pdf, "G2.4")
    t = doc_text(r)
    ctx.check("unmapped_glyphs" in issue_codes(r), "có cảnh báo unmapped_glyphs", issue_codes(r))
    ctx.check(inv(r)["logical_images"] > 0, "glyph không đọc được giữ bằng ảnh tại chỗ", inv(r)["logical_images"])
    ctx.check("(cid:" not in t and "\ufffd" not in t, "không lọt mã (cid:N)/U+FFFD vào Word", t[:120])
    ctx.check(r["quality_status"] != "passed", "không được 'passed'", r["quality_status"])
    ctx.info = {"text_head": t[:100], "images": inv(r)["logical_images"]}


@case("G2", "G2.5", "PDF ToUnicode SAI có chủ đích (dịch mã) — có phát hiện được không")
def g2_pdf_shifted(ctx: Ctx) -> None:
    pdf = ctx.fixture("vn_shifted", lambda: F.pdf_broken_tounicode(ctx.fx, mode="shifted"))
    r = ctx.convert_file(pdf, "G2.5")
    detected = "text_mapping_suspect" in issue_codes(r)
    ctx.check(r["quality_status"] != "passed", "không được 'passed'", r["quality_status"])
    ctx.info = {"detected": detected, "text_head": doc_text(r)[:80],
                "note": "Phát hiện bằng heuristic (chữ trích ra không giống văn bản tự nhiên); ánh xạ sai 'hợp lý' "
                        "(vd. đổi chữ thành chữ khác vẫn có nguyên âm, có dấu cách) KHÔNG phát hiện được bằng chữ."}
    if not detected:
        raise Limitation("không phát hiện được ánh xạ ToUnicode sai trong fixture này")
    ctx.check(detected, "phát hiện (heuristic) và báo text_mapping_suspect")


@case("G2", "G2.6", "Tên tệp tiếng Việt / dấu cách / tên cấm của Windows / đường dẫn lạ")
def g2_filenames(ctx: Ctx) -> None:
    data = F.TEX_G1.encode("utf-8")
    r1 = ctx.convert("Đề thi số 1 (bản nháp).tex", data, "G2.6")
    d1 = Path(r1["output"]["docx"])
    ctx.check(d1.exists() and unicodedata.normalize("NFC", d1.name).startswith("Đề thi số 1"),
              "tên .docx giữ tiếng Việt + dấu cách", d1.name)
    r2 = ctx.convert("CON.tex", data, "G2.6")
    ctx.check(Path(r2["output"]["docx"]).stem.upper() not in ("CON", "PRN", "AUX", "NUL"),
              "tên cấm của Windows được đổi", Path(r2["output"]["docx"]).name)
    r3 = ctx.convert("..\\..\\thoat/ra:ngoai?.tex", data, "G2.6")
    d3 = Path(r3["output"]["docx"]).resolve()
    ctx.check(str(d3).startswith(str(ctx.jobs.resolve())), "tệp kết quả nằm trong thư mục job", str(d3))
    ctx.check(not any(ch in d3.name for ch in '<>:"|?*\\'), "không còn ký tự cấm trong tên", d3.name)


# ═════════════════════════════ G3. toán 2D ═════════════════════════════
@case("G3", "G3.1", "TeX: toán 2D phức tạp → OMML native đúng cấu trúc")
def g3_tex_math(ctx: Ctx) -> None:
    r = ctx.convert("toan_2d.tex", F.TEX_G3_MATH.encode("utf-8"), "G3.1")
    root = doc_root(r)

    def has(xpath: str) -> bool:
        return bool(root.xpath(xpath, namespaces=NS))

    ctx.check(layer(r, 2)["status"] == "PASS", "L2 PASS (khối toán khớp OMML)", layer(r, 2)["summary"])
    ctx.check(has(".//m:acc/m:accPr/m:chr[@m:val='\u20d7']"), "\\vec/\\overrightarrow → dấu kết hợp U+20D7")
    ctx.check(has(".//m:acc/m:accPr/m:chr[@m:val='\u20d6']"), "\\overleftarrow → U+20D6")
    ctx.check(has(".//m:acc/m:accPr/m:chr[@m:val='\u20e1']"), "\\overleftrightarrow → U+20E1")
    ctx.check(has(".//m:acc/m:accPr/m:chr[@m:val='\u0302']"), "\\hat/\\widehat → U+0302")
    ctx.check(has(".//m:d/m:dPr/m:begChr[@m:val='{']"), "hệ (cases/\\heva) → ngoặc nhọn co giãn")
    ctx.check(has(".//m:d/m:dPr/m:begChr[@m:val='[']"), "\\hoac → ngoặc vuông co giãn")
    ctx.check(has(".//m:m"), "ma trận → m:m")
    ctx.check(has(".//m:f"), "phân số → m:f")
    ctx.check(has(".//m:nary"), "tích phân/tổng → m:nary")
    ctx.check(has(".//m:rad/m:deg/m:r"), "căn bậc ba có chỉ số")
    ctx.check(has(".//m:limLow"), "giới hạn → m:limLow")
    oms = omml_texts(root)
    ctx.check(any("a→b" in o for o in oms), "\\rightarrow không cắt cụt công thức (lỗi gốc)", [o for o in oms if "→" in o][:3])
    ctx.check(any("⌈x⌉" in o or o == "x" for o in oms) or has(".//m:d/m:dPr/m:begChr[@m:val='⌈']"),
              "\\left\\lceil … \\right\\rceil dựng ngoặc trần")
    ctx.check(any("mod" in o for o in oms), "\\pmod in ra (mod 3)")
    ctx.check(not any("label" in o or "nonumber" in o or "&" in o for o in oms),
              "không lộ \\label/\\nonumber/& trong OMML", [o for o in oms if "&" in o or "label" in o][:2])
    iv = inv(r)
    ctx.check(iv["tex_residue_count"] == 0, "không rò lệnh TeX ra chữ", iv["tex_residue_samples"])
    ctx.info = {"equation_nodes": iv["equation_nodes"], "L2": layer(r, 2)["summary"]}


@case("G3", "G3.2", "PDF biên dịch từ TeX (đề có công thức) → công thức OMML hoặc ảnh vùng, không mất nội dung")
def g3_pdf_math(ctx: Ctx) -> None:
    src = REPO / "tests_l2w" / "fixtures" / "de_thi_mau.tex"
    if not src.exists():
        raise Skip("thiếu tests_l2w/fixtures/de_thi_mau.tex")
    pdf = ctx.fixture("de_thi_pdf", lambda: F.compile_tex_pdf(src, ctx.fx, "xelatex"))
    r = ctx.convert_file(pdf, "G3.2")
    l1, l2 = layer(r, 1), layer(r, 2)
    ctx.check(l1["status"] in ("PASS", "WARN"), "L1 không FAIL (chữ/số của PDF có trong Word)", l1["summary"])
    ctx.check(l2["status"] in ("PASS", "WARN"), "L2 không FAIL (công thức: OMML hoặc ảnh vùng)", l2["summary"])
    bs = r["blocks_summary"]
    unres = [b for b in r["blocks"]["items"] if b.get("method") == "unresolved"]
    ctx.check(not unres or any(i["code"].endswith("unresolved") for i in r["issues"]),
              "khối unresolved (nếu có) đều được báo", len(unres))
    ctx.info = {"blocks": bs, "L1": l1["summary"], "L2": l2["summary"],
                "math_as_image": r.get("pdf_manifest_summary", {}).get("inline_math_images"),
                "inline_math_omml": r.get("pdf_manifest_summary", {}).get("inline_math_omml")}


# ═════════════════════════════ G4. trắc nghiệm ═════════════════════════════
def _mcq_table(ctx: Ctx, kind: str, cid: str) -> dict[str, Any]:
    r = ctx.convert(f"tn_{kind}.tex", F.tex_mcq(kind).encode("utf-8"), cid)
    t = tables(r, "mcq")
    ctx.check(len(t) == 1, "có đúng 1 bảng trắc nghiệm (role=mcq)", inv(r)["tables_by_role"])
    return {"rep": r, "table": t[0] if t else {}}


@case("G4", "G4.1", "Trắc nghiệm phương án ngắn → 1×4 (4 cột)")
def g4_short(ctx: Ctx) -> None:
    x = _mcq_table(ctx, "short", "G4.1")
    ctx.check(str(x["table"].get("cols")) == "4", "4 cột", x["table"])


@case("G4", "G4.2", "Trắc nghiệm phương án vừa → 2×2")
def g4_medium(ctx: Ctx) -> None:
    x = _mcq_table(ctx, "medium", "G4.2")
    ctx.check(str(x["table"].get("cols")) == "2", "2 cột", x["table"])


@case("G4", "G4.3", "Trắc nghiệm phương án dài → 4×1")
def g4_long(ctx: Ctx) -> None:
    x = _mcq_table(ctx, "long", "G4.3")
    ctx.check(str(x["table"].get("cols")) == "1", "1 cột", x["table"])


@case("G4", "G4.4", "Đáp án đánh dấu TƯỜNG MINH (\\True) → answer_key=explicit; Đúng/Sai (\\choiceTF)")
def g4_explicit(ctx: Ctx) -> None:
    x = _mcq_table(ctx, "short", "G4.4")
    ctx.check(x["table"].get("answer_key") == "explicit", "answer_key=explicit", x["table"])
    r = ctx.convert("dung_sai.tex", F.TEX_TF.encode("utf-8"), "G4.4")
    tf = tables(r, "true_false")
    ctx.check(len(tf) == 1 and tf[0].get("answer_key") == "explicit", "bảng Đúng/Sai có đáp án tường minh", tf)
    txt = doc_text(r)
    ctx.check(txt.count("Đ") >= 2 and "S" in txt, "điền Đ/S theo \\True", txt[-200:])


@case("G4", "G4.5", "Chỉ in đậm nhãn (không có dấu tường minh) → answer_key=unknown, không suy đáp án")
def g4_bold_only(ctx: Ctx) -> None:
    r = ctx.convert("in_dam.tex", F.TEX_BOLD_ONLY.encode("utf-8"), "G4.5")
    t = tables(r, "mcq")
    ctx.check(len(t) == 1, "nhận bảng 4 phương án", inv(r)["tables_by_role"])
    if t:
        ctx.check(t[0].get("answer_key") == "unknown", "answer_key=unknown (không suy từ in đậm)", t[0])


@case("G4", "G4.6", "Chế độ học sinh: không lộ đáp án (gạch chân, cột Đ/S, lời giải)")
def g4_student(ctx: Ctx) -> None:
    r = ctx.convert("hs.tex", F.TEX_G1.encode("utf-8"), "G4.6", answer_mode="student", show_solutions=False)
    root = doc_root(r)
    ctx.check(not root.xpath(".//w:tbl[.//w:tblDescription[contains(@w:val,'role=mcq')]]//w:u", namespaces=NS),
              "bảng trắc nghiệm không có gạch chân đánh dấu")
    t = tables(r, "mcq")
    ctx.check(bool(t) and all(x.get("answer_display") == "hidden" for x in t), "answer_display=hidden", t)
    ctx.check("Lời giải" not in doc_text(r) and "suy ra" not in doc_text(r), "ẩn lời giải")
    r2 = ctx.convert("hs_tf.tex", F.TEX_TF.encode("utf-8"), "G4.6", answer_mode="student")
    root2 = doc_root(r2)
    cells = root2.xpath(".//w:tbl[.//w:tblDescription[contains(@w:val,'role=true_false')]]/w:tr[position()>1]/w:tc[position()>1]",
                        namespaces=NS)
    filled = ["".join(c.itertext()).strip() for c in cells if "".join(c.itertext()).strip()]
    ctx.check(cells and not filled, "cột Đúng/Sai để trống", filled)


@case("G4", "G4.7", "PDF trắc nghiệm (A./B./C./D. cùng dòng và 2×2) → bảng phương án, answer_key=unknown")
def g4_pdf_mcq(ctx: Ctx) -> None:
    pdf = ctx.fixture("mcq_pdf", lambda: F.pdf_mcq(ctx.fx))
    r = ctx.convert_file(pdf, "G4.7")
    t = tables(r, "layout_mcq_pdf")
    ctx.check(len(t) == 2, "2 bảng phương án", inv(r)["tables_by_role"])
    ctx.check(all(x.get("answer_key") in ("unknown", "answer_key_unknown") for x in t), "answer_key=unknown", t)
    txt = doc_text(r)
    for s in ("A. 4.", "D. 7.", "Hà Nội", "Cần Thơ"):
        ctx.check(s in txt, f"có '{s}'")
    ctx.info = {"cols": [x.get("max_cols") for x in t]}


# ═════════════════════════════ G5. bảng ═════════════════════════════
@case("G5", "G5.1", "Bảng biến thiên nhiều dòng (2 dòng biến thiên) → bảng Word native")
def g5_bbt(ctx: Ctx) -> None:
    r = ctx.convert("bbt.tex", F.TEX_BBT.encode("utf-8"), "G5.1")
    v = tables(r, "variation")
    ctx.check(len(v) == 2, "2 bảng biến thiên (role=variation)", inv(r)["tables_by_role"])
    if v:
        ctx.check(v[0]["rows"] >= 4, "bảng 1 có ≥4 hàng (x, f', f, g)", v[0])
    ctx.check("bbt_unresolved" not in issue_codes(r), "không có BBT chưa dựng", issue_codes(r))
    ctx.check(layer(r, 4)["status"] == "PASS", "L4 PASS", layer(r, 4)["summary"])
    ctx.info = {"variation": v}


@case("G5", "G5.2", "Tiệm cận đứng hai phía (-D+, d) → một ô có giá trị hai bên vạch kép")
def g5_asymptote(ctx: Ctx) -> None:
    r = ctx.convert("bbt_tc.tex", F.TEX_BBT.encode("utf-8"), "G5.2")
    root = doc_root(r)
    tbls = root.xpath(".//w:tbl[.//w:tblDescription[contains(@w:val,'role=variation')]]", namespaces=NS)
    ctx.check(len(tbls) == 2, "có bảng tiệm cận", len(tbls))
    if len(tbls) == 2:
        both = [tc for tc in tbls[1].iter(f"{{{W_NS}}}tc")
                if "−∞" in "".join(tc.itertext()) and "+∞" in "".join(tc.itertext())]
        ctx.check(bool(both), "một ô chứa cả −∞ và +∞ (hai phía tiệm cận)")
        if both:
            b = both[0].find(f"{{{W_NS}}}tcPr/{{{W_NS}}}tcBorders")
            ok = b is not None and b.find(f"{{{W_NS}}}left") is not None and b.find(f"{{{W_NS}}}right") is not None
            ctx.check(ok, "ô tiệm cận có hai vạch dọc (vạch kép)")


@case("G5", "G5.3", "Bảng gộp ô (\\multirow + \\multicolumn) giữ đủ ô")
def g5_merged(ctx: Ctx) -> None:
    r = ctx.convert("gop_o.tex", F.TEX_MERGED.encode("utf-8"), "G5.3")
    root = doc_root(r)
    ctx.check(bool(root.xpath(".//w:vMerge", namespaces=NS)), "có gộp dọc (vMerge)")
    ctx.check(bool(root.xpath(".//w:gridSpan", namespaces=NS)), "có gộp ngang (gridSpan)")
    txt = doc_text(r)
    for s in ("Nhóm", "Kết quả", "Đạt", "Không đạt", "Lớp A", "12", "3"):
        ctx.check(s in txt, f"có ô '{s}'", txt)
    ctx.check(inv(r)["empty_tables"] == 0, "không có bảng rỗng")


@case("G5", "G5.4", "longtable dài → lặp dòng tiêu đề khi sang trang (tblHeader)")
def g5_longtable(ctx: Ctx) -> None:
    r = ctx.convert("bang_dai.tex", F.tex_longtable(40).encode("utf-8"), "G5.4")
    root = doc_root(r)
    ctx.check(bool(root.xpath(".//w:trPr/w:tblHeader", namespaces=NS)), "dòng tiêu đề lặp (tblHeader)")
    d = tables(r, "data")
    ctx.check(bool(d) and int(d[0]["rows"]) >= 41, "đủ 41 hàng", d[:1])
    ctx.check("Học sinh 40" in doc_text(r), "giữ hàng cuối")


@case("G5", "G5.5", "PDF bảng trải nhiều trang → bảng native từng trang, không mất hàng")
def g5_pdf_table_pages(ctx: Ctx) -> None:
    pdf = ctx.fixture("long_table_pdf", lambda: F.pdf_long_table(ctx.fx))
    r = ctx.convert_file(pdf, "G5.5")
    txt = doc_text(r)
    missing = [i for i in range(1, 71) if f"Học sinh số {i}\n" not in txt + "\n"]
    ctx.check(not missing, "đủ 70 hàng học sinh", missing[:10])
    n_tab = r["blocks_summary"]["by_method"].get("table_native", 0)
    ctx.check(n_tab >= 2, "bảng native ở mỗi trang", r["blocks_summary"])
    ctx.check(layer(r, 4)["status"] == "PASS", "L4 PASS", layer(r, 4)["summary"])
    ctx.info = {"note": "Bảng trải nhiều trang được dựng thành MỘT bảng Word cho MỖI trang PDF (giữ dòng tiêu đề lặp như "
                        "bản in); chưa tự nối thành một bảng liền.", "tables": n_tab}


# ═════════════════════════════ G6. hình ═════════════════════════════
def _fig_blocks(rep: dict[str, Any]) -> list[dict[str, Any]]:
    return [b for b in (rep.get("blocks") or {}).get("items", []) if b.get("kind") in ("figure", "includegraphics")]


@case("G6", "G6.1", "Hình TikZ đơn giản → DrawingML (đã kiểm chứng) hoặc ảnh vector; không mất hình")
def g6_simple(ctx: Ctx) -> None:
    r = ctx.convert("hinh_don_gian.tex", F.TEX_FIG_SIMPLE.encode("utf-8"), "G6.1")
    fb = _fig_blocks(r)
    ctx.check(len(fb) == 1, "1 khối hình", [b.get("method") for b in fb])
    if fb:
        ctx.check(fb[0]["method"] in ("drawingml", "vector_image", "raster_image"), "dựng native hoặc giữ ảnh", fb[0])
    ctx.check(layer(r, 3)["status"] in ("PASS", "WARN"), "L3 không FAIL", layer(r, 3)["summary"])
    ctx.check("Đoạn sau hình." in doc_text(r), "giữ đoạn sau hình")
    ctx.info = {"figure": [(b["method"], b.get("status"), b.get("reason")) for b in fb]}


@case("G6", "G6.2", "Hình cong/tô/nét đứt/3D → giữ ảnh vùng (snapshot TeX), không dựng DrawingML sai")
def g6_hard(ctx: Ctx) -> None:
    r = ctx.convert("hinh_kho.tex", F.TEX_FIG_HARD.encode("utf-8"), "G6.2")
    fb = _fig_blocks(r)
    ctx.check(len(fb) == 1, "1 khối hình", [b.get("method") for b in fb])
    m = fb[0]["method"] if fb else None
    ctx.info = {"figure": [(b["method"], b.get("status"), b.get("reason")) for b in fb]}
    if m in ("vector_image", "raster_image"):
        ctx.check(True, f"giữ ảnh vùng ({m})")
        ctx.check(inv(r)["logical_images"] >= 1, "ảnh có trong Word")
    else:
        ctx.check(inv(r)["l2w_notices"] or r["quality_status"] != "passed", "không dựng được → có thông báo/không passed",
                  inv(r)["l2w_notices"])
        raise Limitation(f"máy này không chụp được hình bằng TeX → khối hình: {m}")


@case("G6", "G6.3", "\\immini (chữ + hình bên cạnh) → bảng bố cục; SVG+PNG tính MỘT ảnh logic")
def g6_immini(ctx: Ctx) -> None:
    r = ctx.convert("immini.tex", F.TEX_IMMINI.encode("utf-8"), "G6.3", figure_policy="snapshot")
    iv = inv(r)
    ctx.check(iv["tables_by_role"].get("layout_immini", 0) == 1, "1 bảng bố cục layout_immini", iv["tables_by_role"])
    ctx.check("Tính độ dài cạnh" in doc_text(r), "giữ phần chữ bên trái")
    if iv["logical_images"]:
        ctx.check(iv["logical_images"] == 1, "đúng 1 ảnh logic", iv["logical_images"])
        if iv["svg_backed_images"]:
            ctx.check(iv["media_parts"] == 2 * iv["logical_images"], "SVG + PNG dự phòng = 2 media cho 1 ảnh logic",
                      {k: iv[k] for k in ("logical_images", "svg_backed_images", "media_parts")})
    else:
        raise Limitation("không chụp được hình bằng TeX trên máy này (snapshot) → hình giữ dạng khác")
    ctx.info = {k: iv[k] for k in ("logical_images", "svg_backed_images", "media_parts", "tables_by_role")}


@case("G6", "G6.4", "PDF: hình đơn giản + hình cong/tô; nhãn đỉnh sát mép nằm trong vùng cắt")
def g6_pdf_figures(ctx: Ctx) -> None:
    pdf = ctx.fixture("fig_pdf", lambda: F.pdf_figures(ctx.fx))
    r = ctx.convert_file(pdf, "G6.4")
    figs = [b for b in r["blocks"]["items"] if b.get("type") in ("figure", "image")]
    ctx.check(len(figs) == 2, "2 khối hình", [(b["type"], b["method"]) for b in figs])
    ctx.check(all(b["method"] in ("vector_image", "raster_image", "drawingml") for b in figs), "hình được giữ")
    ctx.check(layer(r, 1)["status"] == "PASS", "nhãn A/B/C/M không bị mất (L1 tính cả chữ/số, trừ vùng ảnh)",
              layer(r, 1)["summary"])
    ctx.check(layer(r, 3)["status"] == "PASS", "L3 PASS", layer(r, 3)["summary"])
    lines = [ln.strip() for ln in doc_text(r).split("\n")]
    ctx.check(not any(ln in ("A", "B", "C", "M") for ln in lines), "nhãn không bị tách thành đoạn chữ rời", lines)
    ctx.info = {"figures": [(b["method"], b.get("bbox")) for b in figs]}


# ═════════════════════════════ G7. cấu trúc PDF ═════════════════════════════
@case("G7", "G7.1", "PDF 2 cột (đường cơ sở thẳng hàng) → đọc hết cột trái rồi cột phải")
def g7_two_columns(ctx: Ctx) -> None:
    pdf = ctx.fixture("two_col", lambda: F.pdf_two_columns(ctx.fx))
    r = ctx.convert_file(pdf, "G7.1")
    t = doc_text(r)
    a, b = t.find("Cột trái câu số 8"), t.find("Cột phải dòng thứ 1")
    ctx.check(a >= 0 and b >= 0 and a < b, "toàn bộ cột trái đứng trước cột phải", t[:200])
    ctx.check(t.find("Cột phải dòng thứ 1") < t.find("Cột phải dòng thứ 8"), "thứ tự trong cột phải đúng")


@case("G7", "G7.2", "Trang xoay (/Rotate 90) + trang có CropBox nhỏ")
def g7_rotate_crop(ctx: Ctx) -> None:
    pdf = ctx.fixture("rot_crop", lambda: F.pdf_rotated_and_cropbox(ctx.fx))
    r = ctx.convert_file(pdf, "G7.2")
    t = doc_text(r)
    ctx.check("Trang xoay chín mươi độ vẫn đọc thẳng." in t, "chữ trang xoay đọc thẳng, đúng thứ tự", t[:120])
    ctx.check(t.find("Trang xoay") < t.find("Dòng thứ hai"), "thứ tự dòng đúng")
    ctx.check("Chữ trong vùng cắt hiển thị." in t, "giữ chữ trong CropBox")
    ctx.check("NGOAIVUNGCAT" not in t, "bỏ chữ nằm ngoài CropBox (không hiển thị)")
    ctx.check("outside_cropbox" in issue_codes(r), "báo outside_cropbox", issue_codes(r))
    figs = [b for b in r["blocks"]["items"] if b.get("page") == 1 and b.get("method") in ("vector_image", "raster_image")]
    ctx.check(len(figs) == 1, "hình trên trang xoay được cắt từ PDF gốc", len(figs))
    if figs:
        x0, y0, x1, y1 = figs[0]["bbox"]
        pg = next(p for p in r.get("blocks", {}).get("items", []) if p.get("page") == 1)
        ctx.check(0 <= x0 < x1 <= 596 and 0 <= y0 < y1 <= 843, "toạ độ vùng cắt trong khổ trang hiển thị (dọc)",
                  figs[0]["bbox"])
        ctx.info = {"figure_bbox_display_pt": figs[0]["bbox"], "first_block": pg.get("block_id")}


@case("G7", "G7.3", "Chữ nằm trong Form XObject")
def g7_form_xobject(ctx: Ctx) -> None:
    pdf = ctx.fixture("xobj", lambda: F.pdf_form_xobject(ctx.fx))
    r = ctx.convert_file(pdf, "G7.3")
    t = doc_text(r)
    ctx.check("Chữ nằm trong Form XObject." in t and "Dòng thứ hai trong XObject." in t, "đọc được chữ trong XObject", t)
    ctx.check("Chữ thường ngoài XObject." in t, "giữ chữ ngoài XObject")


@case("G7", "G7.4", "Trang scan, OCR mặc định (cần gói 'vie') → giữ ảnh trang + cảnh báo, không bịa chữ")
def g7_scan_default(ctx: Ctx) -> None:
    pdf = ctx.fixture("scan", lambda: F.pdf_scan(ctx.fx))
    r = ctx.convert_file(pdf, "G7.4")
    types = r["source"]["page_types"]
    ctx.check(types.get("1") == "scan", "nhận trang scan", types)
    bm = r["blocks_summary"]["by_method"]
    codes = issue_codes(r)
    if bm.get("page_image"):
        ctx.check(any(c in codes for c in ("ocr_lang_missing", "ocr_unavailable", "ocr_low_confidence", "ocr_failed")),
                  "có cảnh báo lý do không OCR", codes)
        ctx.check(layer(r, 1)["status"] == "NOT_MEASURED", "L1 NOT_MEASURED (không có chữ để đối chiếu)",
                  layer(r, 1)["summary"])
    else:
        ocr_blocks = [b for b in r["blocks"]["items"] if b.get("ocr")]
        ctx.check(bool(ocr_blocks) and all(b["status"] == "review" for b in ocr_blocks), "chữ OCR gắn trạng thái review")
    ctx.check(r["quality_status"] != "passed", "không 'passed'", r["quality_status"])
    ctx.info = {"by_method": bm, "issues": codes}


@case("G7", "G7.5", "Trang scan tiếng Anh với ocr_lang=eng → chữ OCR (review) hoặc ảnh trang")
def g7_scan_eng(ctx: Ctx) -> None:
    if not tool("tesseract"):
        raise Skip("máy không có tesseract")
    pdf = ctx.fixture("scan", lambda: F.pdf_scan(ctx.fx))
    r = ctx.convert_file(pdf, "G7.5", ocr_lang="eng")
    ocr_blocks = [b for b in r["blocks"]["items"] if b.get("ocr")]
    t = doc_text(r).lower()
    if not ocr_blocks:
        ctx.check(any(c.startswith("ocr_") for c in issue_codes(r)), "không OCR được → có lý do", issue_codes(r))
        raise Limitation("OCR không đạt ngưỡng tin cậy trên fixture này → giữ ảnh trang")
    ctx.check("quick brown fox" in t, "đọc được 'quick brown fox'", t[:160])
    ctx.check(all(b["status"] == "review" for b in ocr_blocks), "khối OCR ở trạng thái review")
    ctx.check(all((b.get("confidence") or {}).get("kind") == "tesseract_conf_mean" for b in ocr_blocks),
              "độ tin cậy ghi rõ nguồn (tesseract_conf_mean)")
    ctx.check(r["quality_status"] != "passed", "không 'passed' (chữ OCR cần soát)", r["quality_status"])


@case("G7", "G7.6", "Trang hỗn hợp (chữ + ảnh chụp)")
def g7_mixed(ctx: Ctx) -> None:
    pdf = ctx.fixture("mixed", lambda: F.pdf_mixed(ctx.fx))
    r = ctx.convert_file(pdf, "G7.6")
    bm = r["blocks_summary"]["by_method"]
    ctx.check(bm.get("raster_image", 0) + bm.get("vector_image", 0) >= 1, "ảnh chụp giữ dạng ảnh", bm)
    t = doc_text(r)
    for s in ("Trang hỗn hợp có chữ và ảnh chụp.", "Chú thích dưới ảnh chụp minh hoạ.", "Đoạn văn tiếp theo sau ảnh."):
        ctx.check(s in t, f"có '{s[:30]}'")
    ctx.check(t.find("Trang hỗn hợp") < t.find("Chú thích dưới"), "thứ tự chữ đúng")


@case("G7", "G7.7", "Scan có sẵn lớp chữ OCR ẩn → không nhân đôi chữ, báo lớp OCR")
def g7_hidden_ocr(ctx: Ctx) -> None:
    pdf = ctx.fixture("scan_hidden", lambda: F.pdf_scan(ctx.fx, hidden_text=True))
    r = ctx.convert_file(pdf, "G7.7")
    ctx.check(r["source"]["page_types"].get("1") == "scan_with_hidden_text", "nhận scan có lớp chữ ẩn",
              r["source"]["page_types"])
    t = doc_text(r).lower()
    ctx.check(t.count("quick brown fox") <= 1, "không lặp chữ (ẩn + OCR)", t[:120])
    ctx.check("scan_has_ocr_layer" in issue_codes(r), "báo scan_has_ocr_layer", issue_codes(r))


# ═════════════════════════════ G8. gói ZIP ═════════════════════════════
@case("G8", "G8.1", "ZIP: \\input/\\include lồng nhiều cấp, main trong thư mục con")
def g8_nested(ctx: Ctx) -> None:
    z = ctx.fixture("zip_nested", lambda: F.zip_nested(ctx.fx))
    r = ctx.convert_file(z, "G8.1", main_file="goc/main.tex")
    t = doc_text(r)
    for s in ("Mở đầu tài liệu.", "Chương một nội dung.", "Phần lồng sâu", "Chương hai được include.", "Kết thúc tài liệu."):
        ctx.check(s in t, f"có '{s}'", t[:300])
    ctx.check(t.find("Chương một") < t.find("Phần lồng sâu") < t.find("Chương hai") < t.find("Kết thúc"),
              "đúng thứ tự include")
    ctx.check(r["source"]["main"] == "goc/main.tex", "main đúng đường dẫn tương đối", r["source"].get("main"))


@case("G8", "G8.2", "ZIP: main_tex SAI → lỗi 409 kèm danh sách ứng viên (không tự chọn)")
def g8_wrong_main(ctx: Ctx) -> None:
    z = ctx.fixture("zip_nested", lambda: F.zip_nested(ctx.fx))
    err = expect_error(lambda: ctx.convert_file(z, "G8.2", main_file="khong_co.tex"))
    ctx.check(err is not None, "ném ConversionError")
    if err:
        ctx.check(err.http_status == 409 and err.code == "main_tex_not_found", "409 main_tex_not_found",
                  (err.http_status, err.code))
        cands = [c.get("path") for c in err.extra.get("candidates") or []]
        ctx.check("goc/main.tex" in cands, "ứng viên có goc/main.tex", cands)


@case("G8", "G8.3", "ZIP: nhiều tệp chính → 409 main_tex_ambiguous; chọn rõ thì đúng tệp")
def g8_multi_main(ctx: Ctx) -> None:
    z = ctx.fixture("zip_multi", lambda: F.zip_multi_main(ctx.fx))
    err = expect_error(lambda: ctx.convert_file(z, "G8.3"))
    ctx.check(err is not None and err.http_status == 409 and err.code == "main_tex_ambiguous", "409 main_tex_ambiguous",
              (err.http_status, err.code) if err else None)
    r = ctx.convert_file(z, "G8.3", main_file="de2.tex")
    t = doc_text(r)
    ctx.check("Tài liệu thứ hai." in t and "Tài liệu thứ nhất." not in t, "chuyển đúng tệp đã chọn", t)


@case("G8", "G8.4", "ZIP: ảnh thiếu → thông báo hiện trong Word, giữ chữ trước/sau, không 'passed'")
def g8_missing_asset(ctx: Ctx) -> None:
    z = ctx.fixture("zip_missing", lambda: F.zip_missing_asset(ctx.fx))
    r = ctx.convert_file(z, "G8.4")
    t = doc_text(r)
    ctx.check("Trước ảnh." in t and "Sau ảnh." in t, "giữ chữ trước/sau")
    ctx.check(bool(inv(r)["l2w_notices"]) or "[L2W" in t, "có thông báo hiện trong Word", inv(r)["l2w_notices"])
    ctx.check(any("unresolved" in c or "missing" in c for c in issue_codes(r)), "issue khối/tệp thiếu", issue_codes(r))
    ctx.check(layer(r, 3)["status"] in ("FAIL", "WARN"), "L3 không PASS", layer(r, 3)["summary"])
    ctx.check(r["quality_status"] != "passed", "không 'passed'", r["quality_status"])


@case("G8", "G8.5", "ZIP: ảnh có thật qua \\graphicspath → ảnh trong Word")
def g8_image_ok(ctx: Ctx) -> None:
    z = ctx.fixture("zip_img", lambda: F.zip_with_image(ctx.fx))
    r = ctx.convert_file(z, "G8.5")
    ctx.check(inv(r)["logical_images"] == 1, "1 ảnh", inv(r)["logical_images"])
    ctx.check(layer(r, 3)["status"] == "PASS", "L3 PASS", layer(r, 3)["summary"])


@case("G8", "G8.6", "ZIP: vòng lặp \\input → phát hiện, không treo, nội dung không lặp vô hạn")
def g8_cycle(ctx: Ctx) -> None:
    z = ctx.fixture("zip_cycle", lambda: F.zip_cycle(ctx.fx))
    t0 = time.time()
    r = ctx.convert_file(z, "G8.6")
    ctx.check(time.time() - t0 < 120, "kết thúc nhanh")
    ctx.check("include_cycle" in issue_codes(r), "issue include_cycle", issue_codes(r))
    t = doc_text(r)
    ctx.check(t.count("Tệp A.") == 1 and t.count("Tệp B.") == 1, "mỗi tệp xuất hiện một lần", t)


@case("G8", "G8.7", "ZIP: gói LaTeX không tồn tại → vẫn chuyển nội dung (không cần biên dịch)")
def g8_bad_package(ctx: Ctx) -> None:
    z = ctx.fixture("zip_badpkg", lambda: F.zip_bad_package(ctx.fx))
    r = ctx.convert_file(z, "G8.7")
    ctx.check(r["conversion_completed"], "tạo được .docx")
    ctx.check("Tài liệu dùng gói không tồn tại" in doc_text(r), "giữ nội dung")
    ctx.check(inv(r)["equation_nodes"] >= 1, "công thức vẫn OMML")


@case("G8", "G8.8", "ZIP độc hại: thoát thư mục (../), đường dẫn tuyệt đối, symlink, tỉ lệ nén cao → 400")
def g8_unsafe(ctx: Ctx) -> None:
    for key, fn in (("zip_slip", F.zip_slip), ("zip_abs", F.zip_absolute), ("zip_link", F.zip_symlink),
                    ("zip_bomb", F.zip_bomb)):
        z = ctx.fixture(key, lambda fn=fn: fn(ctx.fx))
        err = expect_error(lambda z=z: ctx.convert_file(z, "G8.8"))
        ctx.check(err is not None and err.http_status == 400, f"{z.name}: từ chối (400)",
                  (err.http_status, err.code, str(err)[:120]) if err else "KHÔNG bị từ chối")
    ctx.check(not (ctx.out.parent / "thoat_ra_ngoai.txt").exists() and not (ctx.jobs.parent.parent / "thoat_ra_ngoai.txt").exists(),
              "không có tệp nào bị ghi ra ngoài")


@case("G8", "G8.9", "Nhận diện theo NỘI DUNG: tệp TeX đặt đuôi .zip; ZIP nhiều PDF")
def g8_sniff(ctx: Ctx) -> None:
    r = ctx.convert("thuc_ra_la_tex.zip", F.TEX_G1.encode("utf-8"), "G8.9")
    ctx.check(r["input"]["kind"] in ("tex", "zip_tex") and "extension_mismatch" in issue_codes(r),
              "nhận là TeX + cảnh báo đuôi sai", (r["input"]["kind"], issue_codes(r)))
    p1 = ctx.fixture("vn_pdf", lambda: F.pdf_vietnamese(ctx.fx))
    p2 = ctx.fixture("two_col", lambda: F.pdf_two_columns(ctx.fx))
    data = zbytes({"a.pdf": p1.read_bytes(), "b.pdf": p2.read_bytes()})
    err = expect_error(lambda: ctx.convert("hai_pdf.zip", data, "G8.9"))
    ctx.check(err is not None and err.http_status == 409 and len(err.extra.get("candidates") or []) == 2,
              "ZIP 2 PDF → 409 + 2 ứng viên", (err.http_status, err.extra) if err else None)
    r2 = ctx.convert("hai_pdf.zip", data, "G8.9", main_file="b.pdf")
    ctx.check("Cột trái câu số 1" in doc_text(r2), "chọn rõ b.pdf → đúng nội dung")


# ═════════════════════════════ G9. thiếu công cụ / khối hỏng ═════════════════════════════
_DRIVER = r"""
import json, sys
sys.path.insert(0, sys.argv[1])
from backend.pipeline import convert_bytes, ConversionError
from backend.capabilities import capability_summary
args = json.loads(sys.argv[2])
data = open(args["path"], "rb").read()
try:
    rep = convert_bytes(args["name"], data, jobs_root=args["jobs"], **args["opts"])
    out = {"ok": True, "quality_status": rep["quality_status"], "layers": rep["layers"],
           "issues": rep["issues"], "blocks": (rep.get("blocks") or {}).get("items", []),
           "blocks_summary": rep["blocks_summary"], "docx": rep["output"]["docx"],
           "caps": capability_summary()["derived"]}
except ConversionError as e:
    out = {"ok": False, "code": e.code, "error": str(e)}
print("@@L2W@@" + json.dumps(out, ensure_ascii=True, default=str))
"""


def _run_without_tools(ctx: Ctx, path: Path, opts: dict[str, Any]) -> dict[str, Any]:
    empty = ctx.out / "_path_rong"
    empty.mkdir(exist_ok=True)
    env = {k: v for k, v in os.environ.items() if k not in ("PATH", "L2W_SOFFICE", "L2W_TESSERACT", "L2W_XELATEX")}
    env["PATH"] = str(empty)
    env["L2W_PATH_ONLY"] = "1"  # không tìm LibreOffice/Tesseract ở vị trí cài mặc định
    args = {"path": str(path), "name": path.name, "jobs": str(ctx.jobs), "opts": dict(opts, dpi=150)}
    p = subprocess.run([sys.executable, "-I", "-c", _DRIVER, str(REPO), json.dumps(args)], capture_output=True,
                       timeout=900, env=env)
    outp = p.stdout.decode("utf-8", "replace")
    if "@@L2W@@" not in outp:
        raise RuntimeError("tiến trình con lỗi: " + p.stderr.decode("utf-8", "replace")[-800:])
    res = json.loads(outp.split("@@L2W@@", 1)[1].strip().splitlines()[0])
    if res.get("docx"):
        ctx.docx_seen.append(("G9", res["docx"]))
    return res


@case("G9", "G9.1", "Không có TeX (PATH rỗng): hình cần chụp → không mất, báo rõ, không 'passed'")
def g9_no_tex(ctx: Ctx) -> None:
    src = ctx.write("g9_hinh_kho.tex", F.TEX_FIG_HARD)
    res = _run_without_tools(ctx, src, {"qa_pdf": False})
    ctx.check(res["ok"], "vẫn tạo được .docx", res.get("error"))
    if not res["ok"]:
        return
    ctx.check(not res["caps"]["tex_reference_compile"]["available"], "năng lực TeX báo KHÔNG có",
              res["caps"]["tex_reference_compile"])
    figs = [b for b in res["blocks"] if b.get("kind") == "figure"]
    ctx.check(len(figs) == 1, "khối hình vẫn được ghi nhận", [b.get("method") for b in figs])
    if figs:
        ctx.check(figs[0]["method"] != "vector_image" and (figs[0].get("status") in ("unverified", "unresolved", "review")
                                                           or figs[0]["method"] == "unresolved"),
                  "không giả vờ đã chụp ảnh; trạng thái chưa kiểm chứng/unresolved", figs[0])
    ctx.check(res["quality_status"] != "passed", "không 'passed'", res["quality_status"])
    ctx.info = {"figure": [(b["method"], b.get("status"), b.get("reason")) for b in figs]}


@case("G9", "G9.2", "Không có LibreOffice mà bật QA → nhóm E NOT_MEASURED (không tính là đạt)")
def g9_no_lo(ctx: Ctx) -> None:
    src = ctx.write("g9_qa.tex", F.TEX_G1)
    res = _run_without_tools(ctx, src, {"qa_pdf": True})
    ctx.check(res["ok"], "vẫn tạo được .docx", res.get("error"))
    if res["ok"]:
        l5 = [x for x in res["layers"] if x.get("layer") == 5]
        ctx.check(bool(l5) and l5[0]["status"] == "NOT_MEASURED", "L5 NOT_MEASURED", l5[0] if l5 else None)
        ctx.check(res["quality_status"] != "passed", "không 'passed'", res["quality_status"])


@case("G9", "G9.3", "Không có tesseract + trang scan → ảnh trang, báo ocr_unavailable")
def g9_no_ocr(ctx: Ctx) -> None:
    pdf = ctx.fixture("scan", lambda: F.pdf_scan(ctx.fx))
    res = _run_without_tools(ctx, pdf, {"qa_pdf": False})
    ctx.check(res["ok"], "vẫn tạo được .docx", res.get("error"))
    if res["ok"]:
        ctx.check(res["blocks_summary"]["by_method"].get("page_image") == 1, "giữ ảnh trang", res["blocks_summary"])
        ctx.check("ocr_unavailable" in [i["code"] for i in res["issues"]], "issue ocr_unavailable")


@case("G9", "G9.4", "Khối TikZ cố ý hỏng → khối báo lỗi, chữ trước/sau còn nguyên, không 'passed'")
def g9_broken_block(ctx: Ctx) -> None:
    r = ctx.convert("khoi_hong.tex", F.TEX_TIKZ_ERROR.encode("utf-8"), "G9.4")
    t = doc_text(r)
    ctx.check("Đoạn TRƯỚC khối hỏng phải còn nguyên." in t, "giữ đoạn trước")
    ctx.check("Đoạn SAU khối hỏng phải còn nguyên." in t, "giữ đoạn sau")
    fb = _fig_blocks(r)
    ctx.check(len(fb) == 1, "khối hình được ghi nhận", [b.get("method") for b in fb])
    if fb:
        ok_state = fb[0]["method"] == "unresolved" or fb[0].get("status") in ("unverified", "unresolved", "review")
        ctx.check(ok_state or fb[0]["method"] in ("vector_image", "raster_image"),
                  "trạng thái trung thực (unresolved/chưa kiểm chứng)", fb[0])
    ctx.check(r["quality_status"] != "passed", "không 'passed'", r["quality_status"])
    ctx.info = {"figure": [(b["method"], b.get("status"), (b.get("reason") or "")[:160]) for b in fb]}


# ═════════════════════════════ G10. phản ví dụ cho validator ═════════════════════════════
def _docx_with(ctx: Ctx, name: str, build: Callable[[Any], None]) -> Path:
    from docx import Document
    d = Document()
    build(d)
    p = ctx.out / "validator_docx" / name
    p.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(p))
    return p


@case("G10", "G10.1", "100 công thức nguồn, Word chỉ có 1 → L2 FAIL")
def g10_math(ctx: Ctx) -> None:
    from backend.omml_math import latex_to_omml_element
    tex = "\\documentclass{article}\\begin{document}\n" + "\n".join(f"Công thức ${i}x+{i}=0$." for i in range(1, 101)) \
        + "\n\\end{document}"

    def build(d):
        p = d.add_paragraph("Công thức ")
        p._p.append(latex_to_omml_element("1x+1=0"))
    p = _docx_with(ctx, "mot_cong_thuc.docx", build)
    lays = V.run_all_layers(tex, p, None)
    l2 = next(x for x in lays if x["layer"] == 2)
    ctx.check(l2["status"] == "FAIL", "L2 FAIL", l2["summary"])
    q = V.summarize_quality(lays, conversion_completed=True, package_ok=True)
    ctx.check(q["quality_status"] != "passed", "chất lượng không 'passed'", q["quality_status"])


@case("G10", "G10.2", "PDF 2 trang, Word chỉ có trang 1 → L1 không PASS")
def g10_missing_page(ctx: Ctx) -> None:
    from backend.pdf_reader import analyze_pdf
    pdf = ctx.fixture("two_pages", lambda: F.pdf_two_pages(ctx.fx))
    man = analyze_pdf(pdf, ctx.out / "validator_docx" / "_an", dpi=100)

    def build(d):
        for i in range(1, 7):
            d.add_paragraph(f"Trang một dòng {i}: nội dung kiểm tra số {i}.")
    p = _docx_with(ctx, "thieu_trang_cuoi.docx", build)
    lays = V.run_pdf_layers(man, p, None)
    l1 = next(x for x in lays if x["layer"] == 1)
    ctx.check(l1["status"] in ("FAIL", "WARN"), "L1 FAIL/WARN khi thiếu trang cuối", l1["summary"])
    ctx.info = {"L1": l1["status"] + ": " + l1["summary"]}


@case("G10", "G10.3", "Nguồn có hình, Word không có → L3 FAIL")
def g10_figure(ctx: Ctx) -> None:
    def build(d):
        d.add_paragraph("Hình đơn giản:")
        d.add_paragraph("Đoạn sau hình.")
    p = _docx_with(ctx, "mat_hinh.docx", build)
    lays = V.run_all_layers(F.TEX_FIG_SIMPLE, p, None)
    l3 = next(x for x in lays if x["layer"] == 3)
    ctx.check(l3["status"] == "FAIL", "L3 FAIL", l3["summary"])


@case("G10", "G10.4", "Bảng nguồn có dữ liệu, Word có bảng rỗng → L4 không PASS")
def g10_empty_table(ctx: Ctx) -> None:
    def build(d):
        d.add_table(rows=3, cols=3)
    p = _docx_with(ctx, "bang_rong.docx", build)
    lays = V.run_all_layers(F.TEX_MERGED, p, None)
    l4 = next(x for x in lays if x["layer"] == 4)
    ctx.check(l4["status"] in ("FAIL", "WARN"), "L4 FAIL/WARN", l4["summary"])


@case("G10", "G10.5", "Chỉ có ảnh Word, không có nguồn đối chiếu → NOT_MEASURED, similarity=None, không 'passed'")
def g10_no_reference(ctx: Ctx) -> None:
    word_pdf = ctx.fixture("vn_pdf", lambda: F.pdf_vietnamese(ctx.fx))  # "PDF render từ Word" giả lập
    info, pages = V.generate_visual_comparison(ctx.out / "validator_docx" / "_vis", None, word_pdf,
                                               compile_reference=False, mode="tex")
    ctx.check(info.get("measured") is False and info.get("similarity_percent") is None,
              "không đo → similarity_percent=None", info)
    l5 = V.run_layer5_visual(info)
    ctx.check(l5.status == "NOT_MEASURED", "L5 NOT_MEASURED", l5.summary)
    lays = [V.LayerResult(n, f"L{n}", "PASS", "ok").to_dict() for n in (1, 2, 3, 4)] + [l5.to_dict()]
    q = V.summarize_quality(lays, conversion_completed=True, package_ok=True)
    ctx.check(q["quality_status"] == "review_required", "4 PASS + 1 NOT_MEASURED → review_required", q["quality_status"])
    q2 = V.summarize_quality([V.LayerResult(n, f"L{n}", "PASS", "ok").to_dict() for n in (1, 2, 3, 4, 5)],
                             conversion_completed=True, package_ok=True)
    ctx.check(q2["quality_status"] == "passed", "đối chứng: 5 PASS → passed", q2["quality_status"])
    q3 = V.summarize_quality([V.LayerResult(n, f"L{n}", "SKIP", "bỏ qua").to_dict() for n in (1, 2, 3, 4, 5)],
                             conversion_completed=True, package_ok=True)
    ctx.check(q3["quality_status"] != "passed", "toàn SKIP → không 'passed'", q3["quality_status"])


# ═════════════════════════════ G11. gói DOCX ═════════════════════════════
@case("G11", "G11.1", "Mọi .docx tạo ra trong lượt chạy: quan hệ/media/ID hợp lệ")
def g11_package(ctx: Ctx) -> None:
    seen = {p for _c, p in ctx.docx_seen if Path(p).exists()}
    if not seen:
        raise Skip("chưa có .docx nào (chạy kèm các nhóm khác)")
    bad = {}
    for p in sorted(seen):
        v = validate_package(Path(p))
        if not v["ok"]:
            bad[Path(p).name] = v["errors"][:3]
    ctx.check(not bad, f"{len(seen)} tệp hợp lệ (content types, rels, r:embed, docPr id duy nhất, mở lại được)", bad)
    leaks = {Path(p).name: docx_inventory(Path(p))["marker_leaks"] for p in seen
             if docx_inventory(Path(p))["marker_leaks_total"]}
    ctx.check(not leaks, "không lộ chỗ đánh dấu nội bộ", leaks)
    ctx.info = {"docx_checked": len(seen)}


@case("G11", "G11.2", "Mở + lưu lại bằng python-docx: giữ công thức, bảng, ảnh, nhóm hình")
def g11_resave(ctx: Ctx) -> None:
    from docx import Document
    seen = [p for _c, p in ctx.docx_seen if Path(p).exists()]
    if not seen:
        raise Skip("chưa có .docx nào")
    keys = ("equation_nodes", "word_tables", "logical_images", "drawingml_groups", "media_parts")
    diff = {}
    for p in seen[:40]:
        a = docx_inventory(Path(p))
        q = ctx.out / "resave" / (Path(p).parent.parent.name + "_" + Path(p).name)
        q.parent.mkdir(parents=True, exist_ok=True)
        Document(p).save(str(q))
        b = docx_inventory(q)
        d = {k: (a[k], b[k]) for k in keys if a[k] != b[k]}
        if d:
            diff[q.name] = d
    ctx.check(not diff, f"{min(40, len(seen))} tệp: số đối tượng không đổi sau khi lưu lại", diff)


@case("G11", "G11.3", "Ca mẫu có sẵn trong repo vẫn chuyển được (không lỗi nội bộ, gói hợp lệ)")
def g11_baseline(ctx: Ctx) -> None:
    files = [REPO / "example.tex"] + sorted((REPO / "examples").glob("*.tex"))
    files = [f for f in files if f.exists()]
    if not files:
        raise Skip("repo không có example.tex / examples/*.tex")
    summary = {}
    for f in files:
        try:
            r = ctx.convert_file(f, "G11.3")
            summary[f.name] = {"ok": r["conversion_completed"],
                               "layers": [layer(r, n)["status"] for n in (1, 2, 3, 4)],
                               "blocks": r["blocks_summary"]["by_method"]}
            ctx.check(r["conversion_completed"], f"{f.name}: tạo .docx hợp lệ")
            ctx.check(inv(r)["marker_leaks_total"] == 0, f"{f.name}: không lộ marker")
        except ConversionError as exc:
            ctx.check(False, f"{f.name}: lỗi chuyển đổi", str(exc))
    ctx.info = summary


# ═════════════════════════════ G12. hai job cùng tên tệp ═════════════════════════════
@case("G12", "G12.1", "Hai job song song, CÙNG tên tệp, khác nội dung → không lẫn kết quả")
def g12_same_name(ctx: Ctx) -> None:
    srcs = {k: F.PREAMBLE + f"\\begin{{document}}\nMã riêng của job {k}: MARKER{k}{k}{k}.\n$x_{{{k}}}^2$\n\\end{{document}}\n"
            for k in ("A", "B")}
    res: dict[str, Any] = {}

    def run(k: str) -> None:
        try:
            res[k] = ctx.convert("de.tex", srcs[k].encode("utf-8"), "G12.1")
        except Exception as exc:  # pragma: no cover
            res[k] = exc

    th = [threading.Thread(target=run, args=(k,)) for k in srcs]
    for t in th:
        t.start()
    for t in th:
        t.join()
    ok = all(isinstance(res.get(k), dict) for k in srcs)
    ctx.check(ok, "cả hai job hoàn tất", {k: str(v)[:100] for k, v in res.items() if not isinstance(v, dict)})
    if not ok:
        return
    ctx.check(res["A"]["job_id"] != res["B"]["job_id"], "job_id khác nhau")
    ta, tb = doc_text(res["A"]), doc_text(res["B"])
    ctx.check("MARKERAAA" in ta and "MARKERBBB" not in ta, "job A chỉ có nội dung A", ta[:80])
    ctx.check("MARKERBBB" in tb and "MARKERAAA" not in tb, "job B chỉ có nội dung B", tb[:80])
    for k in srcs:
        inp = res[k]["input"]
        stored = Path(inp.get("stored") or inp.get("path") or "")
        if stored.exists():
            ctx.check(hashlib.sha256(stored.read_bytes()).hexdigest() == hashlib.sha256(srcs[k].encode()).hexdigest(),
                      f"tệp vào của job {k} giữ nguyên bytes")
        ctx.check(res[k].get("input_preserved") is True, f"job {k}: input_preserved")


@case("G12", "G12.2", "Hai PDF khác nhau cùng tên de.pdf → mỗi job đúng nội dung của mình")
def g12_same_name_pdf(ctx: Ctx) -> None:
    p1 = ctx.fixture("vn_pdf", lambda: F.pdf_vietnamese(ctx.fx))
    p2 = ctx.fixture("two_col", lambda: F.pdf_two_columns(ctx.fx))
    r1 = ctx.convert("de.pdf", p1.read_bytes(), "G12.2")
    r2 = ctx.convert("de.pdf", p2.read_bytes(), "G12.2")
    t1, t2 = doc_text(r1), doc_text(r2)
    ctx.check("Nguyễn Huệ" in t1 and "Cột trái" not in t1, "job 1 đúng nội dung")
    ctx.check("Cột trái" in t2 and "Nguyễn Huệ" not in t2, "job 2 đúng nội dung")
    ctx.check(r1["output"]["docx"] != r2["output"]["docx"], "đường dẫn kết quả khác nhau")


# ───────────────────────────── chạy test cũ ─────────────────────────────
def run_old_tests(out: Path) -> list[dict[str, Any]]:
    res = []
    for name in ("kiem_tra_smart_paste.py", "kiem_tra_universal.py"):
        f = REPO / name
        if not f.exists():
            res.append({"script": name, "status": "SKIP", "reason": "không có tệp"})
            continue
        t0 = time.time()
        try:
            p = subprocess.run([sys.executable, str(f)], cwd=str(REPO), capture_output=True, timeout=1200)
            tail = (p.stdout.decode("utf-8", "replace") + "\n" + p.stderr.decode("utf-8", "replace"))[-2500:]
            res.append({"script": name, "returncode": p.returncode, "seconds": round(time.time() - t0, 1),
                        "status": "PASS" if p.returncode == 0 else "FAIL", "output_tail": tail})
        except subprocess.TimeoutExpired:
            res.append({"script": name, "status": "FAIL", "reason": "quá thời gian"})
    (out / "test_cu.json").write_text(json.dumps(res, ensure_ascii=False, indent=2), encoding="utf-8")
    return res


# ───────────────────────────── báo cáo ─────────────────────────────
def environment() -> dict[str, Any]:
    from backend.capabilities import capability_summary
    caps = capability_summary()
    return {"python": sys.version.split()[0], "platform": platform.platform(),
            "tools": {t: tool(t) for t in ("xelatex", "pdflatex", "lualatex", "pdftocairo", "soffice", "tesseract", "gs")},
            "derived_capabilities": {k: v.get("available") for k, v in caps["derived"].items()},
            "preprocess_default": os.environ.get("L2W_PREPROCESS", "engine"),
            "tesseract_languages": caps.get("tesseract_languages")}


def write_markdown(out: Path, results: list[dict[str, Any]], meta: dict[str, Any], old: list[dict[str, Any]] | None) -> Path:
    cnt: dict[str, int] = {}
    for r in results:
        cnt[r["status"]] = cnt.get(r["status"], 0) + 1
    lines = ["# Báo cáo kiểm thử PDF / .tex / .zip → Word", "",
             f"- Lệnh: `{meta['command']}`", f"- Thời điểm: {meta['started']} — {meta['seconds']} s",
             f"- Môi trường: Python {meta['env']['python']}, {meta['env']['platform']}",
             "- Công cụ: " + ", ".join(f"{k}={'có' if v else 'KHÔNG'}" for k, v in meta["env"]["tools"].items()),
             f"- tesseract: {meta['env'].get('tesseract_languages')}",
             f"- Tiền xử lý TeX: {meta['env'].get('preprocess_default')} · smart_paste nạp được: "
             f"{meta['env']['derived_capabilities'].get('paste_smart')}", "",
             "**Tổng:** " + ", ".join(f"{k} {v}" for k, v in sorted(cnt.items())), "",
             "Ý nghĩa: PASS = mọi tiêu chí ghi trong ca đúng trên máy này với fixture này (không phải \"đúng 100%\"); "
             "SKIP = thiếu điều kiện, KHÔNG tính là đạt; LIMITATION = hành vi chưa đạt mức lý tưởng, ca kiểm việc báo "
             "cáo trung thực.", ""]
    group = None
    for r in results:
        if r["group"] != group:
            group = r["group"]
            lines += ["", f"## {group}", "", "| Ca | Nội dung | Kết quả | Thời gian | Ghi chú |", "|---|---|---|---|---|"]
        bad = [c["name"] for c in r.get("checks", []) if not c["ok"]]
        note = r.get("reason") or ("; ".join(bad[:3]) if bad else f"{sum(c['ok'] for c in r.get('checks', []))} tiêu chí đúng")
        note = str(note).replace("|", "/").replace("\n", " ")[:220]
        lines.append(f"| {r['id']} | {r['title']} | **{r['status']}** | {r['seconds']} s | {note} |")
    if old is not None:
        lines += ["", "## Test cũ của repo", ""]
        for o in old:
            lines.append(f"- `{o['script']}`: **{o['status']}** (mã thoát {o.get('returncode')}) "
                         + (o.get("reason") or ""))
    lines += ["", "Chi tiết từng tiêu chí: `ket_qua_kiem_tra.json`.", ""]
    p = out / "BAO_CAO_KIEM_TRA.md"
    p.write_text("\n".join(lines), encoding="utf-8")
    return p


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, default=REPO / "ket_qua_kiem_tra")
    ap.add_argument("--nhom", default="", help="vd. G1,G7 (mặc định: tất cả)")
    ap.add_argument("--ca", default="", help="vd. G2.4,G7.1")
    ap.add_argument("--qa", action="store_true", help="bật render Word→PDF cho ca QA (chậm, cần LibreOffice/Word)")
    ap.add_argument("--chay-test-cu", action="store_true")
    args = ap.parse_args()
    out = args.out.resolve()
    if out.exists():
        _force_rmtree(out)
    out.mkdir(parents=True)
    ctx = Ctx(out, args.qa)
    groups = {g.strip() for g in args.nhom.split(",") if g.strip()}
    only = {c.strip() for c in args.ca.split(",") if c.strip()}
    started = time.strftime("%Y-%m-%d %H:%M:%S")
    t_all = time.time()
    results = []
    # G11 cần chạy SAU các nhóm khác (kiểm mọi .docx đã tạo)
    order = [c for c in CASES if c[0] != "G11"] + [c for c in CASES if c[0] == "G11"]
    for group, cid, title, fn in order:
        if groups and group not in groups and not (group == "G11" and groups):
            continue
        if only and cid not in only:
            continue
        ctx.checks, ctx.info = [], {}
        t0 = time.time()
        rec: dict[str, Any] = {"group": group, "id": cid, "title": title}
        try:
            fn(ctx)
            rec["status"] = "PASS" if ctx.checks and all(c["ok"] for c in ctx.checks) else ("FAIL" if ctx.checks else "SKIP")
        except Skip as exc:
            rec["status"], rec["reason"] = "SKIP", str(exc)
        except Limitation as exc:
            failed = [c for c in ctx.checks if not c["ok"]]
            rec["status"] = "FAIL" if failed else "LIMITATION"
            rec["reason"] = str(exc)
        except Exception as exc:
            rec["status"] = "FAIL"
            rec["reason"] = f"lỗi khi chạy ca: {type(exc).__name__}: {exc}"
            rec["trace"] = traceback.format_exc()[-2500:]
        rec["seconds"] = round(time.time() - t0, 1)
        rec["checks"] = list(ctx.checks)
        if ctx.info:
            rec["info"] = _short(ctx.info, 600)
        results.append(rec)
        print(f"[{rec['status']:<10}] {cid:<6} {title[:70]} ({rec['seconds']} s)"
              + (f" — {rec.get('reason', '')[:120]}" if rec.get("reason") else ""), flush=True)
    old = run_old_tests(out) if args.chay_test_cu else None
    meta = {"command": " ".join([Path(sys.executable).name] + sys.argv), "started": started,
            "seconds": round(time.time() - t_all, 1), "env": environment()}
    summary: dict[str, int] = {}
    for r in results:
        summary[r["status"]] = summary.get(r["status"], 0) + 1
    (out / "ket_qua_kiem_tra.json").write_text(json.dumps({"meta": meta, "summary": summary, "results": results,
                                                            "old_tests": old}, ensure_ascii=False, indent=2,
                                                           default=str), encoding="utf-8")
    md = write_markdown(out, results, meta, old)
    print(f"\nTổng: {summary} — báo cáo: {md}")
    return 1 if summary.get("FAIL") else 0


if __name__ == "__main__":
    raise SystemExit(main())
